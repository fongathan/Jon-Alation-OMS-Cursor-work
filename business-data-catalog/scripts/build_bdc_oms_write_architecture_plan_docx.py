#!/usr/bin/env python3
"""Build high-level BDC-on-OMS write architecture plan (.docx)."""
from __future__ import annotations

from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

DIR = Path(__file__).resolve().parents[1]
OUT = DIR / "BDC-OMS-Business-Metadata-Write-Architecture-Plan.docx"

NAVY = RGBColor(0x1B, 0x2B, 0x44)
MUTED = RGBColor(0x64, 0x74, 0x8B)


def set_run_font(run, *, bold=False, size=11, color=None, italic=False) -> None:
    run.bold = bold
    run.italic = italic
    run.font.size = Pt(size)
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    if color:
        run.font.color.rgb = color


def add_heading(doc: Document, text: str, level: int = 1) -> None:
    p = doc.add_heading(text, level=level)
    for run in p.runs:
        run.font.name = "Calibri Light"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri Light")
        if level <= 2:
            run.font.color.rgb = NAVY


def add_para(doc: Document, text: str, *, bold: bool = False, italic: bool = False) -> None:
    p = doc.add_paragraph()
    run = p.add_run(text)
    set_run_font(run, bold=bold, italic=italic)


def add_bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        run = p.add_run(item)
        set_run_font(run)


def add_table(doc: Document, headers: list[str], rows: list[list[str]]) -> None:
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    for i, h in enumerate(headers):
        table.rows[0].cells[i].text = h
        for para in table.rows[0].cells[i].paragraphs:
            for run in para.runs:
                set_run_font(run, bold=True, size=10, color=NAVY)
    for r_idx, row in enumerate(rows):
        for c_idx, val in enumerate(row):
            table.rows[r_idx + 1].cells[c_idx].text = val
            for para in table.rows[r_idx + 1].cells[c_idx].paragraphs:
                for run in para.runs:
                    set_run_font(run, size=10)
    doc.add_paragraph()


def build() -> Path:
    doc = Document()
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    title = doc.add_paragraph()
    title.alignment = WD_PARAGRAPH_ALIGNMENT.LEFT
    tr = title.add_run("Business Data Catalog on OMS\nBusiness Metadata Write Architecture Plan")
    set_run_font(tr, bold=True, size=22, color=NAVY)

    sub = doc.add_paragraph()
    sr = sub.add_run(
        f"High-level architecture for editable business metadata when OMS is the shared backbone. "
        f"Draft — {date.today().strftime('%B %d, %Y')}."
    )
    set_run_font(sr, size=10, color=MUTED, italic=True)

    doc.add_paragraph()

    add_heading(doc, "1. Purpose and scope", 1)
    add_para(
        doc,
        "This document describes how to plan and implement write capability for the Business Data Catalog (BDC) "
        "when Operational Metadata Store (OMS) is the system backbone. OMS today primarily aggregates technical "
        "metadata from connected platforms (Snowflake, Databricks Unity Catalog, and similar sources). The BDC "
        "replaces Alation for business-facing catalog experiences; the existing OMS UI continues to serve "
        "technical and operational stakeholders. Both UIs must read from the same OMS-backed metadata fabric; "
        "only the presentation and workflows differ.",
    )
    add_bullets(
        doc,
        [
            "In scope: architecture principles, metadata layering, write paths, APIs, phasing, and ownership.",
            "In scope for MVP business writes: table (and asset) descriptions, steward assignment, and closely related stewardship fields (e.g., trust/endorsement, domain tagging where program-owned).",
            "Out of scope for this plan: detailed schema DDL, vendor-specific push-back to Snowflake comments, and full Alation parity (glossary workflows, articles, custom fields beyond MVP).",
        ],
    )

    add_heading(doc, "2. Context: from Alation to OMS + BDC", 2)
    add_bullets(
        doc,
        [
            "Alation is a standalone catalog service with its own metadata store and write model.",
            "OMS is an aggregator: connectors crawl sources on a schedule and populate OMS with technical metadata (schemas, tables, columns, lineage signals, operational attributes).",
            "BDC is a business-first UI slice on top of the same OMS identity graph—not a second catalog database for technical facts.",
            "The gap: OMS was not originally designed as the authoritative store for steward-edited business metadata; BDC requires that capability to replace Alation for stewardship workflows.",
        ],
    )

    add_heading(doc, "3. UX intent: two slices, one backbone", 2)
    add_table(
        doc,
        ["Surface", "Primary audience", "What they need from metadata"],
        [
            [
                "BDC (Navigator / business catalog)",
                "Stewards, analysts, product, governance",
                "Descriptions, ownership, domain, trust, policy context, search tuned for discovery",
            ],
            [
                "OMS UI (current operational UI)",
                "Engineers, platform, data ops",
                "Lineage, jobs, storage, profiling, SQL history, connector health, technical detail",
            ],
            [
                "Shared requirement",
                "Both",
                "Same asset identity, consistent search results, deep links between surfaces, unified audit of catalog actions",
            ],
        ],
    )
    add_para(
        doc,
        "Architecture should not fork asset identity. Each table, column, or catalog object must have a stable OMS identifier "
        "(for example, a canonical URI or FQN registered during ingestion). BDC and OMS UI both resolve that identifier; "
        "writes attach business metadata to that identifier rather than inventing parallel keys.",
    )

    add_heading(doc, "4. The core problem: writes on a read-oriented aggregator", 1)
    add_para(
        doc,
        "Connector-driven ingestion is inherently source-authoritative for technical fields: column types, locations, "
        "Unity Catalog tags from the lakehouse, etc. If stewards edit those fields in BDC, the next crawl may "
        "overwrite or conflict with OMS. Conversely, if OMS never persists steward input, BDC cannot replace Alation.",
    )
    add_para(doc, "Design goal:", bold=True)
    add_bullets(
        doc,
        [
            "Separate system-of-record by metadata class (technical vs business), not by UI.",
            "Make merge rules explicit at read time and enforce them at write time.",
            "Keep OMS as the integration hub: one search/index and one identity graph, with an extensible business overlay.",
        ],
    )

    add_heading(doc, "5. Recommended pattern: layered metadata on OMS", 1)
    add_para(
        doc,
        "Adopt a three-layer model. This fits OMS’s connector architecture without pretending that Snowflake is "
        "the store for steward narratives or program-owned assignments.",
    )

    add_heading(doc, "5.1 Layer A — Source technical metadata (read-only in BDC)", 2)
    add_bullets(
        doc,
        [
            "Origin: Snowflake, Databricks, Delta/Hive, and other registered connections.",
            "Population: existing OMS ingest pipelines (scheduled or event-driven sync).",
            "Characteristics: high volume, volatile, owned by the source platform.",
            "BDC behavior: display for context; do not offer edit controls except where the platform is explicitly the system of record (rare for MVP).",
        ],
    )

    add_heading(doc, "5.2 Layer B — Business metadata overlay (OMS-owned, writable)", 2)
    add_bullets(
        doc,
        [
            "Origin: steward and catalog-admin actions via BDC (and optionally OMS UI later).",
            "Population: new OMS capability—Business Metadata Service (name TBD)—persisting overlay records keyed by OMS asset ID.",
            "MVP fields: business title (optional), business description, assigned steward(s), domain / 360 assignment, trust or endorsement flag, deprecation status and reason.",
            "Optional later: glossary term links, custom fields, workflow state, Alation-import provenance.",
            "Characteristics: lower volume, audit-heavy, versioned, authorization-sensitive.",
        ],
    )

    add_heading(doc, "5.3 Layer C — Merged catalog view (served to both UIs)", 2)
    add_bullets(
        doc,
        [
            "Origin: computed at query time or materialized in OMS search/index (preferred for performance).",
            "Rules: technical fields from Layer A; business fields from Layer B when present; documented fallbacks (e.g., show source comment until a business description exists).",
            "Both BDC and OMS UI consume the same merge API so stewards and engineers never see divergent truth for the same asset.",
        ],
    )

    add_heading(doc, "6. Write path (logical architecture)", 1)
    add_para(doc, "Target flow for a steward saving a table description or steward assignment:")
    add_bullets(
        doc,
        [
            "1. BDC UI sends a patch to a Catalog Metadata Write API (REST or GraphQL—align with OMS platform standards).",
            "2. API validates identity (asset exists in OMS), authorization (steward role, domain scope), and field-level policy (which fields are editable in BDC).",
            "3. Service persists an overlay record (create or new version) in the OMS business metadata store—not in the connector snapshot tables.",
            "4. Service emits an audit event and optional domain event (for search reindex, notifications, Data Governance Portal work queue).",
            "5. Search/index pipeline updates the merged document for that asset.",
            "6. BDC and OMS UI read the updated merged view on next fetch.",
        ],
    )
    add_para(
        doc,
        "Important: writes must not mutate connector ingestion tables in place. Ingest jobs should continue to "
        "refresh Layer A independently. Layer B references Layer A by stable ID only.",
    )

    add_heading(doc, "7. Architecture options considered", 1)
    add_table(
        doc,
        ["Option", "Summary", "Pros", "Cons", "Recommendation"],
        [
            [
                "A. Extend OMS with native overlay store",
                "OMS team adds business_metadata tables + APIs alongside existing graph",
                "Single platform, unified auth, one ops footprint",
                "Requires OMS roadmap capacity; careful separation from ingest schema",
                "Preferred for enterprise backbone",
            ],
            [
                "B. Sidecar metadata microservice",
                "Separate service DB; registers with OMS identity service",
                "Faster BDC pilot if OMS core is constrained",
                "Two systems to operate; merge/search integration work",
                "Acceptable as Phase 0 only with explicit sunset to A",
            ],
            [
                "C. Write back to source only",
                "Steward edits Snowflake/UC comments",
                "No new OMS store",
                "Slow, permission-heavy, inconsistent across sources; poor Alation replacement",
                "Not recommended as primary model",
            ],
            [
                "D. Retain Alation as write store",
                "BDC reads Alation for business fields",
                "Short-term migration ease",
                "Defeats decommission goal; dual truth",
                "Reject except temporary migration bridge",
            ],
        ],
    )
    add_para(
        doc,
        "Program recommendation: pursue Option A as the target state. If OMS delivery bandwidth is limited, "
        "Option B may host MVP overlay storage only if it uses OMS asset IDs, the same auth model, and a "
        "committed merge contract consumed by OMS search—never a BDC-private database.",
    )

    add_heading(doc, "8. Data model (conceptual, MVP)", 1)
    add_para(doc, "Overlay entities (illustrative—not implementation DDL):")
    add_bullets(
        doc,
        [
            "business_asset_overlay: oms_asset_id (PK/FK), object_type, business_title, business_description, updated_by, updated_at, source_of_record = 'bdc'.",
            "steward_assignment: oms_asset_id, steward_principal_id (or group), role (primary/backup), domain_id, effective_dates.",
            "trust_signal: oms_asset_id, level (endorsed/deprecated/warning), reason, actor, timestamp.",
            "field_provenance (optional): per-field source (ingest vs overlay vs import), supports conflict UI.",
        ],
    )
    add_para(
        doc,
        "Column-level descriptions can follow the same overlay pattern keyed by oms_column_id. "
        "Defer wide custom-field matrices until enterprise taxonomy is stable.",
    )

    add_heading(doc, "9. Read / merge rules", 1)
    add_table(
        doc,
        ["Field class", "System of record", "On conflict"],
        [
            ["Column type, database, schema, technical tags from UC", "Source ingest (Layer A)", "Ingest wins; surface “changed at source” in UI"],
            ["Business description, steward, domain, trust", "Overlay (Layer B)", "Overlay wins; show last editor and timestamp"],
            ["Description empty in overlay", "Fallback", "Optional: show source comment as read-only hint, not editable as business description"],
            ["Deprecated in overlay", "Overlay", "Display prominently in BDC and OMS merged view"],
        ],
    )

    add_heading(doc, "10. Security, governance, and audit", 1)
    add_bullets(
        doc,
        [
            "Reuse enterprise SSO; map IdP groups to catalog roles (viewer, contributor, steward, catalog admin) as defined in BDC admin requirements.",
            "Enforce domain-scoped stewardship: stewards write only within assigned 360 domains unless enterprise admin.",
            "Immutable audit log for every overlay change (who, what, before/after, correlation id).",
            "Align classification and policy fields with Data Governance Portal as read-mostly in MVP unless policy team owns writes elsewhere.",
            "Rate-limit and validate bulk edit APIs to protect index and connector stability.",
        ],
    )

    add_heading(doc, "11. Integration with existing OMS infrastructure", 1)
    add_bullets(
        doc,
        [
            "Connectors: unchanged responsibility for Layer A; document that ingest must not delete overlay rows when assets are renamed—use identity resolution / alias table.",
            "Search: extend OMS index mapping to include merged business fields for BDC-oriented facets (domain, trust, has description).",
            "Lineage and operational metadata: remain OMS-native; BDC deep-links to OMS UI for technical drill-down.",
            "Migration from Alation: one-time import maps Alation object keys to OMS asset IDs; load into Layer B with provenance alation_import; stewards reconcile conflicts in BDC.",
            "Optional Phase 2: selective publish of business description to Snowflake COMMENT (async job, steward opt-in, failure visible in UI).",
        ],
    )

    add_heading(doc, "12. Phased delivery plan", 1)
    add_table(
        doc,
        ["Phase", "Objective", "Deliverables"],
        [
            [
                "Phase 0 — Foundation",
                "Identity + read merge",
                "Stable asset IDs exposed to BDC; read API returns merged view; no user writes",
            ],
            [
                "Phase 1 — MVP writes",
                "Replace Alation core stewardship",
                "Write API for description + steward; audit; BDC edit UX; search reflects updates within SLA",
            ],
            [
                "Phase 2 — Trust & workflow",
                "Governance parity",
                "Endorse/deprecate, domain bulk assignment, steward queue integration with DGP",
            ],
            [
                "Phase 3 — Enrichment",
                "Scale and optional sync",
                "Glossary links, custom fields, optional push to source comments, analytics on curation",
            ],
        ],
    )

    add_heading(doc, "13. Workstreams and ownership (RACI sketch)", 1)
    add_table(
        doc,
        ["Workstream", "OMS platform", "BDC product/UI", "Data governance", "Source platform teams"],
        [
            ["Overlay storage + write API", "R/A", "C", "C", "I"],
            ["Merge/read API + search index", "R/A", "C", "I", "I"],
            ["BDC edit UX + validation", "C", "R/A", "C", "I"],
            ["AuthZ / domain scope", "R", "C", "A", "I"],
            ["Alation migration mapping", "C", "C", "R/A", "I"],
            ["Connector ingest (Layer A)", "R/A", "I", "I", "C"],
        ],
    )
    add_para(doc, "R = Responsible, A = Accountable, C = Consulted, I = Informed.", italic=True)

    add_heading(doc, "14. Risks and open decisions", 1)
    add_bullets(
        doc,
        [
            "Identity stability: asset renames and orphan overlays—need alias/merge strategy before Phase 1.",
            "OMS team capacity vs BDC timeline—resolve Option A vs time-boxed Option B early.",
            "Whether OMS UI will expose edit controls for business fields or remain read-only with link to BDC.",
            "Search SLA after writes (near-real-time vs batch reindex).",
            "Enterprise vs domain-specific required fields (Ads/Atlas namespaces) without fragmenting overlay schema.",
            "Legal/compliance retention for audit and deprecated assets.",
        ],
    )

    add_heading(doc, "15. Success criteria", 1)
    add_bullets(
        doc,
        [
            "A steward can assign themselves and publish a business description on an OMS-known table entirely within BDC; change appears in merged API and search without a connector run.",
            "Technical metadata refresh from Snowflake does not erase business overlay.",
            "OMS UI and BDC show the same steward and description for the same asset ID.",
            "Audit trail satisfies governance review for Alation decommission.",
            "Migration imports ≥ agreed percentage of Alation descriptions into Layer B with steward sign-off workflow.",
        ],
    )

    add_heading(doc, "16. Summary recommendation", 1)
    add_para(
        doc,
        "Treat OMS as the backbone for identity, technical metadata, search, and operations—but add an "
        "OMS-owned business metadata overlay with explicit merge semantics rather than overloading connector "
        "tables or writing only to external sources. BDC becomes the primary write experience for business "
        "stakeholders; the current OMS UI remains the primary technical experience, both consuming the same "
        "merged catalog API. Deliver MVP writes for description and steward assignment first; expand trust, "
        "workflow, and glossary after the overlay and audit foundation is proven.",
    )

    doc.add_paragraph()
    p = doc.add_paragraph()
    run = p.add_run(
        "Related workspace references: Product-Brief-Alation-to-OMS-Migration.md; "
        "business-data-catalog/BDC-OMS-integrated-vs-on-top-decision.md; "
        "business-data-catalog/Alation-OMS-626-Capability-Comparison-Chart.md."
    )
    set_run_font(run, italic=True, size=10, color=MUTED)

    doc.save(OUT)
    return OUT


if __name__ == "__main__":
    path = build()
    print(f"Wrote {path}")
