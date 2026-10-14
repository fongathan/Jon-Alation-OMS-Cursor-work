#!/usr/bin/env python3
"""Build Business Data Catalog PRFAQ (.docx).

Amazon-style working-backwards document: future press release + internal FAQ,
grounded in Navigator / BDC program artifacts.
"""
from __future__ import annotations

from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

DIR = Path(__file__).resolve().parents[1]
OUT = DIR / "PRFAQ-Business-Data-Catalog.docx"

# Palette — aligned with Navigator / BDC deck styling
NAVY = RGBColor(0x1E, 0x2B, 0x52)
INDIGO = RGBColor(0x2C, 0x47, 0x8C)
BLUE = RGBColor(0x2D, 0x6C, 0xE0)
MUTED = RGBColor(0x6B, 0x77, 0x8C)
TEAL = RGBColor(0x10, 0x9A, 0xA6)


def set_run_font(run, *, bold=False, size=11, color=None, italic=False) -> None:
    run.bold = bold
    run.italic = italic
    run.font.size = Pt(size)
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    if color:
        run.font.color.rgb = color


def add_heading(doc: Document, text: str, level: int = 1) -> None:
    p = doc.add_heading(text, level=level)
    for run in p.runs:
        run.font.name = "Calibri Light"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri Light")
        if level <= 2:
            run.font.color.rgb = NAVY


def add_para(
    doc: Document,
    text: str = "",
    *,
    bold: bool = False,
    size: int = 11,
    color=None,
    italic: bool = False,
    align=None,
) -> None:
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    if text:
        run = p.add_run(text)
        set_run_font(run, bold=bold, size=size, color=color, italic=italic)


def add_bullets(doc: Document, items: list[str], size: int = 11) -> None:
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        run = p.add_run(item)
        set_run_font(run, size=size)


def add_table(doc: Document, headers: list[str], rows: list[list[str]]) -> None:
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    hdr_cells = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr_cells[i].text = h
        for para in hdr_cells[i].paragraphs:
            for run in para.runs:
                set_run_font(run, bold=True, size=10, color=NAVY)
    for r_idx, row in enumerate(rows):
        cells = table.rows[r_idx + 1].cells
        for c_idx, val in enumerate(row):
            cells[c_idx].text = val
            for para in cells[c_idx].paragraphs:
                for run in para.runs:
                    set_run_font(run, size=10)
    doc.add_paragraph()


def add_block_quote(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.35)
    run = p.add_run(text)
    set_run_font(run, italic=True, size=11, color=INDIGO)


def add_faq(doc: Document, question: str, answer_paras: list[str]) -> None:
    q = doc.add_paragraph()
    qr = q.add_run(question)
    set_run_font(qr, bold=True, size=11, color=NAVY)
    for text in answer_paras:
        add_para(doc, text)


def build() -> Path:
    doc = Document()
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    # --- Cover / document control ---
    label = doc.add_paragraph()
    label.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    lr = label.add_run("PRFAQ")
    set_run_font(lr, bold=True, size=12, color=MUTED)

    title = doc.add_paragraph()
    title.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    tr = title.add_run("Business Data Catalog")
    set_run_font(tr, bold=True, size=22, color=NAVY)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    sr = subtitle.add_run("Navigator · Data Governance · DEEPT")
    set_run_font(sr, size=13, color=INDIGO)

    meta = doc.add_paragraph()
    meta.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    mr = meta.add_run(f"Version 0.1 · Draft · {date.today().isoformat()}")
    set_run_font(mr, size=10, color=MUTED)

    doc.add_paragraph()

    add_para(
        doc,
        "This document uses the Amazon PRFAQ (Press Release + Frequently Asked Questions) "
        "format to work backwards from customer value. Part I is a future-dated press release "
        "describing the MVP launch. Part II answers the questions executives, engineers, and "
        "stewards will ask. Part III summarizes MVP capabilities and the business outcomes they enable.",
        color=MUTED,
        size=10,
    )

    doc.add_paragraph()

    # --- Part I: Press Release ---
    add_heading(doc, "Part I — Press Release (future-dated)", 1)

    headline = doc.add_paragraph()
    headline.alignment = WD_PARAGRAPH_ALIGNMENT.LEFT
    hr = headline.add_run(
        "Disney Entertainment & ESPN Technology Launches Business Data Catalog — "
        "Governed Discovery for Every Data Consumer"
    )
    set_run_font(hr, bold=True, size=16, color=NAVY)

    subhead = doc.add_paragraph()
    sr2 = subhead.add_run(
        "New Navigator experience on OMS replaces Alation for business users — "
        "find trusted data, understand what it means, and know who owns it, without tribal knowledge."
    )
    set_run_font(sr2, italic=True, size=12, color=INDIGO)

    add_para(doc, "Burbank, Calif. — October 1, 2026", bold=True, size=11)

    add_para(
        doc,
        "Today Disney Entertainment & ESPN Technology (DEEPT) announced the general availability "
        "of the Business Data Catalog (BDC), the business-first metadata experience within the "
        "Navigator data governance product family. Built on the Operational Metadata Store (OMS), "
        "BDC gives analysts, business consumers, and data stewards a single place to search, browse, "
        "and trust the datasets that power decisions across streaming, parks, studios, and enterprise "
        "analytics — on the same day the organization retires its Alation contract.",
    )

    add_para(
        doc,
        "For years, teams relied on Alation to answer a simple question: \"Can I use this data, and "
        "what does it actually mean?\" As Disney migrated technical metadata into OMS, business users "
        "still needed a surface optimized for governed discovery — not engineering consoles, not "
        "spreadsheets, and not Slack threads. BDC closes that gap.",
    )

    add_para(
        doc,
        "With BDC, a marketing analyst can search for a subscriber metric, read the steward-endorsed "
        "definition linked to the authoritative table, see who owns it, check sensitivity "
        "classifications, and follow lineage before building a dashboard. A data steward can endorse "
        "gold paths, flag deprecated assets, and keep the business glossary connected to live "
        "inventory. Engineers retain full technical depth in OMS — with one-click handoff when they "
        "need SQL, profiling, or pipeline context.",
    )

    add_block_quote(
        doc,
        "\"We did not set out to build another metadata store,\" said a DEEPT Data Governance leader. "
        "\"We set out to make governed metadata usable for the people who actually consume it. "
        "626 improves what we collect. OMS stores and serves it. BDC is how business users "
        "experience it — and how we prove Alation exit without losing trust.\"",
    )

    add_block_quote(
        doc,
        "\"I used to open three tools and ping two stewards before I trusted a number,\" said a "
        "business intelligence analyst in the DNA organization. \"Now I start in BDC, see the gold "
        "definition, the owner, and the policy context on one screen. That is hours back every week.\"",
    )

    add_heading(doc, "What ships at MVP launch", 2)
    add_bullets(
        doc,
        [
            "Unified search and hierarchical browse across data sources, schemas, tables, and columns",
            "Rich asset detail — business and technical descriptions, ownership, tags, and classifications",
            "Business glossary with terms, definitions, synonyms, and links to catalog assets",
            "Trust signals — endorse, warn, and deprecate — so consumers know intended use",
            "Lineage views for impact analysis and compliance traceability (depth per OMS graph coverage)",
            "Documentation articles linked to catalog objects for domain context and how-to knowledge",
            "Open in OMS — deep link to full technical metadata for engineers without duplicating the platform UI",
            "Enterprise SSO and role-aware access for governance-ready operations",
        ],
    )

    add_heading(doc, "BDC is one part of Navigator — not the whole story", 2)
    add_para(
        doc,
        "BDC answers \"where is the governed data and what does it mean?\" The broader Navigator "
        "program also includes the Data Governance Portal (posture, work queues, and evidence across "
        "Privacy, Platform, Security, and Regulatory dimensions), the Policy Drive Governance "
        "Framework (classifications, policies, and metrics in PULS), and deep links to specialist "
        "systems such as Kronos for business inventory. MVP prioritizes catalog parity for Alation "
        "exit; the governance command center ships in a subsequent phase.",
    )

    add_para(
        doc,
        "BDC is available now to in-scope DEEPT users via the Navigator left navigation. "
        "Stewards and analysts should bookmark the Business Data Catalog entry point; engineers "
        "continue to use OMS for platform work and open BDC when business context matters.",
        bold=True,
    )

    add_para(
        doc,
        "About Navigator: Navigator is DEEPT's unified data governance product family — one program "
        "with clear boundaries between what Data Governance builds (experiences and workflows) and "
        "what platform teams operate (OMS, SISU, PULS). Learn more at the internal Navigator program hub.",
        size=10,
        color=MUTED,
    )

    doc.add_paragraph()

    # --- Part II: FAQ ---
    add_heading(doc, "Part II — Frequently Asked Questions", 1)

    add_faq(
        doc,
        "What is the Business Data Catalog?",
        [
            "The Business Data Catalog (BDC) is the in-house, business-first metadata experience on "
            "top of OMS. It is how analysts, business consumers, and stewards discover assets, read "
            "business meaning (definitions, ownership, classifications, quality and trust signals), "
            "and follow lineage — without navigating engineering-oriented platform chrome.",
            "BDC is not a second metadata store. It reads OMS APIs. Data 626 and MCI improve metadata "
            "supply; OMS stores and serves it; BDC is the governed experience layer for the personas "
            "who depended on Alation.",
        ],
    )

    add_faq(
        doc,
        "We talk about BDC — but what else is in Navigator?",
        [
            "Navigator is the umbrella data governance product family. BDC is the catalog spine, but "
            "the program is intentionally broader:",
        ],
    )
    add_table(
        doc,
        ["Component", "What it does", "MVP status"],
        [
            [
                "Business Data Catalog (BDC)",
                "Governed discovery, glossary, stewardship UX, trust signals on OMS assets",
                "MVP — Oct 2026 Alation exit target",
            ],
            [
                "Navigator (Governance Portal)",
                "Posture overview, work queues, evidence export across Privacy, Platform, Security, Regulatory",
                "Phase 2 — after BDC MVP",
            ],
            [
                "Policy Drive Governance Framework",
                "Classifications, policies, metrics (PULS); Access Manager alignment; Gov App backend APIs",
                "Phased — read paths in MVP where modeled; workflow depth in Phase 2",
            ],
            [
                "Kronos (business inventory)",
                "Authoritative business inventory on SISU Postgres — purpose, obligations, regulatory context",
                "Deep-linked from Navigator Regulatory — not rebuilt in BDC",
            ],
            [
                "OMS (Foundations)",
                "Operational metadata store, graph, connectors, APIs — system of record for technical metadata",
                "Platform-owned; BDC consumes, does not duplicate",
            ],
        ],
    )

    add_faq(
        doc,
        "Why not just use OMS?",
        [
            "OMS optimizes for engineers and platform integrators: ingestion, graph completeness, SQL, "
            "profiling, and API stability. BDC optimizes for governed discovery: business language, "
            "glossary links, endorse/deprecate signals, and persona defaults that match how ~500+ "
            "Alation users actually work.",
            "Architecture compromise: one metadata API, one store, two shells (UX only) plus Open in OMS "
            "for full technical detail. Foundations owns store and serve. Data Governance owns the "
            "business-metadata experience backed by OMS.",
        ],
    )

    add_faq(
        doc,
        "Where does 626 fit?",
        [
            "626 and adjacent MCI work improve how metadata is produced, scanned, and curated at the "
            "source. BDC does not own scanners or ingestion pipelines. As 626 raises coverage and "
            "definition quality upstream, BDC surfaces that improvement to consumers — with provenance "
            "so stewards and users know what was migrated vs. human-curated.",
        ],
    )

    add_faq(
        doc,
        "Who is the MVP for?",
        [
            "Primary: business/DNA consumers and analysts who need to find and trust metrics and "
            "datasets quickly.",
            "Secondary: data stewards (thin endorse/reconcile flows) and privacy/compliance reviewers "
            "(read-heavy, role-filtered views of classifications and policy context).",
            "Handoff personas: engineers and DREs use OMS UI and APIs respectively; BDC links to OMS "
            "rather than replacing it.",
        ],
    )

    add_faq(
        doc,
        "What is in MVP vs. explicitly out?",
        [
            "In scope for MVP: search and browse, asset detail with business metadata, glossary, "
            "classifications read, trust flags, ownership, articles, lineage (as OMS graph allows), "
            "SSO, read API for automation, and Open in OMS.",
            "Out of scope for MVP (Phase 2+): DSAR and tracker consoles, DRE remediation UI, full "
            "governance command center and unified exception queue, in-catalog SQL, social/admin "
            "features, and production-strength bulk edit / catalog sets (minimum slice only if "
            "steward validation demands it for exit).",
        ],
    )

    add_faq(
        doc,
        "What business value does MVP deliver?",
        [
            "MVP is not a feature checklist — it is a set of outcomes that protect migration, adoption, "
            "and audit readiness:",
        ],
    )
    add_table(
        doc,
        ["Capability", "Business value", "Who benefits"],
        [
            [
                "Unified search & browse",
                "Answer \"what data exists?\" in minutes, not days of Slack archaeology",
                "Analysts, business consumers",
            ],
            [
                "Glossary linked to assets",
                "One vocabulary for metrics and dimensions — fewer conflicting dashboard definitions",
                "Stewards, analysts, leadership",
            ],
            [
                "Trust signals (endorse / warn / deprecate)",
                "Consumers know which assets are gold vs. legacy — reduces bad decisions and rework",
                "Consumers, stewards, compliance",
            ],
            [
                "Ownership & stewardship visible",
                "Clear accountability — faster escalation and fewer orphaned datasets",
                "Stewards, program leads",
            ],
            [
                "Classifications & tags (read)",
                "Sensitivity and compliance scope visible at point of use — supports privacy-by-design",
                "Privacy, compliance, consumers",
            ],
            [
                "Lineage entry points",
                "Impact analysis and audit traceability without exporting to spreadsheets",
                "Analysts, engineers, compliance",
            ],
            [
                "Articles & domain documentation",
                "Institutional knowledge survives people changes and tool migration",
                "Stewards, new hires, domain teams",
            ],
            [
                "Open in OMS + read APIs",
                "Same truth for UI and automation — no dual-write, no forked inventory",
                "Engineers, DRE, platform",
            ],
            [
                "Alation exit on schedule",
                "Contract risk removed; spend redirected to owned capability on OMS",
                "Executive sponsors, finance, IT",
            ],
        ],
    )

    add_faq(
        doc,
        "What comes after MVP?",
        [
            "Phase 2 (~Q2 FY27 horizon): operate at scale — stewardship requests and lightweight "
            "workflows, policy surfacing in catalog, production bulk edit and catalog sets, usage "
            "signals, AI-readiness attributes (quality tier, lineage completeness, documented "
            "limitations), and expanded domain coverage.",
            "Phase 3 (~Q4 FY27+): differentiate — governance command center shell (unified posture and "
            "exceptions across privacy pillars), semantic layer integration for single governed metric "
            "truth, and deeper automation (recommendations, policy-driven metadata) once trust and "
            "attribution rules are clear.",
        ],
    )

    add_faq(
        doc,
        "How will we know MVP succeeded?",
        [
            "Steward UAT sign-off on parity paths; parallel-run adoption vs. Alation for in-scope domains.",
            "Validated glossary and asset links post-migration (% critical objects with business metadata).",
            "Reduced \"where is X?\" escalations to stewards and data platform teams.",
            "Weekly active users among analyst and consumer personas meeting Alation-exit targets.",
            "Evidence that engineers can still reach full technical context via OMS without friction.",
        ],
    )

    add_faq(
        doc,
        "Who builds what?",
        [
            "Data Governance owns BDC product, UX, and adoption for business-metadata experience.",
            "OMS Foundations owns operational metadata store, connectors, graph, and APIs.",
            "626 / MCI programs own metadata supply-side improvement.",
            "SISU partners on Gov App backend, PULS connectivity, and Navigator workflow services.",
            "Kronos remains the system of record for business inventory — linked, not duplicated.",
        ],
    )

    doc.add_paragraph()

    # --- Part III: MVP summary ---
    add_heading(doc, "Part III — MVP at a glance", 1)

    add_para(
        doc,
        "North star (Oct 1, 2026): Find, trust, and cite governed metadata on OMS via Navigator — "
        "fed by 626/MCI, stored in OMS, experienced through BDC.",
        bold=True,
        color=TEAL,
    )

    add_heading(doc, "Stack positioning", 2)
    add_table(
        doc,
        ["Layer", "Role", "Owner"],
        [
            ["626 / MCI", "Improve metadata supply — scanning, curation, definitions at source", "Metadata programs"],
            ["OMS", "Store, connect, and serve operational & technical metadata", "Data Foundations"],
            ["BDC / Navigator", "Governed experience — discovery, glossary, stewardship, policy surfacing", "Data Governance"],
        ],
    )

    add_heading(doc, "MVP success criteria (acceptance posture)", 2)
    add_bullets(
        doc,
        [
            "A business consumer can search, open an asset, read business description and owner, "
            "see trust status and classifications, and cite the object in a planning doc — without Alation.",
            "A steward can endorse or deprecate an asset and link glossary terms without engineering support.",
            "An engineer can open the same object in OMS from BDC and see full technical lineage and SQL context.",
            "Migration validation: critical glossary links and custom business fields preserved for in-scope domains.",
        ],
    )

    add_heading(doc, "Explicit non-goals for MVP", 2)
    add_bullets(
        doc,
        [
            "Replacing OMS, 626 scanners, Kronos, MAUI, Access Manager, or specialist privacy consoles",
            "Building a second metadata graph or ingestion pipeline",
            "Unified governance command center with live KPI feeds from every downstream system",
            "Feature parity with every Alation module — prioritized by interview-validated daily use",
        ],
    )

    doc.add_paragraph()
    add_para(
        doc,
        "References: Data-Navigator-Approach-One-Pager.md, Data-Navigator-Features-Prioritized-Roadmap.md, "
        "BDC-vs-OMS-Persona-Review-Talking-Sheet.md, PRD-Navigator-Data-Governance-Umbrella.docx, "
        "Navigator-persona-and-626-boundary-brief.md.",
        size=9,
        color=MUTED,
        italic=True,
    )

    doc.save(OUT)
    return OUT


if __name__ == "__main__":
    path = build()
    print(f"Wrote {path}")
