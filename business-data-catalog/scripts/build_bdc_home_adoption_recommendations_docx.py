#!/usr/bin/env python3
"""Build BDC Home adoption & improvement recommendations (.docx)."""
from __future__ import annotations

from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

DIR = Path(__file__).resolve().parents[1]
OUT = DIR / "BDC-Home-Adoption-Improvement-Recommendations.docx"

NAVY = RGBColor(0x1E, 0x2B, 0x52)
MUTED = RGBColor(0x6B, 0x77, 0x8C)


def set_run_font(run, *, bold=False, size=11, color=None, italic=False) -> None:
    run.bold = bold
    run.italic = italic
    run.font.size = Pt(size)
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    if color:
        run.font.color.rgb = color


def add_heading(doc: Document, text: str, level: int = 1) -> None:
    p = doc.add_heading(text, level=level)
    for run in p.runs:
        run.font.name = "Calibri Light"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri Light")
        if level <= 2:
            run.font.color.rgb = NAVY


def add_para(doc: Document, text: str, *, bold: bool = False, italic: bool = False) -> None:
    p = doc.add_paragraph()
    run = p.add_run(text)
    set_run_font(run, bold=bold, italic=italic)


def add_bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        run = p.add_run(item)
        set_run_font(run)


def add_table(doc: Document, headers: list[str], rows: list[list[str]]) -> None:
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    for i, h in enumerate(headers):
        table.rows[0].cells[i].text = h
        for para in table.rows[0].cells[i].paragraphs:
            for run in para.runs:
                set_run_font(run, bold=True, size=10, color=NAVY)
    for r_idx, row in enumerate(rows):
        for c_idx, val in enumerate(row):
            table.rows[r_idx + 1].cells[c_idx].text = val
            for para in table.rows[r_idx + 1].cells[c_idx].paragraphs:
                for run in para.runs:
                    set_run_font(run, size=10)
    doc.add_paragraph()


def build() -> Path:
    doc = Document()
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    label = doc.add_paragraph()
    label.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    set_run_font(label.add_run("Program brief"), bold=True, size=12, color=MUTED)

    title = doc.add_paragraph()
    title.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    set_run_font(
        title.add_run("Navigator BDC Home — Adoption & Improvement Recommendations"),
        bold=True,
        size=22,
        color=NAVY,
    )

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    set_run_font(
        subtitle.add_run("Making the catalog something people open on purpose"),
        italic=True,
        size=12,
        color=MUTED,
    )

    meta = doc.add_paragraph()
    meta.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    set_run_font(
        meta.add_run(f"Prepared {date.today():%B %d, %Y} · Data Governance / BDC (Alation → OMS)"),
        size=10,
        color=MUTED,
    )
    doc.add_paragraph()

    add_heading(doc, "Executive summary", level=1)
    add_para(
        doc,
        "The Home and My Domains direction aligns with how enterprise catalogs succeed: "
        "personalization, trust signals, and a single place to discover governed data. "
        "Adoption will not come from layout alone. Users return when the product closes a loop:",
    )
    add_bullets(
        doc,
        [
            "Pick domains → see a different Home → open an asset → trust it and access it → "
            "come back when something changes.",
        ],
    )
    add_para(doc, "Priority investments:", bold=True)
    add_bullets(
        doc,
        [
            "Real personalization and change signals (not demo telemetry)",
            "Access and semantic actions from the asset page",
            "Team and steward workflows tied to the same metadata",
        ],
    )

    add_heading(doc, "1. Close the gap: demo Home vs. “why I’d return”", level=1)
    add_para(
        doc,
        "Today, Home mixes useful patterns (recent visits, favorites, domains, suggestions) "
        "with static or demo content (generic trending, carousel stats). Users quickly learn "
        "that little is live for them personally.",
    )
    add_para(doc, "Recommended improvements", bold=True)
    add_bullets(
        doc,
        [
            "Wire Home to real signals (start small): recent opens from analytics or persisted "
            "profile history; steward or metadata updates on favorited assets; domain-scoped "
            "“new or changed this week” from OMS crawl metadata.",
            "Replace or defer generic “trending” until you have trustworthy telemetry. Prefer "
            "domain- or role-scoped activity (e.g., “Popular in Subscriber 360 this week”).",
            "One clear primary action above the fold (in addition to search): e.g., “Continue "
            "where you left off” or “Your team’s shared favorites”—not six equal-weight panels.",
        ],
    )
    add_para(
        doc,
        "Design rule: Every Home block should answer “So what do I do next?” with a click that works.",
        italic=True,
    )

    add_heading(doc, "2. Make My Domains earn its place", level=1)
    add_para(
        doc,
        "360 domain selection is the right model for DEEPT, but choosing domains only drives "
        "adoption if the rest of the product visibly reacts.",
    )
    add_para(doc, "Recommended improvements", bold=True)
    add_bullets(
        doc,
        [
            "Immediate payoff after Confirm: confirm message + Home reorder (suggestions, highlights) "
            "scoped to selected domains.",
            "Domain as a lens everywhere—not only Home chips: catalog default filters, glossary/OSSIE "
            "scope, marketplace, breadcrumbs.",
            "Steward-attributed nudges (phased): e.g., tables in your domains lost Gold endorsement; "
            "new data product in a followed 360.",
            "Optional squad defaults: suggest domains from team/org (Workforce) so empty state is not "
            "“pick from 40 chips” cold start.",
        ],
    )
    add_para(doc, "Without visible personalization, domains feel like survey fatigue.", italic=True)
    add_para(doc, "Copy direction (design mock):", bold=True)
    add_bullets(
        doc,
        [
            "Before selection: “Choose your domains to personalize Home—suggestions combine your "
            "picks with your recent catalog activity.”",
            "After selection: short lede “Used to personalize suggestions and highlights on Home.” "
            "and a quiet Change control on the chip row.",
        ],
    )

    add_heading(doc, "3. Search & discovery — one front door", level=1)
    add_para(
        doc,
        "Hero search as primary entry (facet search removed) is correct. Next: trust in results "
        "and paths for users who do not know table names.",
    )
    add_para(doc, "Recommended improvements", bold=True)
    add_bullets(
        doc,
        [
            "Unified ranking across tables, terms, data products, and 360 domains—with type icons "
            "and short “why ranked” explanations (same language as Suggested for you).",
            "Natural language plus filters after search (domain, platform, endorsement level)—not a "
            "second search box on every page.",
            "Guided browse: platform → 360 domain → schema/table (matches Data Sources narrative).",
            "Preserve Alation strengths in migration: steward-curated articles, popular queries, Q&A—"
            "linked from asset pages, not a separate wiki hunt.",
        ],
    )
    add_para(
        doc,
        "If search fails once for a busy analyst, they return to Slack or tribal knowledge.",
        italic=True,
    )

    add_heading(doc, "4. Trust, access, and “can I use this?”", level=1)
    add_para(doc, "Business users do not adopt catalogs that stop at description.")
    add_para(doc, "Recommended improvements", bold=True)
    add_bullets(
        doc,
        [
            "Endorsement, classification, and steward visible on list and detail; one-click request "
            "access or open in Snowflake/Looker where entitlement exists.",
            "Sample data, metric definition, and lineage snippet on the business asset view without "
            "requiring OMS engineering UI for every question.",
            "Glossary/OSSIE and metric conflicts: simplified “what does this number mean?” for consumers; "
            "fuller steward register for governance personas (Me vs Team vs org scope on Home).",
        ],
    )

    add_heading(doc, "5. Team & stewardship loops", level=1)
    add_para(
        doc,
        "Team tab and shared favorites are structurally right but need real workflow hooks.",
    )
    add_para(doc, "Recommended improvements", bold=True)
    add_bullets(
        doc,
        [
            "Shared favorites and comments with notifications (Teams/email) when stewards update "
            "metadata on shared items.",
            "Squad domains and shared data products from marketplace.",
            "Light tasks from DGP work queue surfaced on relevant BDC assets (“Review description for X”)—"
            "one program, two portals.",
        ],
    )

    add_heading(doc, "6. Program & change management", level=1)
    add_para(doc, "Buildable is not the same as adoptable.")
    add_para(doc, "Recommended improvements", bold=True)
    add_bullets(
        doc,
        [
            "Executive metric: monthly active catalog users with meaningful actions (open asset, save "
            "favorite, approved query path)—not login counts alone.",
            "Domain stewards as product owners for their 360 strip on Home (featured products, "
            "“start here” assets, program news).",
            "Alation exit with persona parity checklist (analyst, engineer, steward).",
            "Training carousel evolved into short in-product tasks with completion tracking.",
        ],
    )

    add_heading(doc, "7. Prioritized roadmap (suggested)", level=1)
    add_table(
        doc,
        ["Phase", "Focus", "User outcome"],
        [
            ["Now", "Real recent visits + favorites + domain-filtered catalog/glossary", "“It remembers me.”"],
            ["Next", "Change alerts on favorites; domain-scoped “updated this week”", "“It tells me when to look.”"],
            ["Then", "Access + semantic open (Snowflake/Looker) from asset", "“I can act, not just read.”"],
            [
                "Parallel",
                "Steward Home (conflicts, completeness, work queue links)",
                "“Governance is visible, not a separate tool.”",
            ],
        ],
    )

    add_heading(doc, "Bottom line", level=1)
    add_para(
        doc,
        "Invest first in closed loops: personalization and change signals, then access/semantic actions, "
        "then team/steward collaboration. Treat carousel and generic trending as optional until backed "
        "by real telemetry.",
    )
    add_para(
        doc,
        "Optional next step: Convert this into epics and acceptance criteria for OMS/BDC MVP vs "
        "post-migration—scoped to DEEPT pilot or enterprise cutover.",
        italic=True,
    )

    doc.save(OUT)
    return OUT


if __name__ == "__main__":
    path = build()
    print(path)
