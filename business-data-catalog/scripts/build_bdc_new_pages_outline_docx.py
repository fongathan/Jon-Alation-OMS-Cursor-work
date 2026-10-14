#!/usr/bin/env python3
"""Build BDC new page outlines (.docx) — Data Sources, Home, Admin, Marketplace."""
from __future__ import annotations

from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

DIR = Path(__file__).resolve().parents[1]
OUT = DIR / "BDC-New-Pages-Outline.docx"

NAVY = RGBColor(0x1E, 0x2B, 0x52)
INDIGO = RGBColor(0x2C, 0x47, 0x8C)
MUTED = RGBColor(0x6B, 0x77, 0x8C)


def set_run_font(run, *, bold=False, size=11, color=None, italic=False) -> None:
    run.bold = bold
    run.italic = italic
    run.font.size = Pt(size)
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    if color:
        run.font.color.rgb = color


def add_heading(doc: Document, text: str, level: int = 1) -> None:
    p = doc.add_heading(text, level=level)
    for run in p.runs:
        run.font.name = "Calibri Light"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri Light")
        if level <= 2:
            run.font.color.rgb = NAVY


def add_para(doc: Document, text: str, *, bold: bool = False, italic: bool = False) -> None:
    p = doc.add_paragraph()
    run = p.add_run(text)
    set_run_font(run, bold=bold, italic=italic)


def add_bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        run = p.add_run(item)
        set_run_font(run)


def add_labeled_block(doc: Document, label: str, body: str | list[str]) -> None:
    p = doc.add_paragraph()
    lr = p.add_run(f"{label} ")
    set_run_font(lr, bold=True, color=NAVY)
    if isinstance(body, str):
        add_para(doc, body)
    else:
        add_bullets(doc, body)


def add_table(doc: Document, headers: list[str], rows: list[list[str]]) -> None:
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    for i, h in enumerate(headers):
        table.rows[0].cells[i].text = h
        for para in table.rows[0].cells[i].paragraphs:
            for run in para.runs:
                set_run_font(run, bold=True, size=10, color=NAVY)
    for r_idx, row in enumerate(rows):
        for c_idx, val in enumerate(row):
            table.rows[r_idx + 1].cells[c_idx].text = val
            for para in table.rows[r_idx + 1].cells[c_idx].paragraphs:
                for run in para.runs:
                    set_run_font(run, size=10)
    doc.add_paragraph()


def build() -> Path:
    doc = Document()
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    label = doc.add_paragraph()
    label.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    set_run_font(label.add_run("Requirements outline"), bold=True, size=12, color=MUTED)

    title = doc.add_paragraph()
    title.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    set_run_font(title.add_run("Business Data Catalog — New Pages"), bold=True, size=22, color=NAVY)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    set_run_font(
        subtitle.add_run("Alation migration · OMS backbone · Data Navigator experience"),
        size=13,
        color=INDIGO,
    )

    meta = doc.add_paragraph()
    meta.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    set_run_font(meta.add_run(f"Draft · {date.today().isoformat()}"), size=10, color=MUTED)

    doc.add_paragraph()
    add_para(
        doc,
        "Short page outlines for four BDC surfaces: dual navigation by data source and data domain "
        "(Snowflake and Databricks primary), catalog home and pulse, admin and federated teams, "
        "and a data product marketplace. Aligned with OMS as system of record and a business-first shell.",
    )

    # --- Page 1 ---
    add_heading(doc, "1) Data Source Navigation", level=1)
    add_labeled_block(
        doc,
        "Purpose:",
        "Let users start discovery by where data lives (platform) or what business area it serves "
        "(data domain)—without needing exact table names.",
    )
    add_labeled_block(
        doc,
        "Primary users:",
        "Analysts, business consumers, engineers (hand off to OMS for deep technical work).",
    )
    add_heading(doc, "Layout and behavior", level=2)
    add_bullets(
        doc,
        [
            "Dual entry toggle or tabs: “By data source” | “By data domain” (same catalog grid and asset detail underneath; only the browse tree changes).",
            "By data source (P0): Top-level cards or tree for Snowflake and Databricks (plus “All sources” in Catalog). Drill: source → account/workspace → database/catalog → schema → table/column. Show sync health, last indexed time, and Open in OMS on any asset.",
            "By data domain (P0): Browse using governance domain / proposed 360 buckets (concept map), not legacy folder names. Each domain shows KPIs (coverage, steward-attested %, recent changes) and opens the same filtered catalog.",
            "Shared chrome: Global search, facets (trust, classification, steward, tags), saved views, deep-linkable URL state (e.g. source=snowflake and domain=marketing-audience).",
        ],
    )
    add_labeled_block(
        doc,
        "Alation mapping:",
        "Hierarchical browse (F02), tags/domains (F06), datasource mental model from Alation with Snowflake-first reality.",
    )
    add_labeled_block(
        doc,
        "Acceptance (short):",
        "User can reach a table from either Snowflake path or domain path in three clicks or fewer; "
        "filters persist when switching browse mode; asset IDs match OMS.",
    )

    # --- Page 2 ---
    add_heading(doc, "2) Home (Overview / Catalog pulse)", level=1)
    add_labeled_block(
        doc,
        "Purpose:",
        "Answer “What’s happening in the catalog?” and “How do I use BDC?” in one landing—adoption and trust, not only a link farm.",
    )
    add_labeled_block(
        doc,
        "Primary users:",
        "All personas; role-aware defaults (consumer vs steward vs compliance).",
    )
    add_heading(doc, "Sections (top to bottom)", level=2)
    add_bullets(
        doc,
        [
            "Hero + search — Unified search (tables, glossary, articles, data products); optional persona lens (consumer / steward / compliance read-only).",
            "Trending and popular — Datasets/assets with rising views, saves, or query references (where telemetry exists); Gold / endorsed highlights; do not expose raw PII in rankings.",
            "Recently updated — Business metadata changes (descriptions, owners, classifications, trust flags), glossary links, new/changed data products; filter by domain or source.",
            "News and program updates — Short announcements (migration, new domains onboarded, policy changes); link to articles/docs.",
            "How to use BDC — Curated cards: find data, read trust signals, use glossary, request access / ask a steward, Open in OMS, steward quick actions; link to key articles (Alation article migration target).",
            "Optional KPI strip (steward/leadership) — Domain coverage, active stewardship %, assets missing owners—lightweight, not a full admin console.",
        ],
    )
    add_labeled_block(
        doc,
        "Alation mapping:",
        "Articles (F09), usage signals (F15, phased), trust/endorse (F07), unified search entry (F01).",
    )
    add_labeled_block(
        doc,
        "Acceptance (short):",
        "Home loads with at least one personalized block (recent or domain); every module deep-links to catalog/glossary/article; "
        "“How to” content is editable by admins without deploy.",
    )

    # --- Page 3 ---
    add_heading(doc, "3) Admin / Settings", level=1)
    add_labeled_block(
        doc,
        "Purpose:",
        "Operate the catalog experience—privileges, groups, basic settings, and federated teams (e.g. Ads, Atlas) "
        "in one instance without forcing one tag vocabulary on everyone.",
    )
    add_labeled_block(
        doc,
        "Primary users:",
        "Platform admins, catalog admins, domain stewards (scoped admin), compliance (read/audit).",
    )
    add_heading(doc, "Core areas", level=2)
    add_bullets(
        doc,
        [
            "Identity and access — SSO groups mapped to BDC roles (viewer, contributor, steward, catalog admin); object-level rules for sensitive metadata ops; audit log for admin changes.",
            "Groups and stewardship — Domain/source steward groups; default owners; bulk assignment policies; delegation (who can endorse, deprecate, edit classifications).",
            "Connectors and scope — Snowflake/Databricks connection status, crawl scope, sync schedules; per-source enable/disable for BDC surfacing (626/OMS owns ingestion; BDC configures what appears in business UX).",
            "Taxonomy and custom fields — Global vs namespace fields: enterprise standards (classification, domain) vs team extensions (Ads-specific tags, Atlas conventions). Import/mapping jobs from legacy Alation with preview and conflict resolution.",
            "Multi-team / workspace model (Ads, Atlas, etc.) — Catalog partition metadata (e.g. program, origin_catalog, stewardship_namespace) on assets—not separate databases. Imported tag dictionaries per joining team. “Show only my program” toggle or saved view filtering assets, glossary, and marketplace by namespace; default off for enterprise search. Governance guardrails: enterprise-required fields still apply; cross-namespace assets show tag provenance.",
            "Content and communications — Manage home news, how-to links, featured collections.",
            "Analytics (admin slice) — Adoption, search gaps, import job status.",
        ],
    )
    add_labeled_block(
        doc,
        "Alation mapping:",
        "SSO (F10), tags/custom fields (F06), bulk/dictionary import (F14), admin telemetry (F20, phased).",
    )
    add_labeled_block(
        doc,
        "Acceptance (short):",
        "Admin can assign role by group; steward sees only their domain scope; Ads/Atlas import preserves tags with namespace; "
        "one-click “My program” filter applies across browse, search, and home modules.",
    )

    # --- Page 4 ---
    add_heading(doc, "4) Data Product Marketplace", level=1)
    add_labeled_block(
        doc,
        "Purpose:",
        "Turn curated bundles of tables, metrics, docs, and policies into discoverable, versioned products "
        "others can trust and consume—beyond browsing raw schemas.",
    )
    add_labeled_block(
        doc,
        "Primary users:",
        "Product owners / domain stewards (publish); analysts and consumers (discover/subscribe); engineers (implementation detail via OMS).",
    )
    add_heading(doc, "Lifecycle", level=2)
    add_bullets(
        doc,
        [
            "Draft (prep) — Wizard: name, summary, domain, owners, SLA/refresh expectations, included assets, classifications, trust tier, sample use cases, limitations.",
            "Review — Optional workflow: steward + compliance check for PII/sensitivity; versioning; changelog.",
            "Publish — Listing in Marketplace browse + search; badge on member assets; home can feature new products.",
            "Consume — Subscribe/favorite; access path (request access, linked entitlement, or read-only discovery); links to notebooks/BI/AI-readiness attributes (phased).",
            "Operate — Deprecate/sunset with notice; metrics on subscribers and usage; updates with audit trail.",
        ],
    )
    add_heading(doc, "Browse UX", level=2)
    add_bullets(
        doc,
        [
            "Filters: domain, source (Snowflake/Databricks), trust tier, team namespace, AI-ready.",
            "Product detail page: narrative description, asset graph, owners, policies, changelog, related articles—metadata-driven packaging, not folder placement alone.",
        ],
    )
    add_labeled_block(
        doc,
        "Relation to 360 restructure:",
        "Physical domain ownership stays; Marketplace is the conceptual packaging layer (nine-bucket taxonomy can label products without collapsing Snowflake governance boundaries).",
    )
    add_labeled_block(
        doc,
        "Dependencies:",
        "Stable asset IDs in OMS, ownership/stewardship (F08), trust flags (F07), optional requests workflow (F12).",
    )
    add_labeled_block(
        doc,
        "Acceptance (short):",
        "Steward publishes a v1 product with at least one Snowflake/Databricks asset and glossary link; consumer finds it via Marketplace and search; "
        "deprecation propagates to listings and asset badges; publish requires role and audit entry.",
    )

    # --- Nav table ---
    add_heading(doc, "Suggested nav placement (F30)", level=1)
    add_table(
        doc,
        ["Page", "Typical nav label"],
        [
            ["Data Source Navigation", "Data sources (default browse) + domain switcher"],
            ["Home", "Overview or Home"],
            ["Admin / Settings", "Admin (after divider with Requests/Workflows)"],
            ["Data Product Marketplace", "Marketplace or Data products"],
        ],
    )

    add_para(doc, "Traceability: Data Navigator roadmap, 360 Data Product restructure summary, F30 OMS/BDC navigation IA.", italic=True)

    doc.save(OUT)
    return OUT


if __name__ == "__main__":
    path = build()
    print(path)
