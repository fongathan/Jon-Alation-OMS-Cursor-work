#!/usr/bin/env python3
"""Build simple Word write-up: BDC serving Data, Ads, and Atlas orgs."""
from __future__ import annotations

from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

DIR = Path(__file__).resolve().parents[1]
OUT = DIR / "BDC-Multi-Org-Data-Ads-Atlas-Writeup.docx"

NAVY = RGBColor(0x1E, 0x2B, 0x52)
MUTED = RGBColor(0x6B, 0x77, 0x8C)


def set_run_font(run, *, bold=False, size=11, italic=False, color=None) -> None:
    run.bold = bold
    run.italic = italic
    run.font.size = Pt(size)
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    if color:
        run.font.color.rgb = color


def add_heading(doc: Document, text: str, level: int = 1) -> None:
    p = doc.add_heading(text, level=level)
    for run in p.runs:
        run.font.name = "Calibri Light"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri Light")
        if level <= 2:
            run.font.color.rgb = NAVY


def add_para(doc: Document, text: str, *, bold=False, italic=False) -> None:
    p = doc.add_paragraph()
    run = p.add_run(text)
    set_run_font(run, bold=bold, italic=italic)


def add_bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        run = p.add_run(item)
        set_run_font(run)


def add_table(doc: Document, headers: list[str], rows: list[tuple[str, ...]]) -> None:
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = h
        for p in cell.paragraphs:
            for run in p.runs:
                set_run_font(run, bold=True, size=10)
    for r_idx, row in enumerate(rows):
        for c_idx, val in enumerate(row):
            cell = table.rows[r_idx + 1].cells[c_idx]
            cell.text = val
            for p in cell.paragraphs:
                for run in p.runs:
                    set_run_font(run, size=10)
    doc.add_paragraph()


def build() -> Path:
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    today = date.today().strftime("%B %d, %Y")

    title = doc.add_paragraph()
    title.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    tr = title.add_run("Business Data Catalog on OMS")
    set_run_font(tr, bold=True, size=18, color=NAVY)

    sub = doc.add_paragraph()
    sub.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    sr = sub.add_run("Serving Data, Ads, and Atlas — simple plan")
    set_run_font(sr, size=13, color=MUTED)

    meta = doc.add_paragraph()
    meta.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    mr = meta.add_run(f"Draft · {today}")
    set_run_font(mr, italic=True, size=10, color=MUTED)

    doc.add_paragraph()

    add_heading(doc, "Purpose of this note", 2)
    add_para(
        doc,
        "We are moving from Alation to a Business Data Catalog (BDC) on the Operational Metadata Store (OMS). "
        "Data (DEEPT), Ads, and Atlas may all use the same Navigator experience—but we organize data differently today. "
        "This note explains how we work in Alation now, what we want on OMS, and a simple way to let each org join "
        "without forcing everyone into one domain model.",
    )

    add_heading(doc, "One shared catalog, three orgs", 2)
    add_bullets(
        doc,
        [
            "We should not build three separate catalogs or three copies of Navigator.",
            "We should build one catalog and one UI shell, with org-specific settings for how domains, filters, and home pages behave.",
            "Enterprise governance (classifications, policies, DGP) stays shared where it must; each org owns its business metadata and stewards.",
        ],
    )

    add_heading(doc, "How we work today in Alation (shared instance)", 2)
    add_para(
        doc,
        "Data, Ads, and Atlas all use the same Alation instance. In Alation we are set up as separate "
        "“Domains”—for example Ad Platforms, Atlas Domain, and DEET Data Domain. Those entries behave like org "
        "boundaries even though Alation calls them domains.",
    )
    add_para(doc, "In practice the three orgs differ in how assets are grouped and found:", bold=True)

    add_table(
        doc,
        ["Org", "How we show up in Alation", "Data sources", "How assets are grouped / found", "Access / stewards"],
        [
            (
                "Data (DEEPT)",
                "DEET Data Domain (and related domain objects)",
                "Primarily Snowflake 360 conform and related enterprise sources",
                "Composable Workspaces model: many named 360 Data Products (e.g. Subscriber 360, Content 360)",
                "Business + technical stewards per 360; familiar to DEEPT consumers",
            ),
            (
                "Ads",
                "Ad Platforms domain",
                "Often different sources than core Data 360 (Ads-specific platforms)",
                "Naming convention: “(Ads)” prefix on tables and similar objects so Ads assets are easy to spot",
                "Ads stewards; consumers learn to search and browse within Ads naming and domain",
            ),
            (
                "Atlas",
                "Atlas Domain (e.g. Atlas / CIM Team)",
                "Largely the same technical sources as Data (e.g. Snowflake)",
                "Assets tagged and associated to the Atlas domain; not mainly a 360-folder model",
                "Datasets locked down to Atlas stewards and domain members (e.g. Atlas CIM Team admins)",
            ),
        ],
    )

    add_para(
        doc,
        "Important: Alation “domain” ≈ org workspace for us. On OMS we should use clear language—"
        "org (Data / Ads / Atlas) and domain taxonomy (360 vs tags vs prefixes)—so we do not confuse "
        "Alation product terms with DEEPT 360 Data Products.",
    )

    add_heading(doc, "What we want on OMS / BDC", 2)
    add_bullets(
        doc,
        [
            "One asset graph in OMS with technical metadata from crawlers (Snowflake, Databricks, etc.).",
            "Every governable asset has a primary owning org (Data, Ads, or Atlas)—not guessed only from tags.",
            "Each org can define how consumers browse “domains” without breaking the others’ model.",
            "Global search in Navigator can find assets across orgs, with org and steward visible on every result.",
            "Org context (switcher or default from workforce / access) sets home page, facet labels, and domain lists.",
        ],
    )

    add_heading(doc, "Recommended approach (keep it simple)", 2)

    add_heading(doc, "1. Org profile (configuration)", 3)
    add_para(
        doc,
        "A small, governed config per org: display name, default landing tab, which modules are on "
        "(Home, Catalog, Glossary, Marketplace, Admin), support channel, and which domain taxonomy type to use.",
    )

    add_heading(doc, "2. Domain taxonomy per org (not one size fits all)", 3)
    add_bullets(
        doc,
        [
            "Data: keep 360 Data Products as the main browse axis (today’s Composable Workspaces story).",
            "Ads: support Ads-specific sources plus browse/filter by org and naming pattern; treat “(Ads)” prefix as a display and search convention until metadata catches up.",
            "Atlas: domain membership via tags / steward scope on shared sources; Atlas home and filters emphasize Atlas-tagged assets, not the full 360 list.",
        ],
    )

    add_heading(doc, "3. Asset affiliation in metadata", 3)
    add_para(
        doc,
        "Replace inference-only rules with explicit fields on each asset (stored in OMS): primary org; "
        "domain or product memberships; optional shared-with orgs for cross-cutting assets (identity, governance). "
        "Migration from Alation should map each Alation Domain object to org + membership on assets.",
    )

    add_heading(doc, "4. Same UI patterns, org-specific content", 3)
    add_bullets(
        doc,
        [
            "One header search for discovery across the enterprise (with org on each hit).",
            "Catalog: compact “Search assets…” in the filter column filters the current list; facets come from the active org profile.",
            "Asset detail: always show owning org, stewards, and platform link—same layout for all orgs.",
        ],
    )

    add_heading(doc, "Phased rollout", 2)
    add_table(
        doc,
        ["Phase", "What we deliver", "Who benefits"],
        [
            (
                "A — Federation",
                "Org filter, separate domain lists, explicit primary org on onboarded assets; migrate Alation domain membership",
                "Ads and Atlas can appear in the same BDC without adopting 360 naming",
            ),
            (
                "B — Org home & facets",
                "Custom home blocks and facet labels per org profile",
                "Each org sees a familiar browse path (360 vs Ads prefix vs Atlas tags)",
            ),
            (
                "C — Shared governance",
                "DGP classifications/policies enterprise-wide; glossary core + org extensions",
                "One compliance story; orgs keep local terms and products",
            ),
        ],
    )

    add_heading(doc, "Questions to align with Ads and Atlas", 2)
    add_bullets(
        doc,
        [
            "Should Alation “Domain” map 1:1 to org in OMS, or do we need sub-domains under Atlas or Ads?",
            "For Ads: is “(Ads)” prefix the long-term convention, or will we normalize titles in OMS and keep prefix only in source systems?",
            "For Atlas: which tags or steward groups define membership for migration?",
            "Who approves new top-level org config changes (platform vs org stewards)?",
            "Do any assets require hard isolation (separate search results or access), or is filter + permissions enough?",
        ],
    )

    add_heading(doc, "Suggested next steps", 2)
    add_bullets(
        doc,
        [
            "Document authoritative steward lists and Alation domain IDs for Ads and Atlas (as in Atlas Domain today).",
            "Agree on org profile template and taxonomy types for a short pilot (Catalog + asset detail only).",
            "Inventory Ads-only sources vs shared Snowflake for crawl and platform links.",
            "Add multi-org requirements to the BDC / Navigator PRD as a dedicated section with acceptance criteria.",
        ],
    )

    add_para(doc, "")
    foot = doc.add_paragraph()
    foot.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    fr = foot.add_run("Prepared for Data Governance / BDC program · Navigator on OMS")
    set_run_font(fr, italic=True, size=9, color=MUTED)

    doc.save(OUT)
    return OUT


if __name__ == "__main__":
    path = build()
    print(path)
