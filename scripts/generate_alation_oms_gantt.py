#!/usr/bin/env python3
"""Generate Excel Gantt workbook: Alation → In-house (OMS) data catalog migration.

Styling matches Sustainable_Inventories_Kronos_Sisu_Executive_Gantt (Jon Kronos Work).
"""

from __future__ import annotations

from datetime import date, timedelta
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

OUT_PATH = (
    Path(__file__).resolve().parents[1]
    / "Alation-to-In-House-Business-Data-Catalog-Gantt-Chart.xlsx"
)

# --- Executive styling (Sustainable_Inventories_Kronos_Sisu_Executive_Gantt) ---
HEADER_FILL = PatternFill("solid", fgColor="0F2744")
HEADER_FONT = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
TITLE_BAND_1 = PatternFill("solid", fgColor="082441")
TITLE_BAND_2 = PatternFill("solid", fgColor="135CA9")
TITLE_BAND_3 = PatternFill("solid", fgColor="B9D4F5")
TITLE_FONT = Font(name="Calibri", size=27, bold=True, color="FFFFFF")
SUBTITLE_ON_BAND = Font(name="Calibri", size=11, color="E8F4FF")
META_ON_BAND_LIGHT = Font(name="Calibri", size=10, italic=True, color="1A3A5C")
BODY_FONT = Font(name="Calibri", size=10, color="1F2937")
META_FONT = Font(name="Calibri", size=9, italic=True, color="64748B")
THIN = Side(style="thin", color="CBD5E1")
GRID_BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

# Timeline bar colors (semantic; distinct for scanning)
BAR_COLORS: dict[str, str] = {
    "kickoff": "4A6FA5",
    "requirements": "2E5A87",
    "design": "8B4A9C",
    "development": "1B3A5C",
    "testing": "2D6A4F",
    "migration": "6B7280",
    "change": "135CA9",
    "launch": "0F2744",
}

# Zebra: subtle blue striping (two cool tints, easy to scan)
ZEBRA: tuple[PatternFill, PatternFill] = (
    PatternFill("solid", fgColor="EFF6FF"),
    PatternFill("solid", fgColor="E0F2FE"),
)


def monday_on_or_before(d: date) -> date:
    return d - timedelta(days=d.weekday())


def main() -> None:
    wb = Workbook()

    tasks: list[tuple[str, str, str, date, date, str]] = [
        ("0", "Intake", "Kickoff & program intake (T+0 milestone 4/30/26)", date(2026, 4, 27), date(2026, 5, 8), "kickoff"),
        ("1", "Phase 1", "Requirements — current use, enhancements, OMS leverage & gaps", date(2026, 7, 1), date(2026, 8, 25), "requirements"),
        ("1a", "Phase 1", "Consolidated requirements doc & gate sign-off", date(2026, 8, 18), date(2026, 8, 28), "requirements"),
        ("2", "Design", "Solution design — backend (data model, integrations, lineage, APIs)", date(2026, 8, 26), date(2026, 10, 16), "design"),
        ("2b", "Design", "Solution design — frontend / BDC UX (search, glossary, asset IA)", date(2026, 9, 14), date(2026, 11, 3), "design"),
        ("2c", "Design", "Migration design (extract–transform–load, validation strategy)", date(2026, 10, 7), date(2026, 11, 3), "design"),
        ("2d", "Design", "Architecture / design review & approval", date(2026, 10, 27), date(2026, 11, 3), "design"),
        ("3", "Build", "Development — MVP parity (high-priority catalog: F01–F06, F09–F11, F16, F20–F21, F23–F25, F30, F33)", date(2026, 11, 4), date(2027, 3, 31), "development"),
        ("3b", "Build", "Development — Phase 2 (medium: bulk edit, dataflows, domains, exports, usage, API, policy, etc.)", date(2027, 1, 4), date(2027, 6, 15), "development"),
        ("3c", "Build", "Development — migration tooling, parallel-run hooks, hardening", date(2027, 4, 1), date(2027, 6, 30), "development"),
        ("4", "Quality", "Test planning, automation & regression prep", date(2027, 6, 1), date(2027, 6, 30), "testing"),
        ("4b", "Quality", "Formal testing window (4 weeks, July 2027)", date(2027, 7, 1), date(2027, 7, 28), "testing"),
        ("5", "Cutover", "Migration execution, UAT, stabilization (6 weeks)", date(2027, 7, 29), date(2027, 9, 11), "migration"),
        ("6", "Adoption", "Change management — training, champions, comms calendar", date(2027, 6, 1), date(2027, 9, 15), "change"),
        ("6b", "Cutover", "Go-live readiness, cutover rehearsal, launch", date(2027, 9, 4), date(2027, 9, 16), "launch"),
        ("7", "Hypercare", "Post-launch stabilization (optional window)", date(2027, 9, 16), date(2027, 10, 16), "migration"),
    ]

    grid_start = monday_on_or_before(tasks[0][3])
    grid_end = max(t[4] for t in tasks)
    grid_end_monday = monday_on_or_before(grid_end) + timedelta(weeks=2)
    week_starts: list[date] = []
    cur = grid_start
    while cur <= grid_end_monday:
        week_starts.append(cur)
        cur += timedelta(weeks=1)

    labels = ["ID", "Phase", "Workstream / deliverable", "Start", "End"]
    first_week_col = len(labels) + 1
    ncols = len(labels) + len(week_starts)
    last_letter = get_column_letter(ncols)

    ws = wb.active
    ws.title = "Executive Gantt"

    # --- Title block (full width): same 3-row band as Kronos / Sisu Gantt ---
    ws.merge_cells(f"A1:{last_letter}1")
    c1 = ws["A1"]
    c1.value = "Alation to In-House: Business Data Catalog - Gantt Chart"
    c1.font = TITLE_FONT
    c1.fill = TITLE_BAND_1
    c1.alignment = Alignment(horizontal="left", vertical="center", indent=1, wrap_text=True)
    ws.row_dimensions[1].height = 48

    ws.merge_cells(f"A2:{last_letter}2")
    c2 = ws["A2"]
    c2.value = (
        "Data governance catalog migration — requirements, design, build, testing, cutover "
        "(program timeline; launch target Sep 16, 2027)"
    )
    c2.font = SUBTITLE_ON_BAND
    c2.fill = TITLE_BAND_2
    c2.alignment = Alignment(horizontal="left", vertical="center", indent=1, wrap_text=True)
    ws.row_dimensions[2].height = 22

    ws.merge_cells(f"A3:{last_letter}3")
    c3 = ws["A3"]
    c3.value = (
        "Sources: Product Brief — Alation to In-house (Mar 2026) · BDC Feature Inventory Matrix (PRD F01–F35). "
        "Mitigate contract risk via extension / parallel run per brief."
    )
    c3.font = META_ON_BAND_LIGHT
    c3.fill = TITLE_BAND_3
    c3.alignment = Alignment(horizontal="left", vertical="center", indent=1, wrap_text=True)
    ws.row_dimensions[3].height = 20

    header_row = 5
    for col, label in enumerate(labels, start=1):
        cell = ws.cell(row=header_row, column=col, value=label)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = GRID_BORDER

    for i, wk in enumerate(week_starts):
        col = first_week_col + i
        cell = ws.cell(row=header_row, column=col, value=wk)
        cell.number_format = "mmm d"
        cell.font = Font(name="Calibri", size=9, bold=True, color="FFFFFF")
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", text_rotation=40)
        cell.border = GRID_BORDER
    ws.row_dimensions[header_row].height = 52

    data_start = header_row + 1
    for idx, (tid, phase, name, ts, te, bucket) in enumerate(tasks):
        r = data_start + idx
        base_fill = ZEBRA[idx % 2]

        ws.cell(row=r, column=1, value=tid).font = BODY_FONT
        ws.cell(row=r, column=1).alignment = Alignment(vertical="top", wrap_text=True)
        ws.cell(row=r, column=1).fill = base_fill
        ws.cell(row=r, column=1).border = GRID_BORDER

        ws.cell(row=r, column=2, value=phase).font = BODY_FONT
        ws.cell(row=r, column=2).alignment = Alignment(vertical="top", wrap_text=True)
        ws.cell(row=r, column=2).fill = base_fill
        ws.cell(row=r, column=2).border = GRID_BORDER

        ws.cell(row=r, column=3, value=name).font = BODY_FONT
        ws.cell(row=r, column=3).alignment = Alignment(vertical="top", wrap_text=True)
        ws.cell(row=r, column=3).fill = base_fill
        ws.cell(row=r, column=3).border = GRID_BORDER

        for col_idx, d in enumerate([ts, te], start=4):
            cell = ws.cell(row=r, column=col_idx, value=d)
            cell.number_format = "mmm d, yyyy"
            cell.font = BODY_FONT
            cell.alignment = Alignment(horizontal="center", vertical="top")
            cell.fill = base_fill
            cell.border = GRID_BORDER

        bar_fill = PatternFill("solid", fgColor=BAR_COLORS.get(bucket, "1B3A5C"))
        for wi, wk in enumerate(week_starts):
            col = first_week_col + wi
            week_end = wk + timedelta(days=6)
            active = ts <= week_end and te >= wk
            cell = ws.cell(row=r, column=col, value="")
            cell.border = GRID_BORDER
            if active:
                cell.fill = bar_fill
                cell.alignment = Alignment(horizontal="center", vertical="center")
            else:
                cell.fill = base_fill

        ws.row_dimensions[r].height = 36

    # Column widths (reference uses very wide week cols; we scale down for ~80 weeks)
    ws.column_dimensions["A"].width = 10
    ws.column_dimensions["B"].width = 18
    ws.column_dimensions["C"].width = 72
    ws.column_dimensions["D"].width = 16
    ws.column_dimensions["E"].width = 16
    week_w = max(3.2, min(5.5, 280 / max(len(week_starts), 1)))
    for i in range(len(week_starts)):
        ws.column_dimensions[get_column_letter(first_week_col + i)].width = week_w

    ws.freeze_panes = ws.cell(row=data_start, column=first_week_col)

    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToHeight = 1
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToPage = True
    ws.print_options.horizontalCentered = True
    ws.sheet_properties.tabColor = "0F2744"

    # Legend + color key (below chart, Kronos pattern)
    legend_row = header_row + len(tasks) + 2
    ws.merge_cells(start_row=legend_row, start_column=1, end_row=legend_row, end_column=first_week_col - 1)
    leg = ws.cell(row=legend_row, column=1)
    leg.value = (
        "Legend: each timeline column is one calendar week (Monday–Sunday). "
        "Bar shading shows the activity span. Color key identifies workstream type."
    )
    leg.font = META_FONT
    leg.alignment = Alignment(wrap_text=True, vertical="top")

    key_labels = [
        ("Kickoff / intake", "kickoff"),
        ("Requirements", "requirements"),
        ("Design", "design"),
        ("Development / build", "development"),
        ("Testing / QA", "testing"),
        ("Migration / UAT", "migration"),
        ("Change / adoption", "change"),
        ("Launch", "launch"),
    ]
    key_start = legend_row + 2
    ws.cell(row=key_start, column=1, value="Workstream color key").font = Font(
        name="Calibri", size=10, bold=True, color="0F2744"
    )
    r = key_start + 1
    for label, key in key_labels:
        sw = ws.cell(row=r, column=1, value="")
        sw.fill = PatternFill("solid", fgColor=BAR_COLORS[key])
        sw.border = GRID_BORDER
        lbl = ws.cell(row=r, column=2, value=label)
        lbl.font = META_FONT
        lbl.alignment = Alignment(vertical="center")
        ws.row_dimensions[r].height = 18
        r += 1

    # --- Milestones sheet (Kronos-style header + table) ---
    ms = wb.create_sheet("Milestones & dates")
    ms.sheet_properties.tabColor = "2E5A87"
    ms.merge_cells("A1:C1")
    mst = ms["A1"]
    mst.value = "Milestone checkpoints — Business Data Catalog migration"
    mst.font = Font(name="Calibri", size=16, bold=True, color="FFFFFF")
    mst.fill = TITLE_BAND_2
    mst.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ms.row_dimensions[1].height = 26
    ms.merge_cells("A2:C2")
    ms2 = ms["A2"]
    ms2.value = "Alation to In-House · Dates per program plan"
    ms2.font = META_ON_BAND_LIGHT
    ms2.fill = TITLE_BAND_3
    ms2.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ms.row_dimensions[2].height = 20

    hr = 4
    for col, h in enumerate(["Milestone", "Target date", "Notes"], 1):
        cell = ms.cell(row=hr, column=col, value=h)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.border = GRID_BORDER
        cell.alignment = Alignment(horizontal="center", wrap_text=True)

    milestones = [
        ("Kickoff / intake (T+0)", date(2026, 4, 30), "Program start; align DSS / OMS / stakeholders"),
        ("Phase 1 requirements complete", date(2026, 8, 28), "8-week window Jul–Aug 2026 (per plan)"),
        ("Design complete", date(2026, 11, 3), "10-week backend + frontend design Aug–Nov"),
        ("Development window", date(2026, 11, 4), "8 months through Jun 2027 (MVP → parity → tooling)"),
        ("Formal testing", date(2027, 7, 1), "4-week test window (July 2027)"),
        ("Migration + UAT + stabilization", date(2027, 7, 29), "6-week execution ending ahead of launch"),
        ("Launch — Internal Solution primary", date(2027, 9, 16), "Go-live; Alation decommission per contract"),
    ]
    for i, row in enumerate(milestones, start=hr + 1):
        ms.cell(row=i, column=1, value=row[0]).font = BODY_FONT
        ms.cell(row=i, column=2, value=row[1]).number_format = "mmm d, yyyy"
        ms.cell(row=i, column=2).font = BODY_FONT
        ms.cell(row=i, column=3, value=row[2]).font = BODY_FONT
        for c in range(1, 4):
            cell = ms.cell(row=i, column=c)
            cell.border = GRID_BORDER
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        z = ZEBRA[(i - hr) % 2]
        for c in range(1, 4):
            ms.cell(row=i, column=c).fill = z
    ms.column_dimensions["A"].width = 42
    ms.column_dimensions["B"].width = 18
    ms.column_dimensions["C"].width = 72

    # --- Feature matrix sheet ---
    feat = wb.create_sheet("Feature matrix (summary)")
    feat.sheet_properties.tabColor = "1B3A5C"
    feat.merge_cells("A1:F1")
    ft = feat["A1"]
    ft.value = "BDC Feature Inventory — priority snapshot for migration waves"
    ft.font = Font(name="Calibri", size=16, bold=True, color="FFFFFF")
    ft.fill = TITLE_BAND_2
    ft.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    feat.row_dimensions[1].height = 26
    feat.merge_cells("A2:F2")
    ft2 = feat["A2"]
    ft2.value = "See full matrix in Google Doc · High items = MVP wave"
    ft2.font = META_ON_BAND_LIGHT
    ft2.fill = TITLE_BAND_3
    ft2.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    feat.row_dimensions[2].height = 20

    fhr = 4
    hdr = ["#", "Feature", "Priority", "MVP wave", "Primary users", "PRD"]
    for col, h in enumerate(hdr, 1):
        cell = feat.cell(row=fhr, column=col, value=h)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.border = GRID_BORDER
        cell.alignment = Alignment(horizontal="center", wrap_text=True)

    high_rows = [
        ("1", "Search & Discovery", "High", "MVP", "DPMs, PMs, Analysts, Stewards", "F01"),
        ("2", "Business Glossary", "High", "MVP", "DPMs, PMs, Analysts, Stewards", "F02"),
        ("3", "Catalog Browse", "High", "MVP", "DPMs, PMs, Analysts, Stewards", "F03"),
        ("4", "Table / asset detail & descriptions", "High", "MVP", "DPMs, PMs, Analysts, Stewards", "F04"),
        ("5", "Recommendations (AI metadata)", "High", "MVP", "Stewards", "F16"),
        ("6", "Flags (endorse / warn / deprecate)", "High", "MVP", "DPMs, PMs, Analysts, Stewards", "F21"),
        ("7", "Data Lineage", "High", "MVP", "DPMs, Engineers, …", "F05"),
        ("8", "Admin analytics dashboard", "High", "MVP", "Admin", "F25"),
        ("9", "Tags & classification", "High", "MVP", "Governance, DPMs, Stewards", "F06"),
        ("10", "Expanded data sources", "High", "MVP", "Admins", "F30"),
        ("11", "Product pages (authoring)", "High", "MVP", "PMs, DPMs, Stewards", "F33"),
        ("12", "Data sources & connectors", "High", "MVP", "DPMs, Engineers", "F11"),
        ("13", "Authentication / SSO", "High", "MVP", "All users", "F20"),
        ("14", "Folders / source navigation", "High", "MVP", "Cross-functional", "F23"),
        ("15", "Ownership / stewardship", "High", "MVP", "Admin", "F09"),
    ]
    for ri, hrw in enumerate(high_rows, start=fhr + 1):
        for ci, val in enumerate(hrw, start=1):
            cell = feat.cell(row=ri, column=ci, value=val)
            cell.font = BODY_FONT
            cell.border = GRID_BORDER
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        z = ZEBRA[(ri - fhr) % 2]
        for ci in range(1, 7):
            feat.cell(row=ri, column=ci).fill = z

    note_row = fhr + len(high_rows) + 1
    feat.cell(row=note_row, column=1, value="Medium / Phase 2 examples").font = Font(
        name="Calibri", size=10, bold=True, color="0F2744"
    )
    feat.merge_cells(start_row=note_row, start_column=2, end_row=note_row, end_column=6)
    feat.cell(row=note_row, column=2).value = (
        "Bulk edit, dataflows, domains, exports, usage metrics, AI chat, catalog sets, "
        "dictionary import/export, saved queries, overview & social, custom fields, comments, "
        "version control, policy center, …"
    )
    feat.cell(row=note_row, column=2).font = META_FONT
    feat.cell(row=note_row, column=2).alignment = Alignment(wrap_text=True, vertical="top")

    feat.column_dimensions["A"].width = 5
    feat.column_dimensions["B"].width = 34
    feat.column_dimensions["C"].width = 10
    feat.column_dimensions["D"].width = 10
    feat.column_dimensions["E"].width = 28
    feat.column_dimensions["F"].width = 8

    wb.save(OUT_PATH)
    print(f"Wrote {OUT_PATH}")


if __name__ == "__main__":
    main()
