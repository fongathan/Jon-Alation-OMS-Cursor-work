#!/usr/bin/env python3
"""Build a small PPTX with layout options for metadata / OMS solution slides."""

from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

# Theme: dark blue deck, gold accents (aligned with prior slide description)
BG = (18, 32, 72)
PANEL = (32, 48, 98)
ACCENT = (230, 196, 110)
MUTED = (180, 190, 220)
WHITE = (255, 255, 255)


def rgb(c: tuple[int, int, int]) -> RGBColor:
    return RGBColor(c[0], c[1], c[2])


def slide_bg(slide, c: tuple[int, int, int] = BG) -> None:
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = rgb(c)


def tb(
    slide,
    left: float,
    top: float,
    width: float,
    height: float,
    text: str,
    *,
    size: int = 14,
    bold: bool = False,
    color: tuple[int, int, int] = WHITE,
    align: PP_ALIGN = PP_ALIGN.LEFT,
) -> None:
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(size)
    p.font.bold = bold
    p.font.color.rgb = rgb(color)
    p.alignment = align


def round_box(
    slide,
    left: float,
    top: float,
    width: float,
    height: float,
    text: str,
    *,
    fill: tuple[int, int, int] = PANEL,
    font_pt: int = 11,
) -> None:
    sh = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(left),
        Inches(top),
        Inches(width),
        Inches(height),
    )
    sh.fill.solid()
    sh.fill.fore_color.rgb = rgb(fill)
    sh.line.color.rgb = rgb((90, 110, 170))
    tf = sh.text_frame
    tf.clear()
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(font_pt)
    p.font.bold = True
    p.font.color.rgb = rgb(WHITE)
    p.alignment = PP_ALIGN.CENTER


def arrow_label(slide, left: float, top: float, text: str) -> None:
    tb(slide, left, top, 1.2, 0.35, text, size=16, bold=True, color=ACCENT, align=PP_ALIGN.CENTER)


def footer(slide, text: str) -> None:
    tb(slide, 0.4, 6.85, 12.5, 0.35, text, size=9, bold=False, color=MUTED)


def build() -> Path:
    prs = Presentation()
    prs.slide_width = Inches(13.333)  # 16:9
    prs.slide_height = Inches(7.5)

    blank = prs.slide_layouts[6]

    # --- Slide 1: cover / legend ---
    s = prs.slides.add_slide(blank)
    slide_bg(s)
    tb(s, 0.55, 0.45, 12.2, 0.9, "Metadata solution story — layout options", size=32, bold=True)
    tb(
        s,
        0.55,
        1.45,
        12.0,
        2.2,
        "Each following slide is one visual approach we discussed:\n"
        "• Option A — staged pipeline (sync → store → consume)\n"
        "• Option B — same pipeline + Policy and AI as explicit inputs\n"
        "• Option C — map the four capability bullets under the stages\n"
        "• Option D — MVP vs later on one view\n\n"
        "Edit copy freely; boxes are native PowerPoint shapes.",
        size=16,
    )
    footer(s, "Internal draft — layout options for review")

    # --- Slide 2: Option A — layered pipeline ---
    s = prs.slides.add_slide(blank)
    slide_bg(s)
    tb(s, 0.55, 0.35, 12.0, 0.65, "Option A — Staged pipeline (diagram matches depth of story)", size=22, bold=True)
    tb(
        s,
        0.55,
        1.05,
        12.0,
        0.55,
        "Sources → ingestion & sync → catalog in OMS → discovery in apps / search",
        size=13,
        color=MUTED,
    )
    y = 1.85
    w = 2.35
    h = 1.15
    gap = 0.35
    x0 = 0.55
    round_box(s, x0, y, w, h, "Sources\n(Snowflake, Looker, …)", font_pt=12)
    arrow_label(s, x0 + w + 0.05, y + h / 2 - 0.15, "→")
    x1 = x0 + w + gap + 0.35
    round_box(s, x1, y, w + 0.35, h, "Ingestion & sync\n(connectors, scans,\nDDL / API harvest)", font_pt=11)
    arrow_label(s, x1 + w + 0.35 + 0.05, y + h / 2 - 0.15, "→")
    x2 = x1 + w + 0.35 + gap + 0.35
    round_box(s, x2, y, w + 0.45, h, "OMS / metadata store\n(APIs, lineage links,\nsteward workflows)", font_pt=11)
    arrow_label(s, x2 + w + 0.45 + 0.05, y + h / 2 - 0.15, "→")
    x3 = x2 + w + 0.45 + gap + 0.35
    round_box(s, x3, y, w + 0.2, h, "Consumption\n(search, BI context,\ndata products)", font_pt=11)
    tb(
        s,
        0.55,
        3.35,
        12.0,
        1.1,
        "Why: executives see where truth is captured vs where it is experienced — not a single vague “Metadata” bucket.",
        size=14,
        color=MUTED,
    )
    footer(s, "Option A — staged pipeline")

    # --- Slide 3: Option B — Policy + AI rails ---
    s = prs.slides.add_slide(blank)
    slide_bg(s)
    tb(s, 0.55, 0.35, 12.0, 0.65, "Option B — Policy + AI as first-class inputs (not only in bullets)", size=22, bold=True)
    tb(s, 0.55, 1.05, 12.0, 0.5, "Same flow as A, with governance and automation visible", size=13, color=MUTED)
    # Left rail: Policy
    round_box(s, 0.45, 1.75, 1.55, 2.4, "Policy\n• classification rules\n• PII / retention\n• approval gates", font_pt=10)
    # Right rail: AI
    round_box(s, 11.25, 1.75, 1.55, 2.4, "AI assists\n• suggest tags\n• reconcile metric defs\n• rank / personalize search", font_pt=10)
    y = 2.05
    w = 2.05
    h = 1.05
    x0 = 2.35
    round_box(s, x0, y, w, h, "Sources", font_pt=12)
    arrow_label(s, x0 + w + 0.02, y + 0.32, "→")
    x1 = x0 + w + 0.42
    round_box(s, x1, y, w + 0.25, h, "Ingestion\n& sync", font_pt=11)
    arrow_label(s, x1 + w + 0.25 + 0.02, y + 0.32, "→")
    x2 = x1 + w + 0.25 + 0.42
    round_box(s, x2, y, w + 0.35, h, "OMS /\nmetadata", font_pt=11)
    arrow_label(s, x2 + w + 0.35 + 0.02, y + 0.32, "→")
    x3 = x2 + w + 0.35 + 0.42
    round_box(s, x3, y, w + 0.15, h, "Search &\napps", font_pt=11)
    # Curved hint lines: use thin text callouts (connectors are brittle in script)
    tb(s, 2.0, 1.15, 1.8, 0.45, "governs", size=10, bold=True, color=ACCENT, align=PP_ALIGN.CENTER)
    tb(s, 2.0, 3.25, 1.8, 0.45, "audits", size=10, bold=True, color=ACCENT, align=PP_ALIGN.CENTER)
    tb(s, 11.15, 1.15, 1.8, 0.45, "suggests", size=10, bold=True, color=ACCENT, align=PP_ALIGN.CENTER)
    tb(s, 11.15, 3.25, 1.8, 0.45, "reconciles", size=10, bold=True, color=ACCENT, align=PP_ALIGN.CENTER)
    tb(
        s,
        2.35,
        3.45,
        8.6,
        1.0,
        "Tip in deck build: draw short elbow connectors from the side boxes to Ingestion, OMS, and Search "
        "so “governs / suggests” land on specific steps.",
        size=12,
        color=MUTED,
    )
    footer(s, "Option B — policy + AI rails (add connectors by hand for polish)")

    # --- Slide 4: Option C — bullets under stages ---
    s = prs.slides.add_slide(blank)
    slide_bg(s)
    tb(s, 0.55, 0.35, 12.0, 0.65, "Option C — Tie the four capabilities to the diagram (one story)", size=22, bold=True)
    y = 1.2
    w = 2.85
    h = 0.95
    gap = 0.25
    x0 = 0.55
    stages = [
        ("Sync from sources", "Descriptions\n(programmatic tie to\nsource definitions)"),
        ("Scan & harvest", "Classifications\n(scan sources; auto-tags\n+ steward review)"),
        ("Curate in OMS", "Metrics\n(derive definitions;\nAI + policy reconcile\nconflicts)"),
        ("Discover", "Search & discovery\n(contextualize & personalize\nwith AI + policy)"),
    ]
    xs = []
    x = x0
    for title, body in stages:
        round_box(s, x, y, w, h, title, font_pt=11)
        xs.append(x)
        tb(s, x, y + h + 0.08, w, 1.35, body, size=11, bold=False, color=MUTED, align=PP_ALIGN.CENTER)
        x += w + gap
    tb(
        s,
        0.55,
        4.15,
        12.0,
        1.0,
        "Why: the bullets stop floating above a generic arrow — each capability has an obvious home in the flow.",
        size=14,
        color=MUTED,
    )
    footer(s, "Option C — capabilities mapped to stages")

    # --- Slide 5: Option D — MVP vs later ---
    s = prs.slides.add_slide(blank)
    slide_bg(s)
    tb(s, 0.55, 0.35, 12.0, 0.65, "Option D — MVP vs iterative (exec-scannable)", size=22, bold=True)
    round_box(s, 0.55, 1.2, 5.9, 2.55, "MVP (example)\n• Snowflake + Looker → OMS\n• Descriptions from DDL / sync\n• Basic search & browse\n• Steward edit + audit trail", font_pt=12)
    round_box(s, 6.75, 1.2, 5.95, 2.55, "Phase next\n• Automated classification scans\n• Metric reconciliation (AI + policy)\n• Personalized ranking in search\n• Deeper lineage / quality signals", font_pt=12)
    tb(s, 0.55, 4.0, 12.0, 0.45, "Optional: add one measurable outcome per column (e.g. % objects with sourced definitions).", size=13, color=ACCENT)
    tb(
        s,
        0.55,
        4.55,
        12.0,
        1.0,
        "Pair with a thin version of Option A or C above this split so “what we ship first” and “what grows in” share one slide.",
        size=14,
        color=MUTED,
    )
    footer(s, "Option D — MVP vs later")

    out = Path(__file__).resolve().parent.parent / "metadata-slide-layout-options.pptx"
    prs.save(str(out))
    return out


if __name__ == "__main__":
    path = build()
    print(path)
