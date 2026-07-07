#!/usr/bin/env python3
"""Build Word memo and PowerPoint deck for Unity Catalog vs. Alation replacement."""
from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from pptx import Presentation
from pptx.dml.color import RGBColor as PptxRGB
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches as PInches
from pptx.util import Pt as PPt

ROOT = Path(__file__).resolve().parents[1]
OUT_DOCX = ROOT / "Unity-Catalog-Alation-Replacement-Memo.docx"
OUT_PPTX = ROOT / "requirements-gathering" / "slides" / "Unity-Catalog-Alation-Replacement-Deck.pptx"

# --- Word helpers -----------------------------------------------------------------

NAVY = RGBColor(0x0F, 0x27, 0x44)
BLUE = RGBColor(0x1E, 0x5A, 0x96)
MUTED = RGBColor(0x5C, 0x65, 0x70)


def set_run_font(run, *, bold=False, size=11, color=None, italic=False):
    run.bold = bold
    run.italic = italic
    run.font.size = Pt(size)
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    if color:
        run.font.color.rgb = color


def add_heading(doc: Document, text: str, level: int = 1):
    p = doc.add_heading(text, level=level)
    for run in p.runs:
        run.font.name = "Calibri Light"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri Light")
        if level == 1:
            run.font.color.rgb = NAVY
    return p


def add_para(doc: Document, text: str = "", *, bold=False, size=11, color=None, italic=False):
    p = doc.add_paragraph()
    if text:
        run = p.add_run(text)
        set_run_font(run, bold=bold, size=size, color=color, italic=italic)
    return p


def add_bullets(doc: Document, items: list[str], size=11):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        run = p.add_run(item)
        set_run_font(run, size=size)


def add_table(doc: Document, headers: list[str], rows: list[list[str]]):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    hdr_cells = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr_cells[i].text = h
        for p in hdr_cells[i].paragraphs:
            for run in p.runs:
                set_run_font(run, bold=True, size=10, color=NAVY)
    for r_idx, row in enumerate(rows):
        cells = table.rows[r_idx + 1].cells
        for c_idx, val in enumerate(row):
            cells[c_idx].text = val
            for p in cells[c_idx].paragraphs:
                for run in p.runs:
                    set_run_font(run, size=10)
    doc.add_paragraph()
    return table


def add_block_quote(doc: Document, text: str):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.35)
    run = p.add_run(text)
    set_run_font(run, italic=True, size=11, color=BLUE)
    return p


def build_docx() -> Path:
    doc = Document()
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    title = doc.add_paragraph()
    title.alignment = WD_PARAGRAPH_ALIGNMENT.LEFT
    tr = title.add_run("MEMORANDUM")
    set_run_font(tr, bold=True, size=12, color=MUTED)

    add_heading(doc, "Can Unity Catalog Replace Alation?", level=0)

    meta = [
        ("Audience:", "Broad leadership and data platform stakeholders"),
        ("Author:", "Data Governance / Data Catalog program"),
        ("Date:", "June 25, 2026"),
        (
            "Context:",
            "Databricks Data + AI Summit 2026 (June 15–18, 2026); Alation contract end October 1, 2026",
        ),
    ]
    for label, value in meta:
        p = doc.add_paragraph()
        r1 = p.add_run(label + " ")
        set_run_font(r1, bold=True, size=10, color=MUTED)
        r2 = p.add_run(value)
        set_run_font(r2, size=10)

    doc.add_paragraph()

    add_heading(doc, "Executive summary", level=1)
    add_para(
        doc,
        "Unity Catalog (UC) is not a substitute for our Alation replacement. Databricks made meaningful "
        "progress at Data + AI Summit 2026 toward business-aware governance and agentic AI inside the "
        "lakehouse—especially Unity Catalog Metrics, Domains, Unity AI Gateway, and a teased Business "
        "Glossary (preview coming soon). Those capabilities are valuable for Databricks-centric workloads.",
    )
    add_para(
        doc,
        "They do not deliver the enterprise data catalog we use today: a cross-platform discovery and "
        "stewardship hub with a business glossary, articles, bulk governance, and workflows for "
        "non-technical users across Snowflake, pipelines, and BI—not only assets governed inside Databricks.",
    )
    p = add_para(doc)
    r1 = p.add_run("Recommendation: ")
    set_run_font(r1, bold=True, size=11, color=NAVY)
    r2 = p.add_run(
        "Continue the OMS + Navigator path as the Alation replacement. Treat Unity Catalog as a "
        "governed metadata source and execution-plane control for Databricks assets (hub-and-spoke), "
        "not as the enterprise catalog of record."
    )
    set_run_font(r2, size=11)

    add_heading(doc, "Why this question is surfacing now", level=1)
    add_para(
        doc,
        "At Data + AI Summit 2026, Databricks positioned Unity Catalog as the control plane for data, AI, and agents:",
    )
    add_table(
        doc,
        ["Announcement", "Status (per Databricks)", "What it does"],
        [
            [
                "Unity Catalog Metrics / Business Semantics",
                "Available / GA",
                "Governed KPIs and metric views at the data layer; SQL-addressable",
            ],
            [
                "Business Glossary",
                "Preview coming soon (not GA)",
                "Authoritative business concepts, terms, taxonomies—linked to data",
            ],
            ["Domains", "Public Preview", "Organize assets by business area; marketplace browse"],
            [
                "Unity AI Gateway",
                "Beta",
                "Runtime governance for models, agents, MCPs, tools",
            ],
            ["Governance Hub", "Private Preview", "Steward/admin command center for Databricks estate"],
            ["External lineage", "GA", "Ingest lineage from systems outside Databricks"],
            [
                "Certifications, ABAC, classification, DQ monitoring",
                "Beta / expanding",
                "Trust signals and automated policy enforcement",
            ],
            [
                "Cross-region / cross-cloud governance",
                "Preview roadmap",
                "Consistent policies across Databricks footprint",
            ],
        ],
    )
    add_para(
        doc,
        "Short answer for executives: UC is a strong platform governance layer for Databricks. "
        "It is not a drop-in replacement for Alation in a heterogeneous enterprise like ours.",
    )

    add_heading(doc, "What Unity Catalog is—and is not", level=1)
    add_heading(doc, "What UC does well", level=2)
    add_bullets(
        doc,
        [
            "Technical governance inside Databricks — permissions, auditing, lineage for tables, volumes, models, functions",
            "Semantic metrics layer — Metric Views define KPIs once with governance and reuse across SQL, notebooks, Genie, dashboards",
            "Lakehouse-native trust signals — certifications, classification, ABAC, data quality health indicators",
            "Discovery for Databricks users — Catalog Explorer and Discover (preview) within the Databricks experience",
            "Evolving cross-system lineage — external lineage/BYOL with engineering ownership of integrations",
        ],
    )

    add_heading(doc, "What UC is not (for our Alation parity bar)", level=2)
    add_table(
        doc,
        ["Alation capability", "Unity Catalog today", "Gap"],
        [
            [
                "Business glossary",
                "Not available today; Glossary announced as preview coming soon",
                "Critical gap",
            ],
            ["Articles / document hubs", "No equivalent", "Critical gap"],
            [
                "Cross-platform catalog",
                "Databricks-primary; BYOL does not replace managed crawlers",
                "Critical gap",
            ],
            ["Stewardship workflows", "Limited curation; no Alation-style workflow layer", "Major gap"],
            ["Bulk governance", "Tag policies inside UC only", "Major gap"],
            [
                "Consumer-friendly discovery",
                "Databricks UI; not Navigator on OMS",
                "Major gap",
            ],
        ],
    )
    add_para(
        doc,
        "Bottom line: UC does not provide Alation-style business glossary and enterprise catalog capabilities "
        "today—and Databricks only announced a Glossary preview at DAIS 2026, not a shipped product.",
        bold=True,
        color=NAVY,
    )

    add_heading(doc, "Fit for Disney Streaming's context", level=1)
    add_heading(doc, "Our estate is multi-platform", level=2)
    add_para(
        doc,
        "Alation federates metadata across Snowflake accounts, pipeline systems, and documentation—not only "
        "Unity-managed Delta tables. OMS already aggregates technical inventory across Snowflake, Delta, Hive, "
        "Harmony, and Airflow. UC enforcement stays coupled to Databricks compute.",
    )
    add_heading(doc, "Our committed direction", level=2)
    add_para(
        doc,
        "We are building OMS as the metadata store and Navigator as the business-first catalog experience for "
        "Alation exit by October 1, 2026. Pivoting to UC would fragment discovery, duplicate glossary work, "
        "miss the contract window, and shift cost to integration engineering.",
    )

    add_heading(doc, "Recommended architecture: complement, don't substitute", level=1)
    add_para(doc, "Hub-and-spoke model:")
    add_bullets(
        doc,
        [
            "Unity Catalog + 626/MCI → supply technical metadata",
            "OMS → federated metadata store (system of record for enterprise catalog)",
            "Navigator → business-first UX (Alation replacement)",
            "Ingest UC metadata into OMS so Databricks assets appear in the same search as Snowflake tables",
        ],
    )
    add_bullets(
        doc,
        [
            "One system of record per metadata type — avoid dual-writing glossary or PII tags",
            "Revisit when UC Glossary GA's — evaluate feeding OMS, not replacing the enterprise catalog",
        ],
    )

    add_heading(doc, "Talking points for leadership forums", level=1)
    add_para(doc, "If asked: “Can we just use Unity Catalog instead of building in-house?”", bold=True)
    add_block_quote(
        doc,
        "Unity Catalog got materially stronger at DAIS 2026 for governed metrics, domains, and AI agent "
        "governance, but it still doesn't give us a shipped business glossary, cross-platform catalog, or "
        "stewardship workflows across Snowflake and the rest of our stack. It's the right control plane inside "
        "Databricks; OMS and Navigator remain the right answer for replacing Alation by October.",
    )
    add_para(doc, "If asked: “Are we duplicating Databricks?”", bold=True)
    add_block_quote(
        doc,
        "No—if we integrate cleanly. UC governs lakehouse execution; OMS federates the estate; Navigator is the "
        "experience business users and stewards actually need.",
    )
    add_para(doc, "If asked: “Should we pause the OMS program?”", bold=True)
    add_block_quote(
        doc,
        "No. UC's new features reduce risk for Databricks metrics consistency, not for Alation exit.",
    )

    add_heading(doc, "What to watch on the Databricks roadmap", level=1)
    add_table(
        doc,
        ["Item", "Why it matters", "Action"],
        [
            ["UC Glossary (preview / GA)", "Could supply business term definitions", "Track; plan OMS ingestion"],
            ["Discover GA", "Databricks-only business user experience", "Useful for Databricks-heavy personas"],
            ["External lineage maturity", "Reduces lineage gaps", "Evaluate engineering cost vs OMS supply"],
            ["Metric Views in external BI", "KPI consistency on Databricks", "Complements metrics governance"],
        ],
    )

    add_heading(doc, "Conclusion", level=1)
    add_para(
        doc,
        "Databricks Unity Catalog is essential for how we govern data and AI on Databricks. It is insufficient "
        "as the enterprise replacement for Alation. We should accelerate, not redirect, the OMS + Navigator "
        "program—and define integration with Unity Catalog so Databricks assets are first-class citizens in "
        "the catalog we are already building.",
    )

    add_heading(doc, "Sources", level=1)
    sources = [
        "Databricks — What's new with Unity Catalog at Data + AI Summit 2026",
        "Databricks — Unity Catalog Business Semantics GA",
        "Databricks — UC business semantics documentation",
        "Architect Mindset — Is Unity Catalog Enough? (YouTube)",
        "Workspace: Product Brief, Feature Inventory Matrix, Alation Functionality Reference",
    ]
    add_bullets(doc, sources, size=10)

    p = doc.add_paragraph()
    p.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    run = p.add_run("Internal working document — Data Governance / Data Catalog program")
    set_run_font(run, italic=True, size=9, color=MUTED)

    doc.save(OUT_DOCX)
    return OUT_DOCX


# --- PowerPoint helpers (DEET-inspired) -------------------------------------------

DEET_NAVY = PptxRGB(0x15, 0x23, 0x42)
DEET_BLUE = PptxRGB(0x1E, 0x40, 0xAF)
DEET_SLATE = PptxRGB(0x33, 0x41, 0x55)
DEET_MUTED = PptxRGB(0x47, 0x55, 0x69)
DEET_FOOT_MUTED = PptxRGB(0x64, 0x74, 0x8B)
DEET_TOP_BG = PptxRGB(0xEE, 0xF2, 0xFF)
DEET_FOOTER_BG = PptxRGB(0xF1, 0xF5, 0xF9)
DEET_WHITE = PptxRGB(0xFF, 0xFF, 0xFF)
DEET_BORDER = PptxRGB(0xE2, 0xE8, 0xF0)
DEET_GREEN = PptxRGB(0x15, 0x6F, 0x4A)
DEET_TEAL = PptxRGB(0x0E, 0x94, 0x9A)
DEET_AMBER = PptxRGB(0xB4, 0x53, 0x09)
DEET_RED = PptxRGB(0xB9, 0x1C, 0x1C)
SLIDE_W = PInches(13.333333333333334)
FOOTER = "Data Governance  |  June 2026  |  Unity Catalog vs. Alation"


def ppt_para(p, text, *, size=11, color=DEET_SLATE, bold=False, font="Calibri", align=None):
    p.text = text
    p.font.size = PPt(size)
    p.font.color.rgb = color
    p.font.bold = bold
    p.font.name = font
    if align is not None:
        p.alignment = align


def send_back(slide, shape):
    sp_tree = slide.shapes._spTree  # noqa: SLF001
    el = shape._element  # noqa: SLF001
    sp_tree.remove(el)
    sp_tree.insert(2, el)


def rect(slide, left, top, width, height, fill, line=None, radius=False):
    kind = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(kind, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = line if line else fill
    if not line:
        shape.line.fill.background()
    return shape


def footer(slide):
    rect(slide, PInches(0), PInches(7.0), SLIDE_W, PInches(0.5), DEET_FOOTER_BG)
    tb = slide.shapes.add_textbox(PInches(0.5), PInches(7.08), PInches(12.33), PInches(0.34))
    ppt_para(tb.text_frame.paragraphs[0], FOOTER, size=9, color=DEET_FOOT_MUTED)


def title_slide(prs, title, subtitle):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    top = rect(slide, PInches(0), PInches(0), SLIDE_W, PInches(2.4), DEET_TOP_BG)
    send_back(slide, top)
    rect(slide, PInches(0), PInches(2.4), SLIDE_W, PInches(0.06), DEET_NAVY)
    tb = slide.shapes.add_textbox(PInches(0.7), PInches(2.9), PInches(11.9), PInches(1.4))
    ppt_para(tb.text_frame.paragraphs[0], title, size=40, color=DEET_NAVY, font="Calibri Light")
    st = slide.shapes.add_textbox(PInches(0.7), PInches(4.05), PInches(11.9), PInches(0.9))
    ppt_para(st.text_frame.paragraphs[0], subtitle, size=18, color=DEET_MUTED)
    footer(slide)
    return slide


def content_slide(prs, title, subtitle=None):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    header = rect(slide, PInches(0), PInches(0), SLIDE_W, PInches(0.55), DEET_TOP_BG)
    send_back(slide, header)
    tb = slide.shapes.add_textbox(PInches(0.5), PInches(0.88), PInches(12.5), PInches(0.44))
    ppt_para(tb.text_frame.paragraphs[0], title, size=24, color=DEET_NAVY, font="Calibri Light")
    top = PInches(1.36)
    if subtitle:
        st = slide.shapes.add_textbox(PInches(0.5), top, PInches(12.33), PInches(0.3))
        ppt_para(st.text_frame.paragraphs[0], subtitle, size=12, color=DEET_MUTED)
        top = PInches(1.72)
    footer(slide)
    return slide, top


def bullets(slide, left, top, width, height, items, size=11):
    tb = slide.shapes.add_textbox(left, top, width, height)
    tf = tb.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = f"• {item}"
        p.font.size = PPt(size)
        p.font.color.rgb = DEET_SLATE
        p.font.name = "Calibri"
        p.space_before = PPt(4)


def callout(slide, top, text, height=PInches(0.72)):
    rect(slide, PInches(0.5), top, PInches(12.33), height, PptxRGB(0xE0, 0xE7, 0xFF), line=DEET_BLUE, radius=True)
    tb = slide.shapes.add_textbox(PInches(0.65), top + PInches(0.1), PInches(12.0), height - PInches(0.15))
    tf = tb.text_frame
    tf.word_wrap = True
    ppt_para(tf.paragraphs[0], text, size=12, color=DEET_NAVY, bold=True, font="Calibri Light")


def two_col_slide(prs, title, left_title, left_items, right_title, right_items, callout_text=None):
    slide, top = content_slide(prs, title)
    lw = PInches(6.0)
    rect(slide, PInches(0.5), top, lw, PInches(5.1), DEET_WHITE, line=DEET_BORDER, radius=True)
    rect(slide, PInches(6.8), top, lw, PInches(5.1), DEET_WHITE, line=DEET_BORDER, radius=True)
    lt = slide.shapes.add_textbox(PInches(0.65), top + PInches(0.12), PInches(5.7), PInches(0.3))
    ppt_para(lt.text_frame.paragraphs[0], left_title, size=12, color=DEET_BLUE, bold=True)
    bullets(slide, PInches(0.65), top + PInches(0.45), PInches(5.7), PInches(4.5), left_items, size=10)
    rt = slide.shapes.add_textbox(PInches(6.95), top + PInches(0.12), PInches(5.7), PInches(0.3))
    ppt_para(rt.text_frame.paragraphs[0], right_title, size=12, color=DEET_BLUE, bold=True)
    bullets(slide, PInches(6.95), top + PInches(0.45), PInches(5.7), PInches(4.5), right_items, size=10)
    if callout_text:
        callout(slide, top + PInches(5.25), callout_text, height=PInches(0.55))
    return slide


def build_pptx() -> Path:
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = PInches(7.5)

    title_slide(
        prs,
        "Can Unity Catalog replace Alation?",
        "After DAIS 2026 (June 15–18), leadership is asking whether Unity Catalog can substitute for OMS + Navigator.\n"
        "Answer: No — UC complements OMS; it does not replace Alation.  |  Alation contract end: October 1, 2026",
    )

    two_col_slide(
        prs,
        "What Databricks announced (DAIS 2026)",
        "Semantics & discovery",
        [
            "Metric Views — governed KPIs at data layer (available)",
            "Business Glossary — preview coming soon (announced, not shipped)",
            "Domains — business-area marketplace (Public Preview)",
            "Governance Hub — steward command center (Private Preview)",
        ],
        "Agents & governance",
        [
            "Unity AI Gateway — models, agents, MCPs (Beta)",
            "External lineage GA; cross-region governance (roadmap)",
            "ABAC, classification, DQ monitoring (expanding)",
            "UC extends from technical governance toward business semantics—primarily inside Databricks",
        ],
        "Important: Glossary was announced but is not GA. What ships today is Metric Views + agent metadata—not a stewarded enterprise glossary.",
    )

    slide, top = content_slide(prs, "What Unity Catalog is good at")
    cols = [
        ("Governance engine", ["Fine-grained access, audit, lineage", "Models, volumes, functions", "ABAC, masking, classification"], DEET_BLUE),
        ("Semantic metrics", ["Define KPIs once; query via SQL", "Governed, lineaged definitions", "Strong fit for Genie + AI/BI"], DEET_TEAL),
        ("Evolving federation", ["Iceberg / REST interoperability", "External lineage for non-DBX systems", "Integration engineering cost is on us"], DEET_AMBER),
    ]
    cw = PInches(4.05)
    for i, (heading, items, accent) in enumerate(cols):
        left = PInches(0.5) + i * cw
        rect(slide, left, top, cw - PInches(0.12), PInches(5.0), DEET_WHITE, line=accent, radius=True)
        hd = slide.shapes.add_textbox(left + PInches(0.15), top + PInches(0.12), cw - PInches(0.25), PInches(0.35))
        ppt_para(hd.text_frame.paragraphs[0], heading, size=14, color=DEET_NAVY, bold=True, font="Calibri Light")
        bullets(slide, left + PInches(0.15), top + PInches(0.5), cw - PInches(0.25), PInches(4.3), items, size=10)

    two_col_slide(
        prs,
        "What Alation gives us today",
        "Capability pillars (must preserve)",
        [
            "Discovery — search & browse across sources",
            "Business context — glossary, articles, synonyms",
            "Trust — tags, domains, endorse/deprecate",
            "Lineage — tables, columns, dataflows",
            "Scale ops — bulk edit, catalog sets",
            "Governance — stewardship, policy workflows",
        ],
        "Our estate & users",
        [
            "Stewards, analysts, compliance, business consumers",
            "Multiple Snowflake accounts (DSS, Hulu, DTCI…)",
            "Pipelines: Airflow, dbt, 626/MCI supply",
            "Unity for lakehouse—not the only platform",
            "Enterprise catalogs federate metadata above heterogeneous systems",
        ],
    )

    slide, top = content_slide(prs, "Capability comparison (selected)")
    headers = ["Need", "Alation", "Unity Catalog", "Gap"]
    rows = [
        ["Business glossary", "Core", "Preview announced; not shipped", "Critical"],
        ["Articles / knowledge base", "Yes", "No", "Critical"],
        ["Cross-platform search", "Yes", "DBX-primary; BYOL = build/maintain", "Critical"],
        ["Stewardship workflows", "Yes", "Limited curation", "Major"],
        ["Bulk governance", "Mass edit, catalog sets", "Tag policies; not estate-wide", "Major"],
        ["Governed KPIs / metrics", "Glossary + queries", "Metric Views — strength", "Complement OMS"],
        ["DBX permissions & lineage", "Via connector", "Native — strength", "Feed OMS"],
    ]
    tbl = slide.shapes.add_table(len(rows) + 1, 4, PInches(0.5), top, PInches(12.33), PInches(5.0)).table
    col_w = [PInches(2.8), PInches(2.5), PInches(3.8), PInches(1.5)]
    for i, w in enumerate(col_w):
        tbl.columns[i].width = w
    for j, h in enumerate(headers):
        cell = tbl.cell(0, j)
        cell.text = h
        for p in cell.text_frame.paragraphs:
            ppt_para(p, h, size=9, color=DEET_NAVY, bold=True)
    for r, row in enumerate(rows):
        for c, val in enumerate(row):
            cell = tbl.cell(r + 1, c)
            cell.text = val
            color = DEET_RED if c == 3 and val == "Critical" else DEET_SLATE
            if c == 3 and val == "Major":
                color = DEET_AMBER
            if c == 3 and val in ("Complement OMS", "Feed OMS"):
                color = DEET_GREEN
            for p in cell.text_frame.paragraphs:
                ppt_para(p, val, size=8, color=color, bold=(c == 3))

    two_col_slide(
        prs,
        "Business glossary: the decisive gap",
        "What stewards need",
        [
            "Terms with definitions, synonyms, status",
            "Links to tables & columns",
            "Domain & stewardship ownership",
            "Searchable by business language",
            "Examples: Paid Subscribers; Total Time Streamed per Entitled Account",
        ],
        "What UC offers instead",
        [
            "Metric Views — SQL-defined KPIs with display names on measures/dimensions",
            "Agent metadata — helps Genie/AI interpret metrics",
            "Future Glossary — announced DAIS 2026; preview coming soon",
            "Metrics layer ≠ business glossary",
            "UC = structural engine; enterprise catalog = federation layer",
        ],
        "A metrics layer standardizes how KPIs are calculated. A business glossary standardizes what concepts mean—including terms that are not single SQL metrics.",
    )

    slide, top = content_slide(prs, "Recommended architecture — complement, don't substitute")
    arch = [
        ("Unity Catalog", "Execution-plane governance for Databricks: schemas, permissions, lineage, metric views.", DEET_TEAL),
        ("626 / MCI", "Technical metadata supply into OMS.", DEET_AMBER),
        ("OMS", "Federated metadata store — system of record for enterprise catalog.", DEET_BLUE),
        ("Navigator", "Business-first UX — Alation replacement: search, glossary, asset detail, stewardship.", DEET_GREEN),
    ]
    aw = PInches(3.05)
    for i, (name, body, accent) in enumerate(arch):
        left = PInches(0.5) + i * aw
        rect(slide, left, top, aw - PInches(0.08), PInches(2.2), DEET_WHITE, line=accent, radius=True)
        tb = slide.shapes.add_textbox(left + PInches(0.12), top + PInches(0.12), aw - PInches(0.2), PInches(1.95))
        tf = tb.text_frame
        tf.word_wrap = True
        ppt_para(tf.paragraphs[0], name, size=14, color=DEET_NAVY, bold=True, font="Calibri Light")
        ppt_para(tf.add_paragraph(), body, size=9, color=DEET_SLATE)
    callout(
        slide,
        top + PInches(2.45),
        "Principle: one system of record per metadata type — ingest UC → OMS; no dual-writing glossary in UC and OMS.",
        height=PInches(0.55),
    )
    bullets(
        slide,
        PInches(0.5),
        top + PInches(3.15),
        PInches(12.33),
        PInches(2.5),
        [
            "UC governs lakehouse execution; OMS federates the estate; Navigator is the experience business users need",
            "Revisit UC Glossary GA for OMS ingestion—not a program reset",
        ],
        size=10,
    )

    two_col_slide(
        prs,
        "Decision & asks",
        "Decision",
        [
            "Do NOT pivot Alation replacement to Unity Catalog",
            "DO continue OMS + Navigator",
            "Integrate UC as governed Databricks metadata source",
            "Do not pause for UC — Glossary preview not ready; contract ends Oct 1, 2026",
        ],
        "Asks & talking points",
        [
            "“Can UC replace Alation?” — Not for glossary, cross-platform discovery, or stewardship",
            "“Are we duplicating Databricks?” — Hub-and-spoke: UC feeds OMS",
            "Align on OMS + Navigator as Alation replacement strategy",
            "Charter UC → OMS integration for Databricks assets & metric views",
        ],
    )

    OUT_PPTX.parent.mkdir(parents=True, exist_ok=True)
    prs.save(OUT_PPTX)
    return OUT_PPTX


def main():
    docx_path = build_docx()
    pptx_path = build_pptx()
    print(f"Wrote {docx_path}")
    print(f"Wrote {pptx_path}")


if __name__ == "__main__":
    main()
