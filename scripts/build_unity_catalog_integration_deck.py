#!/usr/bin/env python3
"""Build reframed PowerPoint: Unity Catalog + in-house business data catalog (integration focus)."""
from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches, Pt

OUT = (
    Path(__file__).resolve().parents[1]
    / "requirements-gathering"
    / "slides"
    / "Unity-Catalog-Integration-Deck.pptx"
)

NAVY = RGBColor(0x15, 0x23, 0x42)
BLUE = RGBColor(0x1E, 0x40, 0xAF)
SLATE = RGBColor(0x33, 0x41, 0x55)
MUTED = RGBColor(0x47, 0x55, 0x69)
FOOT_MUTED = RGBColor(0x64, 0x74, 0x8B)
TOP_BG = RGBColor(0xEE, 0xF2, 0xFF)
FOOTER_BG = RGBColor(0xF1, 0xF5, 0xF9)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
BORDER = RGBColor(0xE2, 0xE8, 0xF0)
GREEN = RGBColor(0x15, 0x6F, 0x4A)
TEAL = RGBColor(0x0E, 0x94, 0x9A)
AMBER = RGBColor(0xB4, 0x53, 0x09)
RED = RGBColor(0xB9, 0x1C, 0x1C)
SW = Inches(13.333333333333334)
FOOTER = "Data Governance  |  July 2026  |  UC + Business Data Catalog"

# Body text formerly at 10pt — bumped for readability
BODY = 15
BODY_SM = 14
TABLE = 14


def para(p, text, *, size=BODY, color=SLATE, bold=False, font="Calibri"):
    p.text = text
    p.font.size = Pt(size)
    p.font.color.rgb = color
    p.font.bold = bold
    p.font.name = font


def send_back(slide, shape):
    tree = slide.shapes._spTree  # noqa: SLF001
    el = shape._element  # noqa: SLF001
    tree.remove(el)
    tree.insert(2, el)


def box(slide, left, top, width, height, fill, line=None, radius=False):
    kind = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    s = slide.shapes.add_shape(kind, left, top, width, height)
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    s.line.color.rgb = line or fill
    if not line:
        s.line.fill.background()
    return s


def slide_footer(slide):
    box(slide, Inches(0), Inches(7.0), SW, Inches(0.5), FOOTER_BG)
    tb = slide.shapes.add_textbox(Inches(0.5), Inches(7.08), Inches(12.33), Inches(0.34))
    para(tb.text_frame.paragraphs[0], FOOTER, size=9, color=FOOT_MUTED)


def title(prs, h, sub):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    top = box(slide, Inches(0), Inches(0), SW, Inches(2.4), TOP_BG)
    send_back(slide, top)
    box(slide, Inches(0), Inches(2.4), SW, Inches(0.06), NAVY)
    tb = slide.shapes.add_textbox(Inches(0.7), Inches(2.75), Inches(11.9), Inches(1.35))
    para(tb.text_frame.paragraphs[0], h, size=34, color=NAVY, font="Calibri Light")
    st = slide.shapes.add_textbox(Inches(0.7), Inches(4.05), Inches(11.9), Inches(1.35))
    tf = st.text_frame
    tf.word_wrap = True
    para(tf.paragraphs[0], sub, size=16, color=MUTED)
    slide_footer(slide)


def content(prs, h, sub=None):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box(slide, Inches(0), Inches(0), SW, Inches(0.55), TOP_BG)
    tb = slide.shapes.add_textbox(Inches(0.5), Inches(0.88), Inches(12.5), Inches(0.44))
    para(tb.text_frame.paragraphs[0], h, size=24, color=NAVY, font="Calibri Light")
    top = Inches(1.36)
    if sub:
        st = slide.shapes.add_textbox(Inches(0.5), top, Inches(12.33), Inches(0.4))
        tf = st.text_frame
        tf.word_wrap = True
        para(tf.paragraphs[0], sub, size=13, color=MUTED)
        top = Inches(1.78)
    slide_footer(slide)
    return slide, top


def blist(slide, left, top, width, height, items, size=BODY):
    tb = slide.shapes.add_textbox(left, top, width, height)
    tf = tb.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = f"• {item}"
        p.font.size = Pt(size)
        p.font.color.rgb = SLATE
        p.font.name = "Calibri"
        p.space_before = Pt(2)


def callout(slide, top, text, h=Inches(0.78)):
    box(slide, Inches(0.5), top, Inches(12.33), h, RGBColor(0xE0, 0xE7, 0xFF), line=BLUE, radius=True)
    tb = slide.shapes.add_textbox(Inches(0.65), top + Inches(0.08), Inches(12.0), h - Inches(0.12))
    tf = tb.text_frame
    tf.word_wrap = True
    para(tf.paragraphs[0], text, size=14, color=NAVY, bold=True, font="Calibri Light")


def two_col(prs, title_text, lt, li, rt, ri, note=None):
    slide, top = content(prs, title_text)
    w = Inches(6.0)
    box(slide, Inches(0.5), top, w, Inches(4.85), WHITE, line=BORDER, radius=True)
    box(slide, Inches(6.8), top, w, Inches(4.85), WHITE, line=BORDER, radius=True)
    lhd = slide.shapes.add_textbox(Inches(0.65), top + Inches(0.1), Inches(5.7), Inches(0.28))
    para(lhd.text_frame.paragraphs[0], lt, size=14, color=BLUE, bold=True)
    blist(slide, Inches(0.65), top + Inches(0.42), Inches(5.7), Inches(4.2), li)
    rhd = slide.shapes.add_textbox(Inches(6.95), top + Inches(0.1), Inches(5.7), Inches(0.28))
    para(rhd.text_frame.paragraphs[0], rt, size=14, color=BLUE, bold=True)
    blist(slide, Inches(6.95), top + Inches(0.42), Inches(5.7), Inches(4.2), ri)
    if note:
        callout(slide, top + Inches(5.0), note, h=Inches(0.62))
    return slide


def build():
    prs = Presentation()
    prs.slide_width = SW
    prs.slide_height = Inches(7.5)

    title(
        prs,
        "Unity Catalog + Business Data Catalog",
        "Session purpose: clarify and delineate the role and relationships of our internal services to Unity Catalog.\n\n"
        "Better together — not either/or. Unity Catalog governs Databricks execution. OMS + Navigator federate the "
        "estate for business users.\n\n"
        "This deck: inherent UC limits, where our internal services add value, integration patterns, and BYOL constraints.",
    )

    two_col(
        prs,
        "Two complementary roles — same metadata spine",
        "Unity Catalog (execution plane)",
        [
            "Govern permissions, audit, lineage on Databricks assets",
            "Metric Views — governed KPIs for SQL, Genie, AI/BI",
            "Enforce masking, ABAC, classification at query time",
            "Strongest where data is created & consumed on Databricks",
        ],
        "Internal services — OMS + Navigator",
        [
            "Single search across Snowflake, pipelines, BI, Databricks",
            "Business glossary, articles, stewardship workflows",
            "Bulk governance, domains, endorse/deprecate at estate scale",
            "Experience for stewards, analysts, compliance — not Catalog Explorer",
        ],
        "Strategy: UC is essential inside Databricks; our internal services are essential across the enterprise. Integrate both.",
    )

    slide, top = content(
        prs,
        "Inherent limits of Unity Catalog",
        "Why UC alone cannot replace our Alation exit — even after DAIS 2026 advances.",
    )
    limits = [
        ("Platform boundary", "Governance enforcement is coupled to Databricks compute — policies do not automatically apply in Snowflake, Tableau, or internal APIs.", AMBER),
        ("Vendor lock-in", "Business semantics, discovery UX, and agent context deepen dependence on the Databricks platform — harder to keep a tool-agnostic governance story across Snowflake and other estates.", RED),
        ("Business vocabulary", "Glossary announced as preview coming soon; today = Metric Views + agent metadata — not stewarded terms linked across the estate.", RED),
        ("Knowledge & docs", "No articles, document hubs, or collaborative IA equivalent to Alation.", RED),
        ("User experience", "Discover & Governance Hub are Databricks-native and maturing — not Navigator for business consumers.", AMBER),
        ("Estate federation", "UC's primary inventory is Databricks-governed assets; full cross-source discovery needs crawlers + OMS supply.", AMBER),
        ("Stewardship at scale", "Limited vs mass edit, catalog sets, documentation completion tracking across departments.", AMBER),
    ]
    cw, ch = Inches(4.05), Inches(1.22)
    for i, (hd, body, accent) in enumerate(limits):
        col, row = i % 3, i // 3
        left = Inches(0.5) + col * cw
        y = top + row * Inches(1.42)
        box(slide, left, y, cw - Inches(0.12), ch, WHITE, line=accent, radius=True)
        tb = slide.shapes.add_textbox(left + Inches(0.12), y + Inches(0.08), cw - Inches(0.25), ch - Inches(0.12))
        tf = tb.text_frame
        tf.word_wrap = True
        para(tf.paragraphs[0], hd, size=13, color=NAVY, bold=True, font="Calibri Light")
        para(tf.add_paragraph(), body, size=BODY_SM, color=SLATE)

    slide, top = content(prs, "Where our internal services deliver more value")
    rows = [
        ["Business glossary & synonyms", "Preview only in UC", "Core MVP in Navigator / OMS"],
        ["Cross-platform discovery", "Databricks-primary", "626 + OMS federated inventory"],
        ["Articles & onboarding docs", "Not available", "Alation parity path"],
        ["Stewardship workflows", "Limited curation", "Bulk edit, catalog sets, ownership"],
        ["Business-user UX", "Catalog Explorer / Discover", "Navigator — governed, persona-led"],
        ["Governed KPIs on Databricks", "Metric Views — UC strength", "Read/sync; federated metrics governance"],
        ["DBX permissions & lineage", "Native — UC strength", "Ingest UC → OMS; do not re-implement"],
    ]
    tbl = slide.shapes.add_table(len(rows) + 1, 3, Inches(0.5), top, Inches(12.33), Inches(5.0)).table
    for w, idx in zip([Inches(3.5), Inches(4.2), Inches(4.5)], range(3)):
        tbl.columns[idx].width = w
    for j, h in enumerate(["Capability", "Unity Catalog alone", "OMS + Navigator (internal services)"]):
        para(tbl.cell(0, j).text_frame.paragraphs[0], h, size=TABLE, color=NAVY, bold=True)
    for r, row in enumerate(rows):
        for c, val in enumerate(row):
            cell = tbl.cell(r + 1, c)
            color = GREEN if c == 2 else (RED if c == 1 and val.startswith(("Not", "Preview", "Limited", "Databricks")) else SLATE)
            if c == 1 and "strength" in val.lower():
                color = TEAL
            para(cell.text_frame.paragraphs[0], val, size=TABLE, color=color)

    slide, top = content(prs, "Integration playbook — service focus")
    playbook = [
        ("UC focus", "Table/column grants, row/column filters, native Databricks lineage, Metric Views SQL", TEAL),
        (
            "OMS focus",
            "Federated asset graph, cross-source lineage stitch, business descriptions projection for Navigator business data catalog",
            BLUE,
        ),
        (
            "Navigator focus",
            "Glossary, articles, search UX, stewardship signals, bulk business metadata operations",
            GREEN,
        ),
        (
            "Sync direction",
            "UC → OMS (technical metadata, lineage, metrics); governance repo → OMS (classifications)",
            AMBER,
        ),
    ]
    pw = Inches(3.05)
    for i, (label, body, accent) in enumerate(playbook):
        left = Inches(0.5) + i * pw
        box(slide, left, top, pw - Inches(0.08), Inches(2.55), WHITE, line=accent, radius=True)
        tb = slide.shapes.add_textbox(left + Inches(0.12), top + Inches(0.1), pw - Inches(0.22), Inches(2.3))
        tf = tb.text_frame
        tf.word_wrap = True
        para(tf.paragraphs[0], label, size=14, color=accent, bold=True)
        para(tf.add_paragraph(), body, size=BODY_SM, color=SLATE)
    callout(
        slide,
        top + Inches(2.68),
        "Engineering charter: UC → OMS connector + reconciliation reports. Not a second catalog UI inside Databricks.",
    )
    blist(
        slide,
        Inches(0.5),
        top + Inches(3.55),
        Inches(12.33),
        Inches(2.5),
        [
            "Reconciliation: periodic diff reports (UC vs OMS asset counts, stale lineage)",
            "Metric Views: expose in Navigator as governed KPI objects with link to UC definition",
            "When UC Glossary reaches Generally Available (GA): ingest into OMS — one glossary UX, not two",
        ],
    )

    two_col(
        prs,
        "Best ways to integrate with Unity Catalog",
        "Architecture principles",
        [
            "Hub-and-spoke: UC = Databricks System of Record (SoR) for schemas, grants, native lineage, metric views",
            "OMS = enterprise metadata store; Navigator = collaboration & discovery UX",
            "One System of Record (SoR) per metadata type — no dual-writing glossary or ownership",
            "Avoid metadata dual-write — pick where each attribute is authoritative",
        ],
        "Integration patterns (recommended)",
        [
            "Crawl / sync UC → OMS: tables, columns, grants summary, lineage, metric view definitions",
            "626 / MCI continues supply for non-Databricks sources into the same OMS spine",
            "Navigator surfaces UC-backed assets in unified search with Snowflake & pipeline objects",
            "When UC Glossary is Generally Available (GA): ingest into OMS — do not fork a second glossary UX",
            "Optional: propagate select tags (e.g. PII) one direction only — UC or OMS, not both editable",
        ],
        "Goal: Databricks assets are first-class in our catalog; UC stays authoritative for Databricks enforcement.",
    )

    two_col(
        prs,
        "External lineage & BYOL — what it is",
        "What BYOL / external lineage provides",
        [
            "Generally Available (GA) capability to register metadata & lineage from outside Databricks",
            "Programmatic ingestion via APIs — Salesforce, MySQL, Tableau, etc.",
            "Extends UC lineage graph with external objects for audit & impact analysis",
            "Useful for stitching lakehouse assets to upstream/downstream systems",
        ],
        "What BYOL is not",
        [
            "Not a managed enterprise catalog — you build & operate every integration",
            "Not automatic crawlers across 20+ sources like Alation / OMS supply model",
            "Not governance propagation — tags/classifications in BYOL don't enforce in Snowflake/BI",
            "Not a substitute for Navigator glossary, articles, or stewardship workflows",
        ],
        "BYOL narrows the lineage gap; it does not close the business catalog gap.",
    )

    slide, top = content(
        prs,
        "Problems & limits of BYOL in Unity Catalog",
        "Architectural constraints to plan for before relying on BYOL instead of OMS federation.",
    )
    byol = [
        ("Engineering tax", "Every source API change, schema drift, or tool upgrade can break your BYOL pipelines — ongoing maintenance, not one-time setup."),
        ("Incomplete coverage", "Hard to match Alation-grade depth: column lineage, dataflows (Airflow/dbt), BI reports, usage analytics without significant custom work."),
        ("Latency & freshness", "External metadata often polled or batch-loaded — not the same live operational picture as native UC or managed crawlers."),
        ("Governance silos", "Domains, certifications, and PII tags in UC don't translate to Snowflake Horizon or other platforms automatically."),
        ("Double governance risk", "Snowflake + Databricks = two permission models unless you explicitly sync — BYOL doesn't unify enforcement."),
        ("Vendor lock-in via BYOL", "Deepening UC as the lineage hub ties more governance context to Databricks — without delivering enterprise catalog or glossary in Navigator."),
        ("Cost shift", "Savings from not building Navigator can become headcount for integration teams maintaining BYOL graphs."),
    ]
    for i, (hd, body) in enumerate(byol):
        col, row = i % 2, i // 2
        left = Inches(0.5) + col * Inches(6.3)
        y = top + row * Inches(1.38)
        box(slide, left, y, Inches(6.04), Inches(1.28), WHITE, line=BORDER, radius=True)
        tb = slide.shapes.add_textbox(left + Inches(0.15), y + Inches(0.08), Inches(5.75), Inches(1.1))
        tf = tb.text_frame
        tf.word_wrap = True
        para(tf.paragraphs[0], hd, size=13, color=NAVY, bold=True, font="Calibri Light")
        para(tf.add_paragraph(), body, size=BODY_SM, color=SLATE)

    two_col(
        prs,
        "Strategy, talking points & asks",
        "How to talk about it",
        [
            "“We're not choosing UC or our internal services — we're integrating both.”",
            "“UC governs the lakehouse; Navigator lets business users find and trust data everywhere.”",
            "“BYOL helps lineage inside UC — it doesn't replace glossary, articles, or cross-platform search.”",
            "“We maintain connectors once in OMS — not N custom BYOL pipelines per source.”",
        ],
        "Asks",
        [
            "Endorse dual-platform strategy: UC + OMS/Navigator for Alation exit (2027, likely one-year extension)",
            "Charter UC → OMS integration (crawl, lineage, metric views)",
            "Define System of Record (SoR) matrix — prevent dual-write on glossary & ownership",
            "Track UC Glossary Generally Available (GA) for ingestion into OMS — not a standalone UC-only program",
        ],
        "Both platforms, one story: federated discovery and stewardship on internal services; governed execution on Databricks.",
    )

    OUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(OUT)
    return OUT


if __name__ == "__main__":
    path = build()
    print(f"Wrote {path}")
