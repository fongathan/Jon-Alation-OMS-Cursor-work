#!/usr/bin/env python3
"""One-slide talking sheet: BDC vs OMS persona review."""

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.util import Inches, Pt

OUT = (
    Path(__file__).resolve().parents[1]
    / "requirements-gathering"
    / "slides"
    / "BDC-vs-OMS-Persona-Review-Talking-Sheet.pptx"
)

INK = RGBColor(0x10, 0x1A, 0x2E)
INK_SOFT = RGBColor(0x33, 0x40, 0x5C)
MUTE = RGBColor(0x6B, 0x77, 0x8C)
MIST = RGBColor(0xF4, 0xF7, 0xFB)
PAPER = RGBColor(0xFF, 0xFF, 0xFF)
LINE = RGBColor(0xDD, 0xE4, 0xF0)
CLOUD = RGBColor(0xED, 0xF1, 0xF8)
BLUE = RGBColor(0x2D, 0x6C, 0xE0)
TEAL = RGBColor(0x10, 0x9A, 0xA6)
GREEN = RGBColor(0x1F, 0x9D, 0x6B)
CHIP_BG = RGBColor(0xEA, 0xF0, 0xFB)
HEAD = "Helvetica Neue"


def build() -> Presentation:
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    sw, sh = prs.slide_width, prs.slide_height
    s = prs.slides.add_slide(prs.slide_layouts[6])

    def rect(left, top, width, height, fill=None, line=None, lw=Pt(1)):
        shp = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
        shp.shadow.inherit = False
        if fill:
            shp.fill.solid()
            shp.fill.fore_color.rgb = fill
        else:
            shp.fill.background()
        if line:
            shp.line.color.rgb = line
            shp.line.width = lw
        else:
            shp.line.fill.background()
        return shp

    def txt(left, top, width, height, runs, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, spacing=1.05):
        box = s.shapes.add_textbox(left, top, width, height)
        tf = box.text_frame
        tf.word_wrap = True
        tf.auto_size = MSO_AUTO_SIZE.NONE
        tf.vertical_anchor = anchor
        first = True
        for item in runs:
            p = tf.paragraphs[0] if first else tf.add_paragraph()
            first = False
            p.alignment = item.get("align", align)
            p.space_after = item.get("gap", Pt(4))
            p.line_spacing = item.get("spacing", spacing)
            parts = item.get("parts") or [(item.get("text", ""), item)]
            for t, st in parts:
                r = p.add_run()
                r.text = t
                r.font.name = HEAD
                r.font.size = st.get("size", Pt(11))
                r.font.bold = st.get("bold", False)
                r.font.color.rgb = st.get("color", INK)
        return box

    rect(0, 0, sw, sh, MIST)
    txt(
        Inches(0.6),
        Inches(0.38),
        Inches(12),
        Inches(0.28),
        [{"text": "PERSONA REVIEW  \u00b7  DATA GOVERNANCE VS. DATA FOUNDATIONS", "size": Pt(11), "bold": True, "color": TEAL}],
    )
    txt(
        Inches(0.57),
        Inches(0.62),
        Inches(12.2),
        Inches(0.55),
        [{"text": "BDC on OMS — why one store, two optimized experiences", "size": Pt(28), "bold": True, "color": INK}],
    )

    banner = rect(Inches(0.6), Inches(1.28), Inches(12.18), Inches(0.62), CHIP_BG, BLUE, Pt(1))
    rect(Inches(0.6), Inches(1.28), Inches(0.12), Inches(0.62), BLUE)
    txt(
        Inches(0.92),
        Inches(1.28),
        Inches(11.7),
        Inches(0.62),
        [
            {
                "parts": [
                    ("Not a second metadata store. ", {"size": Pt(13.5), "bold": True, "color": INK}),
                    (
                        "OMS = system of record + APIs. Question = who owns the ",
                        {"size": Pt(13.5), "color": INK_SOFT},
                    ),
                    ("business & governance experience", {"size": Pt(13.5), "bold": True, "color": BLUE}),
                    (" for Alation exit.", {"size": Pt(13.5), "color": INK_SOFT}),
                ],
                "align": PP_ALIGN.LEFT,
            }
        ],
        anchor=MSO_ANCHOR.MIDDLE,
    )

    col_w = Inches(5.95)
    top = Inches(2.05)
    h = Inches(3.35)
    lx = Inches(0.6)
    rx = lx + col_w + Inches(0.33)

    rect(lx, top, col_w, h, PAPER, LINE)
    rect(lx, top, col_w, Inches(0.52), CLOUD)
    txt(
        lx + Inches(0.28),
        top,
        col_w - Inches(0.56),
        Inches(0.52),
        [{"text": "OMS  \u00b7  DATA FOUNDATIONS", "size": Pt(12), "bold": True, "color": MUTE, "align": PP_ALIGN.CENTER}],
        align=PP_ALIGN.CENTER,
        anchor=MSO_ANCHOR.MIDDLE,
    )
    oms = [
        "Primary: engineers, platform, integrators",
        "Job: How built? Lineage? SQL? Ops metadata?",
        "Win: platform completeness, one technical portal",
    ]
    txt(
        lx + Inches(0.3),
        top + Inches(0.62),
        col_w - Inches(0.6),
        h - Inches(0.75),
        [{"text": "\u2022  " + b, "size": Pt(11.5), "color": INK_SOFT, "gap": Pt(6)} for b in oms],
    )

    rect(rx, top, col_w, h, PAPER, BLUE, Pt(1.5))
    rect(rx, top, col_w, Inches(0.52), BLUE)
    txt(
        rx + Inches(0.28),
        top,
        col_w - Inches(0.56),
        Inches(0.52),
        [{"text": "BDC  \u00b7  DATA GOVERNANCE", "size": Pt(12), "bold": True, "color": PAPER, "align": PP_ALIGN.CENTER}],
        align=PP_ALIGN.CENTER,
        anchor=MSO_ANCHOR.MIDDLE,
    )
    bdc = [
        "Primary: business users, stewards, privacy/compliance",
        "Job: What does it mean? Trusted? Who owns it?",
        "Win: Alation-exit adoption, glossary, trust, stewardship UX",
    ]
    wr = []
    for b in bdc:
        wr.append(
            {
                "parts": [
                    ("\u25B8  ", {"size": Pt(11.5), "bold": True, "color": TEAL}),
                    (b, {"size": Pt(11.5), "color": INK}),
                ],
                "gap": Pt(6),
            }
        )
    txt(rx + Inches(0.3), top + Inches(0.62), col_w - Inches(0.6), h - Inches(0.75), wr)

    ask = rect(Inches(0.6), Inches(5.55), Inches(12.18), Inches(1.15), PAPER, GREEN, Pt(1))
    rect(Inches(0.6), Inches(5.55), Inches(0.12), Inches(1.15), GREEN)
    txt(
        Inches(0.92),
        Inches(5.68),
        Inches(11.7),
        Inches(0.9),
        [
            {
                "parts": [
                    ("THE ASK: ", {"size": Pt(12), "bold": True, "color": GREEN}),
                    ("Foundations owns ", {"size": Pt(12), "color": INK_SOFT}),
                    ("store & serve", {"size": Pt(12), "bold": True, "color": INK}),
                    ("  \u00b7  Governance owns ", {"size": Pt(12), "color": INK_SOFT}),
                    ("business-metadata experience", {"size": Pt(12), "bold": True, "color": BLUE}),
                    (" on OMS  \u00b7  ", {"size": Pt(12), "color": INK_SOFT}),
                    ("Open in OMS", {"size": Pt(12), "bold": True, "color": TEAL}),
                    (" for full technical detail", {"size": Pt(12), "color": INK_SOFT}),
                ],
                "spacing": 1.08,
            },
            {
                "text": "If one UI only: OMS must commit to business-first defaults + Alation-exit metrics with Governance in the design loop.",
                "size": Pt(10.5),
                "color": MUTE,
                "gap": Pt(0),
            },
        ],
    )

    rect(Inches(0.55), Inches(6.92), Inches(12.23), Pt(1), LINE)
    txt(
        Inches(0.55),
        Inches(7.02),
        Inches(12.23),
        Inches(0.3),
        [
            {
                "text": "626 = supply  \u00b7  OMS = store  \u00b7  BDC = governed business experience  \u00b7  DEEPT \u00b7  July 2026",
                "size": Pt(8.5),
                "color": MUTE,
                "align": PP_ALIGN.CENTER,
            }
        ],
        align=PP_ALIGN.CENTER,
    )
    return prs


def main() -> None:
    OUT.parent.mkdir(parents=True, exist_ok=True)
    build().save(OUT)
    print(f"saved {OUT}")


if __name__ == "__main__":
    main()
