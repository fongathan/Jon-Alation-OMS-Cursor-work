#!/usr/bin/env python3
"""Faithful readable rebuild of Architect Mindset UC video slides (no extra commentary)."""
from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches, Pt

OUT = (
    Path(__file__).resolve().parents[1]
    / "requirements-gathering"
    / "slides"
    / "Architect-Mindset-UC-Video-Rebuilt-Deck.pptx"
)

# Architect Mindset–style palette (dark navy canvas, light panels)
BG = RGBColor(0x0F, 0x17, 0x2A)
PANEL = RGBColor(0xF8, 0xFA, 0xFC)
PANEL_ALT = RGBColor(0xEE, 0xF2, 0xFF)
BORDER = RGBColor(0xCB, 0xD5, 0xE1)
TITLE = RGBColor(0xFF, 0xFF, 0xFF)
TEXT = RGBColor(0x1E, 0x29, 0x3B)
MUTED = RGBColor(0x64, 0x74, 0x8B)
BLUE = RGBColor(0x25, 0x63, 0xEB)
TEAL = RGBColor(0x0D, 0x94, 0x88)
GREEN = RGBColor(0x05, 0x96, 0x69)
RED = RGBColor(0xDC, 0x26, 0x26)
AMBER = RGBColor(0xD9, 0x77, 0x06)
BANNER_BG = RGBColor(0x1E, 0x3A, 0x8A)
SW = Inches(13.333333333333334)
SH = Inches(7.5)
TITLE_SZ = 26
HEADING_SZ = 16
BODY = 17
BODY_SM = 15
LABEL = 13
WATERMARK = "ArchitectMindset"


def p(paragraph, text, *, size=BODY, color=TEXT, bold=False, font="Calibri", align=None):
    paragraph.text = text
    paragraph.font.size = Pt(size)
    paragraph.font.color.rgb = color
    paragraph.font.bold = bold
    paragraph.font.name = font
    if align is not None:
        paragraph.alignment = align


def bg(slide):
    slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), SW, SH).fill.solid()
    slide.shapes[-1].fill.fore_color.rgb = BG
    slide.shapes[-1].line.fill.background()


def watermark(slide):
    tb = slide.shapes.add_textbox(Inches(11.4), Inches(7.05), Inches(1.7), Inches(0.3))
    p(tb.text_frame.paragraphs[0], WATERMARK, size=9, color=MUTED)


def panel(slide, left, top, width, height, *, fill=PANEL, line=BORDER):
    s = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    s.line.color.rgb = line
    return s


def heading(slide, left, top, width, text, *, color=BLUE, size=HEADING_SZ):
    tb = slide.shapes.add_textbox(left, top, width, Inches(0.35))
    p(tb.text_frame.paragraphs[0], text, size=size, color=color, bold=True)
    return tb


def bullets(slide, left, top, width, height, items, *, size=BODY, color=TEXT):
    tb = slide.shapes.add_textbox(left, top, width, height)
    tf = tb.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p(para, f"• {item}", size=size, color=color)
        para.space_before = Pt(3)


def label(slide, left, top, width, text, *, color=MUTED, size=LABEL, bold=True):
    tb = slide.shapes.add_textbox(left, top, width, Inches(0.28))
    p(tb.text_frame.paragraphs[0], text, size=size, color=color, bold=bold)


def banner(slide, top, text):
    panel(slide, Inches(0.45), top, Inches(12.43), Inches(0.58), fill=BANNER_BG, line=BANNER_BG)
    tb = slide.shapes.add_textbox(Inches(0.65), top + Inches(0.12), Inches(12.0), Inches(0.38))
    tf = tb.text_frame
    tf.word_wrap = True
    p(tf.paragraphs[0], text, size=16, color=TITLE, bold=True, font="Calibri Light")


def checks(slide, left, top, width, text):
    tb = slide.shapes.add_textbox(left, top, width, Inches(0.35))
    p(tb.text_frame.paragraphs[0], f"✓  {text}", size=BODY_SM, color=GREEN, bold=True)


def slide_title(slide, text):
    tb = slide.shapes.add_textbox(Inches(0.45), Inches(0.28), Inches(12.4), Inches(0.55))
    p(tb.text_frame.paragraphs[0], text, size=TITLE_SZ, color=TITLE, bold=True, font="Calibri Light")


def section_slide(prs, title, subtitle):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg(slide)
    tb = slide.shapes.add_textbox(Inches(0.8), Inches(2.35), Inches(11.7), Inches(1.2))
    p(tb.text_frame.paragraphs[0], title, size=36, color=TITLE, font="Calibri Light")
    st = slide.shapes.add_textbox(Inches(0.8), Inches(3.75), Inches(11.7), Inches(1.4))
    tf = st.text_frame
    tf.word_wrap = True
    p(tf.paragraphs[0], subtitle, size=22, color=MUTED)
    watermark(slide)


def content_slide(prs, title):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg(slide)
    slide_title(slide, title)
    watermark(slide)
    return slide


def col_box(slide, left, top, width, height, head, items, *, accent=BLUE, foot=None, subhead=None):
    panel(slide, left, top, width, height)
    y = top + Inches(0.12)
    heading(slide, left + Inches(0.14), y, width - Inches(0.28), head, color=accent)
    y += Inches(0.38)
    if subhead:
        label(slide, left + Inches(0.14), y, width - Inches(0.28), subhead, color=accent, bold=False)
        y += Inches(0.28)
    bullets(slide, left + Inches(0.14), y, width - Inches(0.28), height - Inches(0.75), items)
    if foot:
        label(slide, left + Inches(0.14), top + height - Inches(0.38), width - Inches(0.28), foot, color=MUTED)


def build():
    prs = Presentation()
    prs.slide_width = SW
    prs.slide_height = SH

    # --- SECTION: Intro ---
    section_slide(
        prs,
        "Is Unity Catalog Enough? Do You Still Need Enterprise Data Catalogs in 2026?",
        "Architect Mindset — your blueprint for modern systems.",
    )

    # Slide 1
    slide = content_slide(prs, "Do We Still Need Enterprise Data Catalogs?")
    top = Inches(1.05)
    cw, ch = Inches(4.35), Inches(3.55)
    col_box(
        slide, Inches(0.45), top, cw, ch,
        "Unity Catalog",
        ["Metadata", "Lineage", "Access Control", "Auditing"],
        accent=BLUE,
        foot="Platform-Native Governance",
    )
    col_box(
        slide, Inches(8.53), top, cw, ch,
        "Atlan / Alation",
        ["Search", "Discovery", "Glossary"],
        accent=TEAL,
        foot="External Catalog Platforms",
    )
    panel(slide, Inches(5.05), top + Inches(0.55), Inches(3.2), Inches(1.55), fill=PANEL_ALT)
    ct = slide.shapes.add_textbox(Inches(5.2), top + Inches(0.72), Inches(2.9), Inches(1.2))
    tf = ct.text_frame
    tf.word_wrap = True
    for i, line in enumerate(["Why Both?", "What Problem Do External Catalogs Solve?"]):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p(para, line, size=BODY, color=TEXT, bold=True)
    panel(slide, Inches(5.05), top + Inches(2.25), Inches(3.2), Inches(0.55), fill=PANEL_ALT)
    label(
        slide, Inches(5.2), top + Inches(2.35), Inches(2.9),
        "Capability Overlap Emerging", color=AMBER,
    )
    banner(slide, Inches(5.55), "Architectural Decisions Begin with the Right Question")

    # --- SECTION: Automated Data Discovery ---
    section_slide(
        prs,
        "Automated Data Discovery",
        "How Unity Catalog generates metadata — and why that still feeds enterprise catalogs.",
    )

    # Slide 2
    slide = content_slide(prs, "From Manual Discovery to Automated Data Discovery")
    top = Inches(1.0)
    label(slide, Inches(4.8), top - Inches(0.02), Inches(3.7), "Lakehouse · Pipelines · Data Assets", color=TITLE, size=BODY_SM)
    cw, ch = Inches(3.95), Inches(3.85)
    col_box(
        slide, Inches(0.45), top + Inches(0.25), cw, ch,
        "Manual Discovery Model",
        ["Manual Wikis", "Documentation", "Word of Mouth"],
        accent=RED,
    )
    label(slide, Inches(0.55), top + Inches(3.55), cw - Inches(0.2), "Stale, Incomplete, Non-Scalable", color=RED)
    col_box(
        slide, Inches(4.7), top + Inches(0.25), cw, ch,
        "Unity Catalog",
        ["Lineage Metadata", "Access Policies", "Data Classification"],
        accent=BLUE,
        subhead="CONTROL ENGINE",
    )
    col_box(
        slide, Inches(8.95), top + Inches(0.25), cw, ch,
        "Automated Discovery Systems",
        ["atlan — Search · Glossary · Discovery", "Alation — Search · Glossary · Discovery"],
        accent=GREEN,
    )
    checks(slide, Inches(9.05), top + Inches(3.15), cw, "Real-Time · Governed · Searchable")
    banner(slide, Inches(5.35), "Automated Discovery Replaces Manual Knowledge with Governed Metadata")

    # Slide 3
    slide = content_slide(prs, "What Automated Data Discovery Actually Does")
    top = Inches(1.05)
    cw, ch = Inches(3.95), Inches(3.75)
    col_box(slide, Inches(0.45), top, cw, ch, "DATA ECOSYSTEM", ["Tables", "Pipelines", "Data Assets"], accent=BLUE)
    col_box(
        slide, Inches(4.7), top, cw, ch,
        "AUTOMATED DISCOVERY ENGINE",
        ["Active Metadata Engine", "Metadata Capture", "Lineage Recording", "Data Classification", "Asset Indexing"],
        accent=TEAL,
    )
    col_box(
        slide, Inches(8.95), top, cw, ch,
        "ORGANIZATIONAL KNOWLEDGE LAYER",
        ["Searchable Data Ecosystem", "Discovered Assets", "Trusted Data", "Business Context"],
        accent=GREEN,
    )
    checks(slide, Inches(0.55), top + Inches(3.95), Inches(12.2), "Continuous Discovery · Context-Aware Data")
    banner(slide, Inches(5.35), "Discovery Transforms Technical Assets into Organizational Knowledge")

    # Slide 4
    slide = content_slide(prs, "Convergence: Platform Governance vs Enterprise Catalogs")
    top = Inches(1.05)
    col_box(
        slide, Inches(0.45), top, Inches(5.85), Inches(2.95),
        "Unity Catalog",
        ["Metadata Management", "Access Control", "Lineage Tracking", "Auditing"],
        accent=BLUE,
    )
    col_box(
        slide, Inches(7.03), top, Inches(5.85), Inches(2.95),
        "Atlan / Alation",
        ["Data Catalog", "Glossary", "Search", "Discovery"],
        accent=TEAL,
    )
    panel(slide, Inches(3.35), top + Inches(3.15), Inches(6.65), Inches(1.35), fill=PANEL_ALT)
    heading(slide, Inches(3.55), top + Inches(3.28), Inches(6.2), "Convergence", color=AMBER)
    bullets(
        slide, Inches(3.55), top + Inches(3.68), Inches(6.2), Inches(0.75),
        ["Metadata", "Lineage", "Discovery Signals"], size=BODY_SM,
    )
    label(slide, Inches(3.55), top + Inches(4.35), Inches(6.2), "Growing Capability Convergence", color=AMBER)
    banner(slide, Inches(5.35), "Platform-Native Governance Begins to Overlap with External Catalog Functions")

    # Slide 5
    slide = content_slide(prs, "Data Discovery with Trust Signals")
    top = Inches(1.05)
    cw, ch = Inches(3.95), Inches(3.75)
    col_box(slide, Inches(0.45), top, cw, ch, "DATA ECOSYSTEM", ["Datasets", "Dashboards", "ML Models"], accent=BLUE)
    col_box(
        slide, Inches(4.7), top, cw, ch,
        "SEARCH & DISCOVERY LAYER",
        ["Enterprise Search", "Unified Search Interface"],
        accent=TEAL,
    )
    col_box(
        slide, Inches(8.95), top, cw, ch,
        "TRUST EVALUATION SIGNALS",
        ["Certifications", "Popularity Metrics", "Lineage Visibility"],
        accent=GREEN,
        foot="Trust Signals",
    )
    checks(slide, Inches(0.55), top + Inches(3.95), Inches(12.2), "Reliable Data Selection · Confidence in Data")
    banner(slide, Inches(5.35), "Discovery is Guided by Trust Signals — Not Just Search Results")

    # Slide 6
    slide = content_slide(prs, "Turning Data into Business Understanding")
    top = Inches(1.05)
    cw, ch = Inches(3.95), Inches(3.75)
    col_box(
        slide, Inches(0.45), top, cw, ch,
        "DISCOVERED DATA",
        ["Data Assets", "Datasets", "Tables", "Models"],
        accent=BLUE,
    )
    col_box(
        slide, Inches(4.7), top, cw, ch,
        "CONTEXT ENRICHMENT LAYER",
        ["Business Context", "Domains", "Tags", "Descriptions"],
        accent=TEAL,
    )
    col_box(
        slide, Inches(8.95), top, cw, ch,
        "USER UNDERSTANDING",
        ["Business-Friendly Data", "Clear Meaning", "Usable Insights", "Context-Aware Data"],
        accent=GREEN,
    )
    checks(slide, Inches(0.55), top + Inches(3.95), Inches(12.2), "Improved Understanding · Faster Decision-Making")
    banner(slide, Inches(5.35), "Context Transforms Data into Meaningful Knowledge")

    # Slide 7
    slide = content_slide(prs, "Active Metadata: From Internal Signals to Enterprise Activation")
    top = Inches(1.05)
    cw, ch = Inches(3.95), Inches(3.75)
    col_box(
        slide, Inches(0.45), top, cw, ch,
        "DATA PLATFORM OPERATIONS",
        ["Query Execution", "Table Usage", "Schema Evolution"],
        accent=BLUE,
        foot="Operational Events",
    )
    col_box(
        slide, Inches(4.7), top, cw, ch,
        "UNITY CATALOG",
        ["Unity Catalog", "Continuous Metadata Generation", "Usage Signals", "Access Patterns", "Lineage Events"],
        accent=TEAL,
        subhead="ACTIVE METADATA ENGINE",
        foot="Automated Metadata Synchronization →",
    )
    col_box(
        slide, Inches(8.95), top, cw, ch,
        "ENTERPRISE ECOSYSTEM",
        ["Atlan / Alation", "Governance Systems", "Discovery Platforms"],
        accent=GREEN,
        foot="Enterprise Metadata Consumers",
    )
    checks(slide, Inches(0.55), top + Inches(3.95), Inches(12.2), "Real-time Visibility · Enterprise-Wide Discovery")
    banner(slide, Inches(5.35), "Metadata Becomes Valuable Only When Activated Beyond the Platform")

    # --- SECTION: Why Dedicated Enterprise Catalog? ---
    section_slide(
        prs,
        "Why Dedicated Enterprise Catalog?",
        "What external catalogs solve that Unity Catalog alone does not — federation, adoption, workflows, and end-to-end lineage.",
    )

    # Slide 8
    slide = content_slide(prs, "Why Enterprise Catalogs Were Created: Federation of Metadata")
    top = Inches(1.05)
    cw, ch = Inches(3.95), Inches(3.75)
    col_box(
        slide, Inches(0.45), top, cw, ch,
        "Distributed Enterprise Systems",
        ["SAP", "Oracle", "Snowflake", "ETL Tools", "BI Platforms"],
        accent=BLUE,
        foot="Fragmented Visibility",
    )
    col_box(
        slide, Inches(4.7), top, cw, ch,
        "Enterprise Data Catalog",
        ["Metadata Aggregation", "Cross-System Visibility"],
        accent=TEAL,
        subhead="Federation Layer",
        foot="Federation — Not Just Governance",
    )
    col_box(
        slide, Inches(8.95), top, cw, ch,
        "Unified Metadata View",
        ["Cross-Platform Discovery", "End-to-End Visibility"],
        accent=GREEN,
        foot="Complete Visibility",
    )
    banner(slide, Inches(5.35), "No Single System Could See Everything — Federation Solves That")

    # Slide 9
    slide = content_slide(prs, "The Adoption Gap: Technical Platforms vs Business Users")
    top = Inches(1.05)
    cw, ch = Inches(3.95), Inches(3.75)
    col_box(
        slide, Inches(0.45), top, cw, ch,
        "DATABRICKS EXPERIENCE",
        ["Notebooks", "SQL Workspaces", "Catalog Explorer"],
        accent=BLUE,
        subhead="Lakehouse Interface",
        foot="Technical Environment",
    )
    col_box(
        slide, Inches(4.7), top, cw, ch,
        "Usability Barrier",
        ["Technical Complexity", "Navigation Friction", "Low Adoption", "Limited Accessibility", "Underutilized Data"],
        accent=RED,
    )
    col_box(
        slide, Inches(8.95), top, cw, ch,
        "BUSINESS USERS",
        ["Business Analysts", "Product Managers", "Compliance Officers"],
        accent=GREEN,
        foot="Non-Technical Users",
    )
    banner(slide, Inches(5.35), "Technical Interfaces Limit Business Adoption")

    # Slide 10
    slide = content_slide(prs, "Enterprise Catalogs as the Collaboration Layer")
    top = Inches(1.05)
    cw, ch = Inches(3.95), Inches(4.05)
    col_box(
        slide, Inches(0.45), top, cw, ch,
        "TECHNICAL METADATA SOURCE",
        ["Unity Catalog", "Metadata", "Lineage", "Policies"],
        accent=BLUE,
    )
    col_box(
        slide, Inches(4.7), top, cw, ch,
        "COLLABORATION PLATFORM — Atlan / Alation",
        [
            "Business-Friendly Interface",
            "Business Glossary",
            "Data Stewardship Workflows",
            "Data Product Marketplace",
            "Certification & Badges",
            "Certified for Financial Audits",
        ],
        accent=TEAL,
    )
    col_box(
        slide, Inches(8.95), top, cw, ch,
        "BUSINESS USERS",
        ["Analysts", "Managers", "Compliance Teams"],
        accent=GREEN,
    )
    checks(slide, Inches(0.55), top + Inches(4.25), Inches(12.2), "Accessible Data · Trusted Assets · Self-Service Discovery")
    banner(slide, Inches(5.55), "Enterprise Catalogs Translate Technical Metadata into Business Value")

    # Slide 11
    slide = content_slide(prs, "Logging Changes Is Not the Same as Governing Them")
    top = Inches(1.05)
    cw, ch = Inches(3.95), Inches(3.75)
    col_box(
        slide, Inches(0.45), top, cw, ch,
        "DATA CHANGE EVENT",
        ["Schema / Data Change", "Definition Update", "Structural Change"],
        accent=BLUE,
    )
    col_box(
        slide, Inches(4.7), top, cw, ch,
        "UNITY CATALOG",
        ["Unity Catalog", "Change Logging", "Lineage Updates", "Audit Records"],
        accent=TEAL,
    )
    col_box(
        slide, Inches(8.95), top, cw, ch,
        "MISSING GOVERNANCE WORKFLOWS",
        ["Ownership Notification", "Approval Processes", "Compliance Sign-Off", "Data Quality Discussions"],
        accent=RED,
        foot="Not Natively Orchestrated",
    )
    label(slide, Inches(4.95), top + Inches(1.55), Inches(3.5), "Workflow Gap", color=AMBER, size=BODY)
    banner(slide, Inches(5.35), "Capturing Changes Does Not Ensure Governance Execution")

    # Slide 12
    slide = content_slide(prs, "Enterprise Catalogs as the Governance Workflow Layer")
    top = Inches(1.05)
    cw, ch = Inches(3.95), Inches(3.95)
    col_box(
        slide, Inches(0.45), top, cw, ch,
        "METADATA INPUT",
        ["Unity Catalog", "Metadata", "Change Events", "Lineage"],
        accent=BLUE,
    )
    col_box(
        slide, Inches(4.7), top, cw, ch,
        "WORKFLOW ORCHESTRATION LAYER",
        [
            "Atlan / Alation",
            "Governance Workflows",
            "Ownership Assignment",
            "Approval Workflows",
            "Documentation Tracking",
            "Stewardship Processes",
        ],
        accent=TEAL,
    )
    col_box(
        slide, Inches(8.95), top, cw, ch,
        "TRUSTED DATA PUBLICATION",
        ["Certified Data Assets", "Approved Data", "Documented Assets", "Policy-Compliant"],
        accent=GREEN,
    )
    checks(slide, Inches(0.55), top + Inches(4.15), Inches(12.2), "Controlled Releases · Accountability · Compliance-Ready Data")
    banner(slide, Inches(5.55), "Governance Requires Workflows — Not Just Metadata Visibility")

    # Slide 13
    slide = content_slide(prs, "Limitation: Lineage Visibility Stops at the Lakehouse Boundary")
    top = Inches(1.05)
    cw, ch = Inches(3.95), Inches(3.75)
    col_box(
        slide, Inches(0.45), top, cw, ch,
        "UPSTREAM SYSTEMS",
        ["Fivetran", "External Sources"],
        accent=BLUE,
        subhead="Data Ingestion Layer",
    )
    col_box(
        slide, Inches(4.7), top, cw, ch,
        "Databricks Lakehouse",
        ["Unity Catalog", "Tables → Lineage → Policies"],
        accent=TEAL,
    )
    col_box(
        slide, Inches(8.95), top, cw, ch,
        "DOWNSTREAM SYSTEMS",
        ["dbt Transformations", "Snowflake", "BI Tools"],
        accent=GREEN,
        subhead="Consumption & Distribution Layer",
    )
    label(slide, Inches(4.95), top + Inches(0.35), Inches(3.5), "Partial Lineage Visibility", color=RED, size=BODY)
    label(slide, Inches(4.95), top + Inches(3.55), Inches(3.5), "Partial Lineage Visibility", color=RED, size=BODY)
    banner(slide, Inches(5.35), "Platform-Native Lineage Does Not Extend Across the Enterprise")

    # Slide 14
    slide = content_slide(prs, "Enterprise Catalogs as the End-to-End Lineage Control Plane")
    top = Inches(1.0)
    col_box(
        slide, Inches(0.45), top, Inches(12.43), Inches(1.05),
        "ENTERPRISE CONTROL PLANE",
        ["Atlan / Alation", "Lineage Ingestion", "Cross-System Stitching", "Metadata Federation"],
        accent=BLUE,
    )
    flow_top = top + Inches(1.2)
    fw, fh = Inches(3.05), Inches(2.55)
    flow = [
        ("UPSTREAM SYSTEMS", ["Fivetran", "External Databases"]),
        ("DATABRICKS", ["Lakehouse", "Unity Catalog Lineage"]),
        ("TRANSFORMATION & DISTRIBUTION", ["dbt", "Snowflake", "Data Pipelines"]),
        ("CONSUMPTION LAYER", ["Tableau", "Power BI"]),
    ]
    for i, (hd, items) in enumerate(flow):
        col_box(slide, Inches(0.45) + i * fw, flow_top, fw - Inches(0.08), fh, hd, items, accent=TEAL)
    checks(
        slide, Inches(0.55), flow_top + Inches(2.7), Inches(12.2),
        "End-to-End Lineage · Cross-Platform Visibility · Single Source of Truth",
    )
    banner(slide, Inches(5.35), "Enterprise Visibility Requires Federated Metadata Across All Systems")

    # Slide 15
    slide = content_slide(prs, "A Shift in the Architecture: Bring Your Own Lineage (BYOL)")
    top = Inches(1.05)
    cw, ch = Inches(3.95), Inches(3.75)
    col_box(
        slide, Inches(0.45), top, cw, ch,
        "BEFORE: PREVIOUS STATE",
        ["Atlan / Alation"],
        accent=BLUE,
        foot="Strong Federation Layer",
    )
    panel(slide, Inches(4.7), top, cw, ch, fill=PANEL_ALT)
    heading(slide, Inches(4.9), top + Inches(0.55), cw - Inches(0.4), "Architectural Shift", color=AMBER)
    label(slide, Inches(4.9), top + Inches(1.15), cw - Inches(0.4), "June 8, 2026 — GA Release", color=MUTED, bold=False, size=BODY_SM)
    col_box(
        slide, Inches(8.95), top, cw, ch,
        "NOW: EVOLVING CAPABILITY",
        ["Unity Catalog", "New Capability: BYOL"],
        accent=GREEN,
        foot="Expanding Beyond Native Boundaries",
    )
    banner(slide, Inches(5.35), "Platform Capabilities Are Rapidly Expanding Into Enterprise Territory")

    OUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(OUT)
    return OUT, len(prs.slides)


if __name__ == "__main__":
    path, count = build()
    print(f"Wrote {path} ({count} slides)")
