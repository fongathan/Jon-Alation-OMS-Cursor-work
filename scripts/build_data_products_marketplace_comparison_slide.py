#!/usr/bin/env python3
"""One-slide comparison: data products & marketplace — Alation vs Snowflake vs OMS/Navigator."""

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.util import Inches, Pt

OUT = (
    Path(__file__).resolve().parents[1]
    / "requirements-gathering"
    / "slides"
    / "Data-Products-Marketplace-Alation-Snowflake-OMS.pptx"
)

INK = RGBColor(0x10, 0x1A, 0x2E)
INK_SOFT = RGBColor(0x33, 0x40, 0x5C)
MUTE = RGBColor(0x6B, 0x77, 0x8C)
MIST = RGBColor(0xF4, 0xF7, 0xFB)
PAPER = RGBColor(0xFF, 0xFF, 0xFF)
LINE = RGBColor(0xDD, 0xE4, 0xF0)
CLOUD = RGBColor(0xED, 0xF1, 0xF8)
BLUE = RGBColor(0x2D, 0x6C, 0xE0)
TEAL = RGBColor(0x10, 0x9A, 0xA6)
GREEN = RGBColor(0x1F, 0x9D, 0x6B)
AMBER = RGBColor(0xCF, 0x8A, 0x00)
INDIGO = RGBColor(0x1E, 0x2B, 0x52)
CHIP_BG = RGBColor(0xEA, 0xF0, 0xFB)
HEAD = "Helvetica Neue"


def build() -> Presentation:
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    sw, sh = prs.slide_width, prs.slide_height
    s = prs.slides.add_slide(prs.slide_layouts[6])

    def rect(left, top, width, height, fill=None, line=None, lw=Pt(1)):
        shp = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
        shp.shadow.inherit = False
        if fill:
            shp.fill.solid()
            shp.fill.fore_color.rgb = fill
        else:
            shp.fill.background()
        if line:
            shp.line.color.rgb = line
            shp.line.width = lw
        else:
            shp.line.fill.background()
        return shp

    def txt(left, top, width, height, runs, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, spacing=1.05):
        box = s.shapes.add_textbox(left, top, width, height)
        tf = box.text_frame
        tf.word_wrap = True
        tf.auto_size = MSO_AUTO_SIZE.NONE
        tf.vertical_anchor = anchor
        first = True
        for item in runs:
            p = tf.paragraphs[0] if first else tf.add_paragraph()
            first = False
            p.alignment = item.get("align", align)
            p.space_after = item.get("gap", Pt(3))
            p.line_spacing = item.get("spacing", spacing)
            parts = item.get("parts") or [(item.get("text", ""), item)]
            for t, st in parts:
                r = p.add_run()
                r.text = t
                r.font.name = HEAD
                r.font.size = st.get("size", Pt(11))
                r.font.bold = st.get("bold", False)
                r.font.color.rgb = st.get("color", INK)
        return box

    def column_card(left, title, accent, subtitle, sections, highlight=False):
        w = Inches(3.95)
        top = Inches(2.02)
        h = Inches(3.72)
        border = accent if highlight else LINE
        bw = Pt(1.5) if highlight else Pt(1)
        rect(left, top, w, h, PAPER, border, bw)
        rect(left, top, w, Inches(0.56), accent if highlight else CLOUD)
        txt(
            left + Inches(0.2),
            top,
            w - Inches(0.4),
            Inches(0.56),
            [
                {
                    "text": title,
                    "size": Pt(11.5),
                    "bold": True,
                    "color": PAPER if highlight else INK,
                    "align": PP_ALIGN.CENTER,
                }
            ],
            align=PP_ALIGN.CENTER,
            anchor=MSO_ANCHOR.MIDDLE,
        )
        y = top + Inches(0.66)
        txt(
            left + Inches(0.22),
            y,
            w - Inches(0.44),
            Inches(0.34),
            [{"text": subtitle, "size": Pt(9.5), "bold": True, "color": accent}],
        )
        y += Inches(0.34)
        for label, bullets in sections:
            txt(
                left + Inches(0.22),
                y,
                w - Inches(0.44),
                Inches(0.2),
                [{"text": label, "size": Pt(8.5), "bold": True, "color": MUTE}],
            )
            y += Inches(0.2)
            for bullet in bullets:
                txt(
                    left + Inches(0.22),
                    y,
                    w - Inches(0.44),
                    Inches(0.38),
                    [{"text": "\u2022  " + bullet, "size": Pt(9.5), "color": INK_SOFT, "gap": Pt(1)}],
                )
                y += Inches(0.36)

    rect(0, 0, sw, sh, MIST)
    txt(
        Inches(0.6),
        Inches(0.38),
        Inches(12),
        Inches(0.28),
        [{"text": "DATA PRODUCTS & MARKETPLACE  \u00b7  PLATFORM COMPARISON", "size": Pt(11), "bold": True, "color": TEAL}],
    )
    txt(
        Inches(0.57),
        Inches(0.62),
        Inches(12.2),
        Inches(0.55),
        [{"text": "Alation vs Snowflake vs OMS/Navigator — who does what?", "size": Pt(27), "bold": True, "color": INK}],
    )

    banner = rect(Inches(0.6), Inches(1.28), Inches(12.18), Inches(0.58), CHIP_BG, BLUE, Pt(1))
    rect(Inches(0.6), Inches(1.28), Inches(0.12), Inches(0.58), BLUE)
    txt(
        Inches(0.92),
        Inches(1.28),
        Inches(11.7),
        Inches(0.58),
        [
            {
                "parts": [
                    ("Different layers, not substitutes. ", {"size": Pt(13), "bold": True, "color": INK}),
                    (
                        "Alation = business trust & discovery  \u00b7  Snowflake = access & delivery  \u00b7  ",
                        {"size": Pt(13), "color": INK_SOFT},
                    ),
                    ("OMS/Navigator = our target federated experience", {"size": Pt(13), "bold": True, "color": BLUE}),
                    (".", {"size": Pt(13), "color": INK_SOFT}),
                ],
            }
        ],
        anchor=MSO_ANCHOR.MIDDLE,
    )

    x1 = Inches(0.6)
    x2 = Inches(4.68)
    x3 = Inches(8.76)

    column_card(
        x1,
        "ALATION",
        MUTE,
        "Catalog-native: use-case offerings",
        [
            (
                "Data product",
                [
                    "YAML-defined, governed package for a business outcome",
                    "Links catalog assets + stewards + usage rules",
                ],
            ),
            (
                "Marketplace",
                [
                    "Internal storefront: browse, badges, request access",
                    "Built for humans and AI agents to find trusted products",
                ],
            ),
            (
                "Disney today",
                [
                    "360s = domain containers (not always Data Products App)",
                    "Marketplace capability only partially adopted",
                ],
            ),
        ],
    )

    column_card(
        x2,
        "SNOWFLAKE",
        TEAL,
        "Platform-native: shareable packages",
        [
            (
                "Data product",
                [
                    "Listing backed by secure share, Native App, or Connected App",
                    "Technical package consumers install or request access to",
                ],
            ),
            (
                "Marketplace",
                [
                    "Internal Marketplace = intra-org discovery & sharing",
                    "Public Marketplace = external/partner/paid distribution",
                ],
            ),
            (
                "Strength",
                [
                    "RBAC, masking, row policies, actual data delivery",
                    "Scoped to Snowflake accounts and listings",
                ],
            ),
        ],
    )

    column_card(
        x3,
        "OMS / NAVIGATOR",
        BLUE,
        "Target: federated business experience",
        [
            (
                "Data product",
                [
                    "Governed product pages on 360/domain taxonomy",
                    "Glossary, stewards, trust signals, cross-platform assets",
                ],
            ),
            (
                "Marketplace",
                [
                    "Navigator browse/search/trust UX for business users",
                    "Federate access requests to Snowflake Internal Marketplace",
                ],
            ),
            (
                "Pattern",
                [
                    "Discover & trust in Navigator; consume via Snowflake",
                    "OMS = metadata store; not a second delivery plane",
                ],
            ),
        ],
        highlight=True,
    )

    compare = rect(Inches(0.6), Inches(5.88), Inches(12.18), Inches(0.72), PAPER, LINE)
    headers = ["Dimension", "Alation", "Snowflake", "OMS / Navigator"]
    rows = [
        ("Starts from", "Business use case", "Technical share/listing", "Consumer trust + federation"),
        ("Access enforced", "Request workflow", "Snowflake RBAC & grants", "Hand off to Snowflake"),
        ("Scope", "Cross-platform catalog", "Snowflake estate", "Enterprise catalog on OMS"),
    ]
    col_x = [Inches(0.78), Inches(2.55), Inches(5.35), Inches(8.15)]
    col_w = [Inches(1.65), Inches(2.65), Inches(2.65), Inches(4.35)]
    y0 = Inches(5.98)
    for i, h in enumerate(headers):
        txt(
            col_x[i],
            y0,
            col_w[i],
            Inches(0.2),
            [{"text": h, "size": Pt(8.5), "bold": True, "color": MUTE}],
        )
    y = y0 + Inches(0.22)
    for row in rows:
        for i, cell in enumerate(row):
            color = INK if i == 0 else INK_SOFT
            bold = i == 0
            txt(
                col_x[i],
                y,
                col_w[i],
                Inches(0.18),
                [{"text": cell, "size": Pt(9), "bold": bold, "color": color, "gap": Pt(0)}],
            )
        y += Inches(0.16)

    ask = rect(Inches(0.6), Inches(6.68), Inches(12.18), Inches(0.42), CHIP_BG, GREEN, Pt(1))
    rect(Inches(0.6), Inches(6.68), Inches(0.12), Inches(0.42), GREEN)
    txt(
        Inches(0.92),
        Inches(6.68),
        Inches(11.7),
        Inches(0.42),
        [
            {
                "parts": [
                    ("TARGET PATTERN: ", {"size": Pt(10.5), "bold": True, "color": GREEN}),
                    ("Navigator = discover & trust", {"size": Pt(10.5), "bold": True, "color": INK}),
                    ("  \u00b7  ", {"size": Pt(10.5), "color": MUTE}),
                    ("Snowflake Internal Marketplace = access delivery", {"size": Pt(10.5), "bold": True, "color": INK}),
                    ("  \u00b7  ", {"size": Pt(10.5), "color": MUTE}),
                    ("360/domain taxonomy = accountability layer", {"size": Pt(10.5), "color": INK_SOFT}),
                ],
            }
        ],
        anchor=MSO_ANCHOR.MIDDLE,
    )

    rect(Inches(0.55), Inches(7.12), Inches(12.23), Pt(1), LINE)
    txt(
        Inches(0.55),
        Inches(7.18),
        Inches(12.23),
        Inches(0.24),
        [
            {
                "text": "Business Data Catalog  \u00b7  Alation exit  \u00b7  DEEPT  \u00b7  September 2026",
                "size": Pt(8.5),
                "color": MUTE,
                "align": PP_ALIGN.CENTER,
            }
        ],
        align=PP_ALIGN.CENTER,
    )
    return prs


def main() -> None:
    OUT.parent.mkdir(parents=True, exist_ok=True)
    build().save(OUT)
    print(f"saved {OUT}")


if __name__ == "__main__":
    main()
