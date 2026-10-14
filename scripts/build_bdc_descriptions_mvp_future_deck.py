#!/usr/bin/env python3
"""Business Data Catalog — descriptions & business metadata (MVP vs future) deck.

Matches the Navigator Business Data Catalog Vision deck styling
(make_navigator_bdc_vision.py compass theme).
"""

from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.util import Inches, Pt

OUT = (
    Path(__file__).resolve().parents[1]
    / "requirements-gathering"
    / "slides"
    / "Business-Data-Catalog-Descriptions-MVP-Future.pptx"
)

PERSONAS_DOC_URL = (
    "https://docs.google.com/document/d/1PPIbiore1hCrWVSgmGtvCBqVYW4AzB5LzugbKPrCGo4/edit"
)

# Palette — aligned with make_navigator_bdc_vision.py
INK = RGBColor(0x10, 0x1A, 0x2E)
INK_SOFT = RGBColor(0x33, 0x40, 0x5C)
MUTE = RGBColor(0x6B, 0x77, 0x8C)
PAPER = RGBColor(0xFF, 0xFF, 0xFF)
MIST = RGBColor(0xF4, 0xF7, 0xFB)
CLOUD = RGBColor(0xED, 0xF1, 0xF8)
LINE = RGBColor(0xDD, 0xE4, 0xF0)
INDIGO = RGBColor(0x1E, 0x2B, 0x52)
INDIGO2 = RGBColor(0x2C, 0x47, 0x8C)
BLUE = RGBColor(0x2D, 0x6C, 0xE0)
TEAL = RGBColor(0x10, 0x9A, 0xA6)
VIOLET = RGBColor(0x6A, 0x57, 0xD6)
AMBER = RGBColor(0xCF, 0x8A, 0x00)
GREEN = RGBColor(0x1F, 0x9D, 0x6B)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
CHIP_BG = RGBColor(0xEA, 0xF0, 0xFB)
HEAD_FONT = "Helvetica Neue"
BODY_FONT = "Helvetica Neue"
FOOTER = (
    "Navigator  \u00b7  Business Data Catalog  \u00b7  Data Governance  "
    "\u00b7  DEEPT \u00b7  2026   |   Confidential"
)

ALATION_METADATA = [
    "Table & column descriptions (business context, formulas, examples)",
    "Business titles, synonyms, and glossary term links",
    "Business glossary — term definitions, aliases, business formulas",
    "Articles & knowledge pages linked to catalog objects",
    "Custom fields — business steward, data domain, business segment, etc.",
    "Tags, domains, and compliance classifications (e.g. GDPR / CCPA scoped)",
    "Trust flags — endorse, warn, deprecate",
    "Ownership — business & technical stewards on tables, columns, and terms",
]

BDC_METADATA = [
    "Table & column descriptions on asset detail (business + technical layers)",
    "Business glossary — terms, definitions, synonyms, linked to catalog assets",
    "Ownership & stewardship — owners, stewards, escalation paths",
    "Tags & classifications — domain, sensitivity, compliance scope",
    "Trust signals — endorse / warn / deprecate on assets",
    "Custom business fields carried from OMS / migration (where modeled)",
    "Provenance badges — who wrote it, when updated, migration vs curated",
    "Policies & defined-vs-observed context (where governance is modeled)",
]


def build_presentation() -> Presentation:
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    sw, sh = prs.slide_width, prs.slide_height

    def slide():
        return prs.slides.add_slide(prs.slide_layouts[6])

    def rect(s, left, top, width, height, fill=None, line=None, line_w=None, shape=MSO_SHAPE.RECTANGLE):
        shp = s.shapes.add_shape(shape, left, top, width, height)
        shp.shadow.inherit = False
        if fill is None:
            shp.fill.background()
        else:
            shp.fill.solid()
            shp.fill.fore_color.rgb = fill
        if line is None:
            shp.line.fill.background()
        else:
            shp.line.color.rgb = line
            shp.line.width = line_w or Pt(1)
        return shp

    def gradient(shp, c1, c2, angle=90):
        shp.fill.gradient()
        stops = shp.fill.gradient_stops
        stops[0].position = 0.0
        stops[0].color.rgb = c1
        stops[1].position = 1.0
        stops[1].color.rgb = c2
        try:
            shp.fill.gradient_angle = angle
        except Exception:
            pass
        shp.line.fill.background()
        shp.shadow.inherit = False
        return shp

    def text(s, left, top, width, height, runs, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, space_after=Pt(5), line_spacing=1.0):
        box = s.shapes.add_textbox(left, top, width, height)
        tf = box.text_frame
        tf.word_wrap = True
        tf.auto_size = MSO_AUTO_SIZE.NONE
        tf.vertical_anchor = anchor
        tf.margin_left = 0
        tf.margin_right = 0
        tf.margin_top = 0
        tf.margin_bottom = 0
        first = True
        for item in runs:
            p = tf.paragraphs[0] if first else tf.add_paragraph()
            first = False
            p.alignment = item.get("align", align)
            p.space_after = item.get("space_after", space_after)
            p.space_before = item.get("space_before", Pt(0))
            p.line_spacing = item.get("line_spacing", line_spacing)
            parts = item.get("parts")
            if parts is None:
                parts = [(item.get("text", ""), item)]
            for txt, st in parts:
                r = p.add_run()
                r.text = txt
                f = r.font
                f.size = st.get("size", Pt(14))
                f.bold = st.get("bold", False)
                f.italic = st.get("italic", False)
                f.color.rgb = st.get("color", INK)
                f.name = st.get("font", BODY_FONT)
                url = st.get("url")
                if url:
                    r.hyperlink.address = url
        return box

    def bg(s, color=MIST):
        rect(s, 0, 0, sw, sh, fill=color)

    def footer(s):
        rect(s, Inches(0.55), Inches(6.92), Inches(12.23), Pt(1), fill=LINE)
        text(
            s,
            Inches(0.55),
            Inches(7.02),
            Inches(12.23),
            Inches(0.3),
            [{"text": FOOTER, "size": Pt(8.5), "color": MUTE, "align": PP_ALIGN.CENTER}],
            align=PP_ALIGN.CENTER,
        )

    def title_block(s, kicker, title, kicker_color=BLUE, subtitle=None):
        text(
            s,
            Inches(0.6),
            Inches(0.42),
            Inches(12),
            Inches(0.3),
            [{"text": kicker, "size": Pt(11), "bold": True, "color": kicker_color, "font": HEAD_FONT}],
        )
        text(
            s,
            Inches(0.57),
            Inches(0.66),
            Inches(12.2),
            Inches(0.9),
            [{"text": title, "size": Pt(30), "bold": True, "color": INK, "font": HEAD_FONT}],
        )
        if subtitle:
            text(
                s,
                Inches(0.6),
                Inches(1.22),
                Inches(12),
                Inches(0.4),
                [{"text": subtitle, "size": Pt(13.5), "italic": True, "color": MUTE}],
            )

    def bullets(items, color=INK_SOFT, size=Pt(12.5), gap=Pt(8), marker="\u2022  "):
        return [
            {
                "text": marker + item,
                "size": size,
                "color": color,
                "space_after": gap,
                "line_spacing": 1.06,
            }
            for item in items
        ]

    def speaker_slide(kicker, title, kicker_color=BLUE):
        s = slide()
        bg(s)
        title_block(s, kicker, title, kicker_color)
        rect(s, Inches(0.6), Inches(1.55), Inches(12.18), Inches(4.95), fill=PAPER, line=LINE, line_w=Pt(1))
        rect(s, Inches(0.6), Inches(1.55), Inches(0.14), Inches(4.95), fill=kicker_color)
        text(
            s,
            Inches(0.95),
            Inches(3.0),
            Inches(11.6),
            Inches(1.2),
            [
                {
                    "text": "Alation Metadata Example",
                    "size": Pt(18),
                    "bold": True,
                    "color": MUTE,
                    "font": HEAD_FONT,
                    "align": PP_ALIGN.CENTER,
                }
            ],
            align=PP_ALIGN.CENTER,
            anchor=MSO_ANCHOR.MIDDLE,
        )
        footer(s)

    # --- Title slide (vision deck layout) ---
    s = slide()
    bg(s, PAPER)
    panel = rect(s, 0, 0, Inches(4.35), sh)
    gradient(panel, INDIGO, INDIGO2, angle=120)
    rect(s, Inches(4.35), 0, Inches(0.06), sh, fill=TEAL)
    rect(
        s,
        Inches(1.05),
        Inches(1.45),
        Inches(2.05),
        Inches(2.05),
        fill=None,
        line=TEAL,
        line_w=Pt(2.25),
        shape=MSO_SHAPE.OVAL,
    )
    text(
        s,
        Inches(1.05),
        Inches(1.52),
        Inches(2.05),
        Inches(1.9),
        [
            {
                "text": "N",
                "size": Pt(78),
                "bold": True,
                "color": WHITE,
                "font": HEAD_FONT,
                "align": PP_ALIGN.CENTER,
            }
        ],
        align=PP_ALIGN.CENTER,
        anchor=MSO_ANCHOR.MIDDLE,
    )
    text(
        s,
        Inches(0.85),
        Inches(3.85),
        Inches(2.6),
        Inches(0.4),
        [
            {
                "text": "N A V I G A T O R",
                "size": Pt(15),
                "bold": True,
                "color": WHITE,
                "font": HEAD_FONT,
                "align": PP_ALIGN.CENTER,
            }
        ],
        align=PP_ALIGN.CENTER,
    )
    text(
        s,
        Inches(0.85),
        Inches(4.22),
        Inches(2.6),
        Inches(0.4),
        [
            {
                "text": "Business Data Catalog",
                "size": Pt(11.5),
                "color": RGBColor(0xB9, 0xC6, 0xE6),
                "font": HEAD_FONT,
                "align": PP_ALIGN.CENTER,
            }
        ],
        align=PP_ALIGN.CENTER,
    )

    rx = Inches(4.95)
    text(
        s,
        rx,
        Inches(1.30),
        Inches(7.9),
        Inches(0.4),
        [
            {
                "text": "DESCRIPTIONS & BUSINESS METADATA  \u00b7  MVP VS. FUTURE STATE",
                "size": Pt(12),
                "bold": True,
                "color": TEAL,
                "font": HEAD_FONT,
            }
        ],
    )
    text(
        s,
        rx - Inches(0.03),
        Inches(1.66),
        Inches(8.0),
        Inches(2.4),
        [
            {
                "text": "Table / Column Descriptions & Other Business Metadata",
                "size": Pt(34),
                "bold": True,
                "color": BLUE,
                "font": HEAD_FONT,
                "space_after": Pt(4),
                "line_spacing": 1.05,
            },
            {
                "text": "MVP vs. Future state",
                "size": Pt(26),
                "bold": True,
                "color": INDIGO2,
                "font": HEAD_FONT,
                "line_spacing": 1.05,
            },
        ],
    )
    text(
        s,
        rx,
        Inches(4.35),
        Inches(7.85),
        Inches(0.95),
        [
            {
                "text": "What we migrate from Alation for MVP — and how descriptions and "
                "other business metadata can be prepopulated automatically in the future.",
                "size": Pt(14.5),
                "color": INK_SOFT,
                "italic": True,
                "line_spacing": 1.12,
            }
        ],
    )
    text(
        s,
        rx,
        Inches(6.78),
        Inches(7.85),
        Inches(0.4),
        [
            {
                "text": "July 2026",
                "size": Pt(12.5),
                "bold": True,
                "color": TEAL,
                "font": HEAD_FONT,
            }
        ],
    )

    # --- Slide 1: What are descriptions and other business metadata? ---
    s = slide()
    bg(s)
    title_block(
        s,
        "BUSINESS METADATA",
        "What are descriptions and other business metadata?",
        TEAL,
        "Types we use today in Alation — and the same business metadata surface in the Business Data Catalog on OMS",
    )

    col_w = Inches(5.95)
    col_top = Inches(1.78)
    col_h = Inches(4.95)
    lx = Inches(0.6)
    rx_col = lx + col_w + Inches(0.33)

    rect(s, lx, col_top, col_w, col_h, fill=PAPER, line=LINE, line_w=Pt(1))
    rect(s, lx, col_top, col_w, Inches(0.62), fill=CLOUD)
    text(
        s,
        lx + Inches(0.3),
        col_top,
        col_w - Inches(0.6),
        Inches(0.62),
        [{"text": "IN ALATION TODAY", "size": Pt(13), "bold": True, "color": MUTE, "font": HEAD_FONT}],
        anchor=MSO_ANCHOR.MIDDLE,
    )
    text(
        s,
        lx + Inches(0.32),
        col_top + Inches(0.85),
        col_w - Inches(0.64),
        col_h - Inches(1.0),
        bullets(ALATION_METADATA, size=Pt(11.5), gap=Pt(7)),
    )

    rect(s, rx_col, col_top, col_w, col_h, fill=PAPER, line=BLUE, line_w=Pt(1.5))
    rect(s, rx_col, col_top, col_w, Inches(0.62), fill=BLUE)
    text(
        s,
        rx_col + Inches(0.3),
        col_top,
        col_w - Inches(0.6),
        Inches(0.62),
        [
            {
                "text": "IN BUSINESS DATA CATALOG (VIA OMS)",
                "size": Pt(13),
                "bold": True,
                "color": WHITE,
                "font": HEAD_FONT,
            }
        ],
        anchor=MSO_ANCHOR.MIDDLE,
    )
    wr = []
    for item in BDC_METADATA:
        wr.append(
            {
                "parts": [
                    ("\u25B8  ", {"size": Pt(11.5), "bold": True, "color": TEAL}),
                    (item, {"size": Pt(11.5), "color": INK}),
                ],
                "space_after": Pt(7),
                "line_spacing": 1.06,
            }
        )
    text(
        s,
        rx_col + Inches(0.32),
        col_top + Inches(0.85),
        col_w - Inches(0.64),
        col_h - Inches(1.0),
        wr,
    )
    footer(s)

    # --- Slide 2: MVP Alation Migration (speaker-led) ---
    speaker_slide("MVP", "MVP: Alation Migration", BLUE)

    # --- Slide 3: Future state ---
    s = slide()
    bg(s)
    title_block(
        s,
        "FUTURE STATE",
        "Future: Prepopulated descriptions and other metadata automatically from scanning the data sources directly",
        VIOLET,
    )
    rect(s, Inches(0.6), Inches(1.55), Inches(12.18), Inches(0.82), fill=CHIP_BG, line=VIOLET, line_w=Pt(1))
    rect(s, Inches(0.6), Inches(1.55), Inches(0.14), Inches(0.82), fill=VIOLET)
    text(
        s,
        Inches(0.95),
        Inches(1.55),
        Inches(11.6),
        Inches(0.82),
        [
            {
                "parts": [
                    (
                        "Beyond migration: metadata supply scans source systems and repositories ",
                        {"size": Pt(14.5), "bold": True, "color": INK},
                    ),
                    (
                        "so stewards start from suggested descriptions and context — not blank fields.",
                        {"size": Pt(14.5), "color": INK_SOFT},
                    ),
                ],
                "line_spacing": 1.08,
            }
        ],
        anchor=MSO_ANCHOR.MIDDLE,
    )

    future_bullets = [
        "Alation Suggest Description Example",
        "626 as possible solution",
        "Other possible options (Tech team \u2014 TBD)",
    ]
    cy = Inches(2.75)
    for i, item in enumerate(future_bullets):
        y = cy + i * Inches(1.15)
        accent = [BLUE, TEAL, AMBER][i]
        rect(s, Inches(0.6), y, Inches(12.18), Inches(0.95), fill=PAPER, line=LINE, line_w=Pt(1))
        rect(s, Inches(0.6), y, Inches(0.12), Inches(0.95), fill=accent)
        text(
            s,
            Inches(0.95),
            y,
            Inches(11.6),
            Inches(0.95),
            [
                {
                    "text": f"{i + 1}.  {item}",
                    "size": Pt(16),
                    "color": INK,
                    "font": HEAD_FONT,
                }
            ],
            anchor=MSO_ANCHOR.MIDDLE,
        )
    footer(s)

    # --- Slide 4: Personas ---
    s = slide()
    bg(s)
    title_block(s, "WHO WE SERVE", "Personas", GREEN)
    rect(s, Inches(0.6), Inches(1.75), Inches(12.18), Inches(4.55), fill=PAPER, line=LINE, line_w=Pt(1))
    rect(s, Inches(0.6), Inches(1.75), Inches(0.14), Inches(4.55), fill=GREEN)
    text(
        s,
        Inches(0.95),
        Inches(2.15),
        Inches(11.6),
        Inches(1.0),
        [
            {
                "text": "Prioritized issues, use cases, and persona context from Alation interviews "
                "and OMS migration discussions.",
                "size": Pt(15),
                "color": INK_SOFT,
                "line_spacing": 1.12,
            }
        ],
    )
    text(
        s,
        Inches(0.95),
        Inches(3.35),
        Inches(11.6),
        Inches(0.5),
        [{"text": "Full personas & use-case matrix:", "size": Pt(14), "bold": True, "color": INK, "font": HEAD_FONT}],
    )
    text(
        s,
        Inches(0.95),
        Inches(3.95),
        Inches(11.6),
        Inches(0.6),
        [
            {
                "parts": [
                    (
                        "Alation_Prioritized_Issues_and_Use_Cases",
                        {
                            "size": Pt(16),
                            "bold": True,
                            "color": BLUE,
                            "font": HEAD_FONT,
                            "url": PERSONAS_DOC_URL,
                        },
                    )
                ]
            }
        ],
    )
    text(
        s,
        Inches(0.95),
        Inches(4.65),
        Inches(11.6),
        Inches(0.8),
        [
            {
                "text": PERSONAS_DOC_URL,
                "size": Pt(11),
                "color": MUTE,
                "url": PERSONAS_DOC_URL,
            }
        ],
    )
    footer(s)

    # --- Slide 5: End ---
    s = slide()
    bg(s, PAPER)
    stripe = rect(s, 0, Inches(3.55), sw, Inches(0.12), fill=TEAL)
    stripe  # noqa: B018 — layout anchor
    text(
        s,
        Inches(0.57),
        Inches(2.55),
        Inches(12.2),
        Inches(1.0),
        [{"text": "End", "size": Pt(54), "bold": True, "color": INK, "font": HEAD_FONT, "align": PP_ALIGN.CENTER}],
        align=PP_ALIGN.CENTER,
        anchor=MSO_ANCHOR.MIDDLE,
    )
    footer(s)

    return prs


def main() -> None:
    out = OUT
    out.parent.mkdir(parents=True, exist_ok=True)
    prs = build_presentation()
    prs.save(out)
    print(f"saved {out} — {len(prs.slides)} slides")


if __name__ == "__main__":
    main()
