#!/usr/bin/env python3
"""Build a non-technical overview of Yzma (.docx)."""
from __future__ import annotations

from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "Yzma-How-It-Works-Simplified.docx"

NAVY = RGBColor(0x1E, 0x2B, 0x52)
INDIGO = RGBColor(0x2C, 0x47, 0x8C)
MUTED = RGBColor(0x6B, 0x77, 0x8C)


def set_run_font(run, *, bold=False, size=11, color=None, italic=False) -> None:
    run.bold = bold
    run.italic = italic
    run.font.size = Pt(size)
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    if color:
        run.font.color.rgb = color


def add_heading(doc: Document, text: str, level: int = 1) -> None:
    p = doc.add_heading(text, level=level)
    for run in p.runs:
        run.font.name = "Calibri Light"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri Light")
        if level <= 2:
            run.font.color.rgb = NAVY


def add_para(doc: Document, text: str, *, bold: bool = False, italic: bool = False) -> None:
    p = doc.add_paragraph()
    run = p.add_run(text)
    set_run_font(run, bold=bold, italic=italic)


def add_bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        run = p.add_run(item)
        set_run_font(run)


def add_table(doc: Document, headers: list[str], rows: list[list[str]]) -> None:
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    for i, h in enumerate(headers):
        table.rows[0].cells[i].text = h
        for para in table.rows[0].cells[i].paragraphs:
            for run in para.runs:
                set_run_font(run, bold=True, size=10, color=NAVY)
    for r_idx, row in enumerate(rows):
        for c_idx, val in enumerate(row):
            table.rows[r_idx + 1].cells[c_idx].text = val
            for para in table.rows[r_idx + 1].cells[c_idx].paragraphs:
                for run in para.runs:
                    set_run_font(run, size=10)
    doc.add_paragraph()


def build() -> Path:
    doc = Document()
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    label = doc.add_paragraph()
    label.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    set_run_font(label.add_run("Plain-language guide"), bold=True, size=12, color=MUTED)

    title = doc.add_paragraph()
    title.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    set_run_font(title.add_run("How Yzma Works"), bold=True, size=22, color=NAVY)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    set_run_font(
        subtitle.add_run("Internal knowledge base for Content Data Solutions"),
        size=13,
        color=INDIGO,
    )

    meta = doc.add_paragraph()
    meta.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    set_run_font(
        meta.add_run(f"Draft · {date.today().isoformat()} · yzma.dssinfra.com"),
        size=10,
        color=MUTED,
    )

    doc.add_paragraph()
    add_para(
        doc,
        "Yzma is a company-internal website that brings together written documentation about "
        "data products, metrics, pipelines, and related concepts—and adds an AI assistant that "
        "can answer questions using that documentation. Think of it as a curated wiki plus a "
        "helpful colleague who reads the wiki for you and points to the exact pages that support "
        "each answer.",
    )

    add_heading(doc, "What problem does it solve?", level=2)
    add_bullets(
        doc,
        [
            "Scattered knowledge: Important definitions and “which table should I use?” guidance "
            "lived in many places. Yzma centralizes it.",
            "Hard to search: Keyword search alone does not always surface the right page. "
            "The chat feature understands questions in everyday language.",
            "Different audiences: Analysts, product partners, and engineers need different "
            "entry points. Yzma offers browsing paths tuned for how people work.",
        ],
    )

    add_heading(doc, "Where does the knowledge come from?", level=2)
    add_para(
        doc,
        "The wiki is not a side feature—it is the foundation Yzma runs on. Browse, graphs, guides, "
        "and especially chat all use the same collection of team documentation stored as Markdown "
        "files in the yzma GitHub repository (the wiki/ folder). The chat assistant does not "
        "invent pipeline facts from memory; it retrieves and reads those pages, then summarizes them.",
    )
    add_para(
        doc,
        "Pages are built and maintained through a mix of human curation and tooling:",
    )
    add_bullets(
        doc,
        [
            "Authors write and update pages in git—tables, metrics, concepts, query patterns, and "
            "how objects link together—using a standard page format (OKF) with title, short description, "
            "tags, and cross-links so search and chat can find the right material.",
            "When CDS ships new data products or models, documentation work is often a deliberate "
            "effort: engineers add pages or expand existing ones in pull requests, similar to shipping code.",
            "Wiki Sync watches selected upstream pipeline repositories. When merged code changes might "
            "affect documentation, Yzma can suggest wiki edits; a person still reviews and accepts or "
            "dismisses those suggestions.",
            "Quality checks are expanding—extracted “ground truth” from pipeline repos, PR review helpers "
            "for wiki changes, and planned automated tests on common questions—to catch drift between "
            "docs and code.",
        ],
    )
    add_para(
        doc,
        "Before Yzma, much of this knowledge lived in engineers’ heads, scattered Confluence, or "
        "comments in SQL. Yzma centralizes it, but keeping it accurate remains ongoing stewardship work.",
    )

    add_heading(doc, "What drives reliability (and what does not)", level=2)
    add_para(
        doc,
        "A smooth chat experience depends first on the wiki corpus, then on retrieval and the AI layer:",
    )
    add_table(
        doc,
        ["Factor", "Typical impact"],
        [
            [
                "Wiki coverage and accuracy",
                "Highest. Missing, outdated, or unclear pages mean wrong or vague answers—even with a good UI.",
            ],
            [
                "Page metadata (titles, descriptions, tags, links)",
                "High. Helps Yzma find the right pages when users ask in everyday language.",
            ],
            [
                "Retrieval (search + page selection)",
                "Medium. Can miss or rank the wrong page among similar tables; fallbacks exist but are not perfect.",
            ],
            [
                "AI summarization",
                "Medium. Can misread or over-simplify even good source pages—citations are there so you can verify.",
            ],
            [
                "Platform (model availability, limits)",
                "Lower for “is this true?” but affects whether chat works at all; very long pages may be truncated.",
            ],
            [
                "Optional live warehouse queries",
                "Separate path. When enabled, adds approval and access constraints; often disabled in some environments.",
            ],
        ],
    )
    add_para(
        doc,
        "Bottom line: Yzma is reliable when the team treats the wiki as a product—kept current, linked, "
        "and reviewed—not when chat is tuned in isolation.",
    )

    add_heading(doc, "The four main areas of the site", level=2)
    add_table(
        doc,
        ["Area (in the app)", "What you do there", "When to use it"],
        [
            [
                "Wiki pages (Browse)",
                "Read documentation organized by topic—metrics, domains, data products, concepts.",
                "You know what you want to learn, or you prefer reading at your own pace.",
            ],
            [
                "Ask the Wiki (Chat)",
                "Type a question; the assistant searches the wiki and replies with citations to specific pages.",
                "You have a question and want a direct answer with sources.",
            ],
            [
                "Pipeline graphs (Graph)",
                "See visual flow diagrams of how data moves through pipelines; click a step to open its wiki page.",
                "You want the “big picture” of how datasets connect.",
            ],
            [
                "Guides (Forms)",
                "Answer a short series of questions (decision trees) to get a recommendation—for example, which table to use.",
                "You want guided help without reading everything upfront.",
            ],
        ],
    )

    add_heading(doc, "General vs. Technical mode", level=2)
    add_para(
        doc,
        "Yzma has two experience levels, usually chosen from the URL or settings:",
    )
    add_bullets(
        doc,
        [
            "General — Simpler navigation aimed at business users and newcomers. You can browse "
            "by audience (“what are you trying to do?”), by product, by business domain, a metrics "
            "glossary, or getting-started hubs. Wiki pages show a short Summary section when one "
            "exists, instead of the full technical article.",
            "Technical — Full wiki structure in the sidebar, aimed at engineers and power users "
            "who want complete detail.",
        ],
    )
    add_para(
        doc,
        "Example link for the chat in General mode: "
        "http://yzma.dssinfra.com/?view=chat&tier=general",
        italic=True,
    )

    add_heading(doc, "How “Ask the Wiki” works (what you see)", level=2)
    add_bullets(
        doc,
        [
            "You ask a question in plain English (for example, “What is the difference between ALID and AVAIL?”).",
            "The status line updates while Yzma works—typically: searching the wiki, choosing pages, "
            "reading those pages, then generating an answer.",
            "You see links to the wiki pages used (references / citations) so you can open and verify them.",
            "The answer streams in word-by-word, like many chat assistants.",
            "After the answer, Yzma may show a confidence level (high, medium, or low).",
            "Chats are stored in your browser; you can start a new chat or share a read-only link on the company network.",
        ],
    )
    add_para(
        doc,
        "Important: The assistant only uses what it finds in the wiki (plus optional approved live queries—see below). "
        "If documentation is missing or outdated, answers may be incomplete—always check the cited pages.",
    )

    add_heading(doc, "What happens behind the scenes (wiki retrieval)", level=2)
    add_para(
        doc,
        "Yzma does not send the entire wiki to the AI on every question. It uses a standard pattern called "
        "Retrieval-Augmented Generation (RAG): first find the right wiki pages, then ask the AI to answer "
        "using only that material. That keeps answers tied to CDS documentation you can verify, instead of "
        "generic model guesses.",
    )
    add_para(
        doc,
        "The team knowledge base (markdown pages under wiki/) is the only source for this chat flow—personal "
        "or private notebooks are not mixed in.",
    )
    add_table(
        doc,
        ["Stage", "What Yzma is doing", "Why it matters to you"],
        [
            [
                "1. Search (~12 candidates)",
                "Scan wiki titles, tags, and page text for words related to your question; rank matches. "
                "A follow-up pass can also pull in structurally related pages (linked neighbors, shared sources) "
                "even when they do not share the same keywords.",
                "Surfaces the right neighborhood of docs, not just exact keyword hits.",
            ],
            [
                "2. Select (3–5 pages)",
                "A separate AI step reads short summaries of those candidates and picks the few pages most "
                "likely to answer you. If that step fails, Yzma falls back to the top keyword matches.",
                "Focuses the final answer on a small, relevant set of sources.",
            ],
            [
                "3. Read (with size limits)",
                "Load the full text of the chosen pages, up to sensible size caps so the model is not overloaded.",
                "Answers reflect real page content, trimmed when pages are very long.",
            ],
            [
                "4. Answer (streaming)",
                "The AI writes a reply using those pages plus scope text (purpose and wiki index). It cites sources "
                "with numbered references like [1], [2], shown as clickable links in the app.",
                "You get readable prose with a clear audit trail to wiki pages.",
            ],
        ],
    )
    add_para(
        doc,
        "If nothing useful is found, Yzma can fall back to a general overview page or tell you that no wiki "
        "pages matched—rather than inventing an answer.",
    )
    add_para(
        doc,
        "Your General vs. Technical setting also affects browsing: General mode may show Summary sections on "
        "wiki pages where those exist; Technical mode shows full articles.",
    )
    add_para(
        doc,
        "Technical reference (full pipeline, scoring, and configuration): "
        "https://github.twdcgrid.net/content-data-solutions/yzma/blob/main/docs/architecture-chat-wiki-retrieval.md",
        italic=True,
    )

    add_heading(doc, "Confidence scores (in plain language)", level=2)
    add_para(
        doc,
        "After an answer, Yzma may label confidence high, medium, or low. That is a heuristic—not a "
        "legal or audit guarantee. It roughly reflects:",
    )
    add_bullets(
        doc,
        [
            "How well the retrieved pages matched your question.",
            "Whether the answer cited the pages it relied on.",
            "Whether the wording is hedged (“might”, “not found”) vs. clearly grounded in the docs.",
            "Whether the answer included numbers from a live warehouse query (when that feature is enabled).",
        ],
    )

    add_heading(doc, "Follow-up questions and conversation memory", level=2)
    add_para(
        doc,
        "For follow-ups, Yzma includes the last several turns of the conversation (up to six messages) "
        "along with your new question so the assistant can handle context (“What about Hulu?” after asking "
        "about Disney+). History is kept in your browser unless you start a new chat.",
    )

    add_heading(doc, "Reading wiki pages: trust and freshness", level=2)
    add_para(
        doc,
        "Many pages include labels to help you judge reliability:",
    )
    add_bullets(
        doc,
        [
            "Human-reviewed — A person has verified the content.",
            "Machine-confirmed — Content was checked by an automated process, not necessarily a human.",
            "Unverified — No verification marker; treat with extra caution.",
            "Deprecated — The page describes something that should no longer be used.",
            "Stale — The page may be past its “stale after” date and due for refresh.",
        ],
    )

    add_heading(doc, "Guides and pipeline graphs", level=2)
    add_para(
        doc,
        "Guides walk you through choices with “if this, then that” logic—for example, "
        "“Which table should I use?” after you answer a few questions.",
    )
    add_para(
        doc,
        "Pipeline graphs show stages of data processing (such as aggregation pipelines). "
        "Each box in the diagram links to the wiki page for that dataset or step.",
    )

    add_heading(doc, "Optional live data queries (Snowflake)", level=2)
    add_para(
        doc,
        "In some configurations, the assistant can propose running a SQL query against Snowflake to "
        "answer with live counts or metrics—not just wiki text. When that capability is enabled:",
    )
    add_bullets(
        doc,
        [
            "Yzma pauses and shows you the SQL and why it wants to run it.",
            "You explicitly approve or decline. If you decline, it continues without running the query.",
            "This protects against expensive or unintended warehouse usage.",
        ],
    )
    add_para(
        doc,
        "In other environments, live Snowflake access is turned off entirely. In that case the "
        "assistant explains that it cannot run queries and may still suggest SQL you can run yourself "
        "based on wiki documentation.",
    )

    add_heading(doc, "Feedback and improvements", level=2)
    add_para(
        doc,
        "You can send feedback (bugs, feature ideas, documentation gaps) from the app. "
        "Feedback is routed to the team via GitHub issues so it can be tracked and prioritized. "
        "Yzma is marked as a beta build, so expect ongoing changes.",
    )

    add_heading(doc, "Quick start for a non-technical reader", level=2)
    add_bullets(
        doc,
        [
            "Open Yzma on the company network (VPN or office network as required).",
            "Start in General mode if you are new.",
            "Try Browse → “Getting started” or “By audience” to explore hubs.",
            "When you have a specific question, open Ask the Wiki, start a new chat, and ask.",
            "Click the citation links in any answer to read the source pages.",
            "Use Guides when you want a recommendation without searching manually.",
        ],
    )

    add_heading(doc, "What Yzma is not", level=2)
    add_bullets(
        doc,
        [
            "Not a replacement for official governance tools or access requests—it explains and documents data.",
            "Not a guarantee of live production numbers unless an answer explicitly ran an approved query.",
            "Not public on the internet; it is an internal DSS infrastructure service.",
        ],
    )

    doc.add_paragraph()
    p = doc.add_paragraph()
    run = p.add_run(
        "Note: This overview summarizes docs/architecture-chat-wiki-retrieval.md in the yzma repository "
        "(Team KB Chat Retrieval Pipeline) and was cross-checked against the live app. Engineers should "
        "use that GitHub document for implementation detail; this file is an on-ramp for non-technical readers."
    )
    set_run_font(run, italic=True, size=10, color=MUTED)

    doc.save(OUT)
    return OUT


if __name__ == "__main__":
    path = build()
    print(path)
