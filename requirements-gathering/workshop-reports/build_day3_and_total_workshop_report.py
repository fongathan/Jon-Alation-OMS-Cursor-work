#!/usr/bin/env python3
"""Build Day 3 + 3-Day Total Summary workshop report (.docx)."""
from __future__ import annotations

from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from docx.shared import Inches, Pt

OUT = Path(__file__).resolve().parent / "Navigator-Portal-Workshop-Day3-and-Total-Summary.docx"


def h1(doc: Document, text: str) -> None:
    doc.add_heading(text, level=1)


def h2(doc: Document, text: str) -> None:
    doc.add_heading(text, level=2)


def h3(doc: Document, text: str) -> None:
    doc.add_heading(text, level=3)


def para(doc: Document, text: str) -> None:
    doc.add_paragraph(text)


def bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        doc.add_paragraph(item, style="List Bullet")


def build() -> Path:
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    title = doc.add_paragraph()
    title.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    tr = title.add_run("Navigator Portal Workshop — Day 3 Report & 3-Day Total Summary")
    tr.bold = True
    tr.font.size = Pt(16)

    sub = doc.add_paragraph()
    sub.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    sr = sub.add_run("Data Governance × Data Foundations Platform alignment · Seattle · July 21–23, 2026")
    sr.bold = True
    sr.font.size = Pt(12)

    meta = doc.add_paragraph()
    meta.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    meta.add_run(f"Draft report · {date.today().strftime('%B %d, %Y')} · Prepared for workshop distribution")
    doc.add_paragraph()

    h1(doc, "Document purpose")
    para(
        doc,
        "This report summarizes Day 3 (Thursday, July 23, 2026) of the Navigator Portal workshop and "
        "synthesizes outcomes across all three days. Format follows the Day 2 workshop report "
        "(Wilson Chaves) posted to the workshop Teams chat.",
    )
    para(doc, "Sources:")
    bullets(
        doc,
        [
            "Teams meeting transcript — Navigator Portal Workshop, July 23, 2026 (recording segments; primary transcript reviewed via Teams MCP).",
            "Workshop Teams chat — Day 3 prep topics (Sarah Smurthwaite), facilitator wrap message, whiteboard photos.",
            "Day 1 summary — Wilson Chaves, workshop chat, July 21, 2026.",
            "Day 2 summary — Wilson Chaves, workshop chat and Word doc link, July 22, 2026.",
            "Facilitator notes — Brenda Villasenor SharePoint folder (linked from Facilitator bot at close of Day 3).",
        ],
    )
    para(
        doc,
        "Note: Day 3 was recorded in multiple segments. The transcript reviewed for this report "
        "captures the metrics/semantic-layer deep dive, governance policy architecture whiteboard "
        "(Roopa Prabhu), and the closing reflection circle. Planned privacy/RIM/DSR agenda items "
        "may appear in earlier segments or facilitator notes not fully extracted here.",
    )

    # --- DAY 3 ---
    h1(doc, "Day 3 — Executive summary")
    para(
        doc,
        "Day 3 shifted from persona and workflow framing (Days 1–2) toward concrete architecture for "
        "enterprise metrics, governance metadata, and policy enforcement. The group aligned on a "
        "Git-backed, OSI-standard metrics artifact model with DNA-led definition and Data Governance–led "
        "change control, OMS-backed drift detection, and phased enforcement for certified publishing. "
        "Roopa Prabhu walked through the Governance Metadata Repository (Postgres SoR for tags and "
        "policies), governance application services, and Iceberg-based metadata publication for platform "
        "consumption—including an end-to-end Access Manager / composable workspace policy-profile example. "
        "Engineering raised valid concerns about service boundaries, version sync between Postgres and "
        "Iceberg readers, and decision-service ownership; follow-up technical design sessions were agreed "
        "with Access Manager and platform engineering (Anne, Eric, Andrew). The day closed with strong "
        "cross-team sentiment that governance–platform collaboration had materially improved and that "
        "backend architecture plans mesh better than expected.",
    )

    h2(doc, "Day 3 session structure (as captured)")
    bullets(
        doc,
        [
            "Part 1 — Metrics creation, governance, and serving layers; OSI / GitHub YAML; unified semantic layer (leadership + product).",
            "Part 2 — Policy & governance architecture whiteboard; metadata repository, publication layer, access decision flow (Roopa + engineering).",
            "Part 3 — Workshop close; learnings, thank-yous, and commitment to follow-on working sessions.",
        ],
    )

    h2(doc, "Main decisions / emerging alignment (Day 3)")
    h3(doc, "1. Three-layer metrics model (create → govern → serve)")
    bullets(
        doc,
        [
            "Creation layer: metric authors (often DNA / domain experts) record definitions in a centralized metrics knowledge base—not ad hoc in dashboards alone.",
            "Governance layer: review, conflict detection, steward approval, and enterprise naming conventions.",
            "Serving layer: discoverability and consumption via OMS / Navigator / semantic views; aligns with Discover–Govern–Monitor–Act–Consume from Day 2.",
            "Workshop gap called out: creators are not always in the room; product must explicitly support creation workflows, not only governance of existing metrics.",
        ],
    )

    h3(doc, "2. OSI + Git as the enterprise metric artifact")
    bullets(
        doc,
        [
            "YAML metric artifacts stored in Git (Open Semantic Interchange / OSI standard), PR-based change control, validation against OSI structure on merge.",
            "DNA owns metric definition negotiation with business leadership; Data Governance owns repository workflow, change control, and enforcement hooks.",
            "Approved metrics propagate to semantic layer implementations (Snowflake semantic views, Databricks metric views) for platform portability.",
            "OMS role evolves: from reverse-engineering scattered metrics to matching implementations against Git SoR and surfacing drift/conflicts.",
            "Rough scale cited: ~240 governed metrics today with backfill of P0 existing definitions; rate of change expected to stabilize after initial backfill.",
        ],
    )

    h3(doc, "3. Enforcement philosophy (pragmatic, phased)")
    bullets(
        doc,
        [
            "Exploration and personal analysis remain unconstrained; gates apply when publishing/sharing for enterprise consumption (360 / marketplace / certified dashboards).",
            "Near-term focus: Looker/DLN dashboards and DNA-controlled surfaces (~60–80% of drift problem); longer tail includes Tableau, Pixie, and other consumers.",
            "Certified vs uncertified reporting: executives may dismiss uncertified numbers in forums; culture + tooling combined.",
            "Metrics Council / stewardship review for conflicts; gated PR merge for net-new enterprise metrics.",
        ],
    )

    h3(doc, "4. Governance Metadata Repository architecture (directional)")
    bullets(
        doc,
        [
            "Postgres-backed Governance Metadata Repository as SoR for classifications, tags, and policy definitions (authored with legal/governance, effective dates, versioning).",
            "Governance application services expose policies to enforcement engines (API) and publish projections to platforms.",
            "Metadata published layer (Iceberg) for zero-ETL consumption inside Snowflake/Databricks catalogs—joint POC direction with Eric/Andrew/Sumanta.",
            "Platform asset tagging (table/CW tags) remains on platform; repository holds policy definitions and mappings—not necessarily every live tag instance.",
            "Access pattern: workspace policy profile + dataset/column tags → rules consumed by Access Manager decision logic (PII allow/mask/deny example validated conceptually).",
        ],
    )

    h3(doc, "5. Cross-org collaboration tone")
    bullets(
        doc,
        [
            "Ryan Malec: backend architecture across Gov and DFP “meshes together better than expected”; unified backend path feels achievable.",
            "Multiple participants noted improved Gov ↔ platform ↔ architecture relationship vs prior years.",
            "Pedro called out as leading Business Data Catalog on OMS; OMS team focused on technical metadata quality and lineage accuracy.",
        ],
    )

    h2(doc, "Key discussion topics (Day 3)")
    h3(doc, "Part 1 — Metrics & semantic layer")
    bullets(
        doc,
        [
            "Industry pattern: “metric server” / single enterprise formula (Uber rides analogy); Disney attempting second major push after prior fragmentation.",
            "Nuanced metrics (e.g., watch hours): multiple valid flavors with explicit naming—not one overloaded definition.",
            "Context in OSI extensions: business dimensions, exclusions, column-level lineage for formulas; machines and agents need richer context than formula alone.",
            "Integration with Kyle / semantic layer POC and Metrics Console; alignment session planned with DNA leadership.",
            "Conflict workflow: detect duplicate/near-duplicate metrics at publish time; steward review and Metrics Council escalation.",
        ],
    )

    h3(doc, "Part 2 — Policy architecture & Access Manager integration")
    bullets(
        doc,
        [
            "Whiteboard: solid boxes = systems/boundaries; dotted groupings = logical functions (not separate monoliths).",
            "Flow: tag/policy workflows → update repository → policies ready for enforcement → governance services + Iceberg publication → platform engines (access, retention, classification).",
            "Chennah (Access Manager): decision service needs both policies and current tags; prefers clear single SoR and one owning service family per database to avoid drift between Iceberg and Postgres readers.",
            "Agreed next step: share written architecture doc; walk through concrete Access Manager request example; include Anne, Eric, Andrew in design follow-up.",
            "Versioning on policy change: impact analysis for affected datasets—not silent propagation.",
        ],
    )

    h3(doc, "Part 3 — Close & reflections")
    bullets(
        doc,
        [
            "Wireframes / governance portal vision cited as especially helpful for engineers (Tim Liu and others).",
            "Agent-first design repeated as success criterion: systems must answer how agents interact, not only human UI.",
            "Niru/Jon/Roopa/Richard/Wilson thanked for facilitation, documentation, and “polite aggression” in working sessions.",
            "Commitment to additional follow-up sessions rather than treating workshop as final spec.",
        ],
    )

    h2(doc, "Day 3 — Open questions")
    bullets(
        doc,
        [
            "Portal count & URL strategy: single vs multiple portals; landing experience (“Google vs ESPN”) — Sarah’s Day 3 prep list; still open.",
            "Where exactly does the policy decision service live, and which team owns reads/writes to the governance Postgres SoR?",
            "Iceberg vs API: which consumers must use which path to avoid version skew?",
            "Full privacy/RIM/DSR scope from original agenda: confirm coverage in earlier Day 3 segments or schedule dedicated follow-up with Norman/privacy leads.",
            "Sisu / external policy engine vs custom governance services — raised on Day 2; still unresolved.",
            "North Star one-liner for Navigator product definition — facilitator follow-up from late Day 2.",
        ],
    )

    h2(doc, "Day 3 — Action items")
    bullets(
        doc,
        [
            "Share governance metadata / policy architecture design doc with Access Manager and platform engineering (Roopa → Chennah, Anne, Eric, Andrew).",
            "Document end-to-end Access Manager example (registry CW request → policy profile → grant/deny/mask) as sequence diagram.",
            "Metrics: continue OSI/Git POC; schedule Kyle/DNA alignment; define backfill plan for P0 metrics.",
            "OMS: document drift-detection model once Git SoR is authoritative cutoff date.",
            "Wilson / TPM: capture whiteboard photos and AI summaries in workshop artifact folder (referenced in room).",
            "Product: portal strategy decision brief for leadership (single vs multiple; persona landing pages).",
            "Facilitator notes: incorporate Brenda Villasenor SharePoint folder content into master action tracker.",
        ],
    )

    h2(doc, "Day 3 — Leadership decisions needed")
    bullets(
        doc,
        [
            "Ratify metrics enforcement phased approach (what requires certification vs allowed exploratory use).",
            "Confirm Governance Metadata Repository as enterprise SoR for policy/tag definitions (vs distributing across platform catalogs).",
            "Portal strategy decision (carried from Days 1–2).",
            "Executive sponsorship for cross-team working sessions post-workshop (Gov + DFP + OMS + Access Manager).",
        ],
    )

    doc.add_page_break()

    # --- 3-DAY TOTAL ---
    h1(doc, "3-Day workshop — total summary")
    para(
        doc,
        "The Navigator Portal workshop brought Data Governance, Data Foundations Platform, OMS, "
        "Access Manager, architecture, and domain partners together in Seattle for 2.5 days to align "
        "on post-Alation catalog/governance experience, shared backend architecture, and operating "
        "model across teams.",
    )

    h2(doc, "Arc across the three days")
    bullets(
        doc,
        [
            "Day 1 (Jul 21): Future-state framing — OMS as aggregation backbone, thin UI layers, single source of truth, portal strategy debate, metadata vs AI context (626), policy lifecycle vs discovery.",
            "Day 2 (Jul 22): Personas (~85% confidence), Discover/Govern/Monitor/Act/Consume workflows, classification/policy/access depth, governance portal expectations, policy translation pipeline.",
            "Day 3 (Jul 23): Metrics OSI/Git model, governance metadata & policy enforcement architecture, Access Manager integration concerns, collaborative close and follow-ups.",
        ],
    )

    h2(doc, "Cross-workshop themes (stable)")
    bullets(
        doc,
        [
            "Dual experience, shared spine: persona-specific portals/UI acceptable if backend and metadata SoR are unified.",
            "Governance is operational, not only reporting: enforcement, audit evidence, workspace-aware access, policy lifecycle.",
            "Human + agent consumers: metadata and context must be machine-readable and validated.",
            "Avoid duplicating OMS: rebuild Alation capabilities deliberately; don’t fork lineage/discovery.",
            "Relationship repair: Gov ↔ platform ↔ architecture collaboration explicitly improved during workshop.",
        ],
    )

    h2(doc, "Cross-workshop decisions & alignment (best current view)")
    bullets(
        doc,
        [
            "OMS remains central technical metadata / lineage / discovery backbone.",
            "Governance policy and tag definitions → Governance Metadata Repository (Postgres) with services + publication layer.",
            "Enterprise metrics → Git + OSI YAML, DNA defines, Gov operationalizes change control and enforcement hooks.",
            "Persona framework and DGMAC workflow categories are default IA for portal design pending leadership sign-off.",
            "Classification, policy, and access are tightly coupled; Access Manager is critical integration point.",
        ],
    )

    h2(doc, "Cross-workshop open questions (prioritized for leadership)")
    bullets(
        doc,
        [
            "1. Portal strategy: one portal vs multiple vs shared shell with persona entry points.",
            "2. Policy platform direction: extend existing stack, Sisu, or net-new governance services.",
            "3. RACI: Gov vs DFP vs OMS vs Access Manager for build/support of each capability.",
            "4. Metrics enforcement timeline and certification rules for enterprise publishing.",
            "5. AI/context strategy (626): validation, canonical promotion, storage—post-workshop track.",
        ],
    )

    h2(doc, "Recommended immediate next steps (program)")
    bullets(
        doc,
        [
            "Publish alignment memo + RACI draft (workshop commitment; target late July 2026).",
            "Technical deep-dives: (a) policy repository ↔ Access Manager, (b) metrics Git/OSI POC, (c) portal IA decision.",
            "Leadership readout: portal strategy, policy SoR, metrics enforcement, and Gov/DFP operating model.",
            "Merge facilitator notes, Wilson Day 2 doc, and this Day 3/total summary into single workshop artifact repository.",
        ],
    )

    h2(doc, "References & links")
    bullets(
        doc,
        [
            "Workshop Teams chat: meeting thread 19:meeting_YzQ2MjIyOWUtNTQ1Yy00OGJiLWFiOTUtOWFhZTY1M2NiYjJi@thread.v2",
            "Day 2 report (Teams): Wilson Chaves message 1784759768700; Word doc on Wilson SharePoint (661115cf-fbed-421c-9f0e-e1ed03fdb352).",
            "Day 1 summary (Teams): Wilson Chaves message 1784676330255.",
            "Day 3 facilitator notes folder: Brenda Villasenor SharePoint (linked from Facilitator message 1784837423579).",
            "Day 3 recording/transcript: Navigator Portal Workshop-20260723 (Niru Sharma Recordings folder).",
            "OSI example (chat): github.com/apache/ossie (flights.yaml sample shared by Brenda).",
        ],
    )

    h2(doc, "Appendix — Whiteboard captures (Day 3)")
    para(
        doc,
        "In-room whiteboards (photos shared to workshop chat) illustrated:",
    )
    bullets(
        doc,
        [
            "Navigator UI areas: site config, user profile, discover, search; Producer vs Consumer capabilities (catalog, lineage, search vs self-service, access).",
            "Classification and Policies feeding automated table/workspace controls; DFP and OMS called out as adjacent platform layers.",
            "Entry URLs discussed: datahq / data portal alignment; React + shared design system (“Shadow”) for thin UI.",
            "Governance hub: auth, domain groupings, metadata services, GraphQL/API backends.",
        ],
    )

    doc.save(OUT)
    return OUT


if __name__ == "__main__":
    print(build())
