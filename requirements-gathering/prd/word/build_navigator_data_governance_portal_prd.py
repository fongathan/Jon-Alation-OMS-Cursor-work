#!/usr/bin/env python3
"""Build Navigator Data Governance Portal PRD (.docx)."""
from __future__ import annotations

from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from docx.shared import Pt

OUT = Path(__file__).resolve().parent / "PRD-Navigator-Data-Governance-Portal.docx"


def add_heading(doc: Document, text: str, level: int = 1) -> None:
    doc.add_heading(text, level=level)


def add_para(doc: Document, text: str) -> None:
    doc.add_paragraph(text)


def add_bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        doc.add_paragraph(item, style="List Bullet")


def add_table(doc: Document, headers: list[str], rows: list[tuple[str, ...]]) -> None:
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    for i, h in enumerate(headers):
        table.rows[0].cells[i].text = h
    for r_idx, row in enumerate(rows):
        for c_idx, val in enumerate(row):
            table.rows[r_idx + 1].cells[c_idx].text = val
    doc.add_paragraph()


def build() -> Path:
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    title = doc.add_paragraph()
    title.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    tr = title.add_run("Product Requirements Document")
    tr.bold = True
    tr.font.size = Pt(16)

    sub = doc.add_paragraph()
    sub.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    sr = sub.add_run("Navigator — Data Governance Portal")
    sr.bold = True
    sr.font.size = Pt(14)

    meta = doc.add_paragraph()
    meta.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    meta.add_run(f"Version 0.1 · Draft · {date.today().isoformat()}")

    doc.add_paragraph()

    # --- 1. Executive summary ---
    add_heading(doc, "1. Executive summary", 1)
    add_para(
        doc,
        "The Navigator Data Governance Portal is the privacy and data governance command center "
        "within the broader Navigator product family. It answers three questions for program operators "
        "and leadership: Are we in policy? What needs a human this week? Can we prove it to auditors?",
    )
    add_para(
        doc,
        "The Portal is not a replacement for specialist systems—consent platforms, DSR orchestration, "
        "tracker remediation, access manager, GIS, Kronos inventory, or the Business Data Catalog. "
        "Design principle: unified experience, federated execution. One front door for posture, "
        "exceptions, and evidence; downstream tools remain authoritative for deep workflow.",
    )
    add_para(
        doc,
        "The Portal sits in a shared Navigator shell alongside Business Data Catalog (discover & steward) "
        "and deep links to Foundations OMS UI (technical metadata for engineers). It reads and updates "
        "governance metadata through Gov Application Services and the Governance Metadata Repository "
        "(PULS Postgres), while catalog assets remain in OMS.",
    )
    add_para(
        doc,
        "Delivery is phased: Business Data Catalog MVP on OMS APIs addresses the Oct 2026 Alation exit "
        "pressure first. The full Governance Portal hub—live KPI feeds, work queues, evidence export—"
        "ships in Phase 2, with metrics automation and resolved OMS↔governance read paths in Phase 3.",
    )

    # --- 2. Problem statement ---
    add_heading(doc, "2. Problem statement", 1)
    add_para(
        doc,
        "DEEPT runs privacy, platform, security, and regulatory governance across many specialist "
        "tools. Each workstream—consent, data-subject requests, tracker findings, access reviews, "
        "GIS compliance, data-use approvals, records management—maintains its own dashboard and "
        "operational queue. Program leads and executives cannot see overall posture or prioritize "
        "across dimensions without assembling screenshots from five URLs.",
    )
    add_para(
        doc,
        "Posture charts alone do not answer what needs a human this week. Privacy officers need "
        "triage and routing, not another static report. Stewards need classification review workflows "
        "connected to catalog assets, not a governance dashboard disconnected from the tables they "
        "curate. Auditors need exportable evidence packages tied to policy definitions and catalog "
        "objects—not ad hoc spreadsheet assembly before every review cycle.",
    )
    add_para(
        doc,
        "As Alation exits, governance and privacy teams need a stable home linked to OMS-backed "
        "catalog assets. Without a dedicated Portal, governance features risk being buried in "
        "Foundations' engineer-first OMS UI—or duplicated in ways that fork truth between OMS and "
        "the Governance Metadata Repository.",
    )
    add_para(doc, "The Portal must solve:")
    add_bullets(
        doc,
        [
            "Visibility: one Governance Overview with PCI-style indices and dimension scores.",
            "Triage: work queues that surface SLA breaches and route to the right downstream tool.",
            "Coherence: defined vs observed—policy intent alongside what OMS sees on catalog assets.",
            "Evidence: auditor-oriented export without rebuilding every compliance backend.",
            "Separation: govern & prove in Portal; discover & steward in Business Data Catalog.",
        ],
    )

    # --- 3. Product vision & positioning ---
    add_heading(doc, "3. Product vision & positioning", 1)
    add_para(
        doc,
        "Vision: The Data Governance Portal is the operating layer for privacy, platform, security, "
        "and regulatory posture—not a monolith that absorbs Kronos, PULS, or consent platforms, but "
        "the shell that makes the whole program legible to non-specialists.",
    )
    add_para(
        doc,
        "Positioning within Navigator:",
    )
    add_table(
        doc,
        ["Product surface", "Primary job", "Primary users"],
        [
            (
                "Data Governance Portal",
                "Govern & prove — posture, exceptions, evidence",
                "Privacy/compliance, program leads, audit, stewards (workflows)",
            ),
            (
                "Business Data Catalog",
                "Discover & steward — search, glossary, asset detail",
                "Analysts, stewards, business consumers",
            ),
            (
                "Foundations OMS UI",
                "Technical metadata — lineage, SQL, ops context",
                "Data engineers, platform operators",
            ),
        ],
    )
    add_para(
        doc,
        "Organizing principle: Portal information architecture follows four functional dimensions—"
        "Privacy, Platform, Security, Regulatory—mapped to how DEEPT runs governance programs, "
        "not internal org-chart names.",
    )

    # --- 4. Scope ---
    add_heading(doc, "4. Scope", 1)
    add_heading(doc, "4.1 In scope", 2)
    add_bullets(
        doc,
        [
            "Governance Overview landing with Privacy Compliance Index (PCI) and Platform / Security / Regulatory section indices.",
            "Four dimension hubs (Privacy, Platform, Security, Regulatory) with Overview + Work queue tabs per function.",
            "Work queue model: unified queue filtered by dimension; handoffs to downstream specialist tools.",
            "Platform dimension: steward coverage, classifications management UI, policies, deep links to Business Data Catalog.",
            "Privacy dimension: consent posture, DSR volume/SLA, tracker remediation KPIs and queue routing.",
            "Security dimension: access reviews, GIS compliance, data-use approval visibility and queues.",
            "Regulatory dimension: laws & regulations, incidents, audits; Kronos inventory deep link.",
            "Defined vs observed views linking policy/classification definitions to OMS catalog state.",
            "Evidence export and auditor-oriented read-only views from Portal header.",
            "Shared Navigator shell: top bar, global search, breadcrumbs, collapsible left rail (v1 two-tab or v2 governance-first IA).",
            "Gov App backend integration for classifications, policies, and metrics (phased).",
            "Role-based landing pages (privacy/compliance → Overview; stewards → Platform workflows; executives → Overview + export).",
        ],
    )
    add_heading(doc, "4.2 Out of scope (unless added by change request)", 2)
    add_bullets(
        doc,
        [
            "Replacing specialist backends: consent platform, DSR orchestration, MAUI, tracker tools, Access Manager, GIS consoles.",
            "Duplicating Business Data Catalog discovery flows (search, browse, glossary edit) inside Portal screens.",
            "Duplicating Kronos business inventory UI—Regulatory dimension links out to Kronos.",
            "Foundations OMS engineer UI (lineage graph, SQL tabs)—deep link only.",
            "Metadata ingestion, connector scraping, or graph store ownership (626/MCI → OMS supply chain).",
            "MCP / agentic catalog access (owned by MCI + Data 66 patterns, not Portal).",
            "Full ITSM or enterprise problem management—incident integration TBD with platform partners.",
        ],
    )

    # --- 5. Personas ---
    add_heading(doc, "5. User personas", 1)
    add_para(
        doc,
        "Primary Portal users operate governance programs—they need triage, attestation, and "
        "cross-pillar visibility. Secondary users consume governed metadata through Business Data "
        "Catalog. Engineers use Foundations OMS UI, not the Portal.",
    )
    add_table(
        doc,
        ["Persona", "Portal role", "Primary surfaces"],
        [
            (
                "Privacy & compliance officers",
                "Primary",
                "DSR queues, consent posture, tracker remediation, compliance reporting",
            ),
            (
                "Governance program leads & executives",
                "Primary",
                "Governance Overview, PCI + section indices, metrics governance, evidence export",
            ),
            (
                "Data stewards / DPMs",
                "Primary (workflows)",
                "Classifications Mgmt UI, tag & policy workflows, classification review",
            ),
            (
                "Access managers & GIS / security",
                "Primary",
                "Access manager handoffs, data-use approvals, Security dimension queues",
            ),
            (
                "Partner & API managers",
                "Primary",
                "Partner API management handoffs, partner data transfer framework",
            ),
            (
                "Records / inventory owners",
                "Primary",
                "Business inventory workflows; Kronos link (Regulatory dimension)",
            ),
            (
                "Internal audit & legal",
                "Primary (read/export)",
                "Evidence aggregation, auditor view, regulatory incident posture",
            ),
            (
                "Business & analytics consumers",
                "Indirect",
                "Business Data Catalog for discovery—not Portal MVP",
            ),
            (
                "Data engineers / platform operators",
                "Not Portal",
                "Foundations OMS UI; deep link from catalog asset",
            ),
        ],
    )

    # --- 6. Information architecture ---
    add_heading(doc, "6. Information architecture", 1)
    add_heading(doc, "6.1 Four dimensions", 2)
    add_table(
        doc,
        ["Dimension", "Scope", "Example KPI / queue items"],
        [
            (
                "Privacy",
                "Consent · guest data requests · tracker remediation",
                "PCI & exception volume; DSR SLA breaches; tracker findings",
            ),
            (
                "Platform",
                "Stewards · classifications · policies",
                "Stewardship coverage; classification sign-offs; policy mappings",
            ),
            (
                "Security",
                "Access reviews · GIS · data use approvals",
                "Access & GIS compliance; data-use intake queue",
            ),
            (
                "Regulatory",
                "Laws · incidents · audits · inventory",
                "Evidence readiness; regulatory incidents; Kronos inventory link",
            ),
        ],
    )
    add_heading(doc, "6.2 Navigation IA — v1 vs v2", 2)
    add_table(
        doc,
        ["", "v1 — Two top-level tabs", "v2 — Governance-first (recommended for Portal)"],
        [
            (
                "Left nav",
                "Business Data Catalog | Data Governance Portal",
                "Governance Overview + Privacy / Platform / Security / Regulatory sections",
            ),
            (
                "Catalog entry",
                "Top-level tab",
                "Platform → Business Data Catalog (nested shell)",
            ),
            (
                "Best for",
                "Clear mode switch; Alation-exit MVP",
                "Privacy/compliance default landing; exec posture first",
            ),
            (
                "Risk",
                "Users may live only in catalog tab",
                "Catalog one click deeper—mitigated by search + Platform CTA",
            ),
        ],
    )
    add_para(
        doc,
        "Recommendation: ship v2 governance-first IA for the Portal Phase 2 hub; keep v1 two-tab "
        "pattern as the simpler MVP path for catalog-only users until Portal launches. Both use "
        "the same shell and APIs.",
    )
    add_heading(doc, "6.3 Shared shell elements", 2)
    add_bullets(
        doc,
        [
            "Top bar: product breadcrumb, global search, role context, Evidence export / Auditor view actions.",
            "Left rail: mode switch (v1) or governance dimensions (v2).",
            "Main panel: Governance Overview, dimension hub, work queue, or specialist handoff.",
            "Deep links: Portal exception → Business Data Catalog asset detail → Open in OMS for technical context.",
        ],
    )

    # --- 7. Functional requirements ---
    add_heading(doc, "7. Functional requirements", 1)

    add_heading(doc, "7.1 Governance Overview", 2)
    add_bullets(
        doc,
        [
            "Display Privacy Compliance Index (PCI) hero metric with trend and status badge.",
            "Display Platform, Security, and Regulatory section indices with drill-down to dimension hubs.",
            "Surface open exceptions by severity and dimension on Overview.",
            "Support time-range selection for trend charts where live feeds exist.",
            "Default landing for privacy/compliance and executive personas.",
        ],
    )

    add_heading(doc, "7.2 Dimension hubs & work queues", 2)
    add_bullets(
        doc,
        [
            "Each dimension exposes Overview tab (KPI cards, trend charts) and Work queue tab (actionable items).",
            "Queue items include: title, severity, SLA state, assignee/owner, linked asset or policy reference, handoff action.",
            "Unified queue model supports cross-dimension filter; dimension-specific views are presets.",
            "Handoff opens downstream tool (MAUI, consent platform, Access Manager, Navigator Requests, Kronos) with context preserved where APIs allow.",
            "Closed/resolved items retained for audit trail with timestamps.",
        ],
    )

    add_heading(doc, "7.3 Privacy dimension", 2)
    add_bullets(
        doc,
        [
            "Consent management posture KPIs and exception queue routing.",
            "DSR / guest data request volume, SLA compliance, and breach highlighting.",
            "Tracker remediation findings queue with severity and remediation status.",
            "Deep links to MAUI and privacy operational tools—not embedded replacement UIs in Phase 2.",
        ],
    )

    add_heading(doc, "7.4 Platform dimension", 2)
    add_bullets(
        doc,
        [
            "Stewardship coverage metrics and steward assignment gaps.",
            "Classifications Management UI: review, approve, reject AI-suggested classifications; link to catalog assets.",
            "Policies management: policy detail, control mappings, linkage to classifications and catalog assets.",
            "Business Data Catalog entry: Platform → Stewards / Catalog for discovery and metadata curation.",
            "Defined vs observed: side-by-side policy/classification definition and OMS-observed state on assets.",
        ],
    )

    add_heading(doc, "7.5 Security dimension", 2)
    add_bullets(
        doc,
        [
            "Access review status and overdue attestation queue items.",
            "GIS compliance posture and exception routing.",
            "Data-use approval intake queue with handoff to approval workflow systems.",
            "Visibility into access/classification alignment (Access Manager integration, phased).",
        ],
    )

    add_heading(doc, "7.6 Regulatory dimension", 2)
    add_bullets(
        doc,
        [
            "Laws & regulations register with applicability and control mapping summaries.",
            "Regulatory incident tracking and posture on Overview + queue.",
            "Audit readiness indicators and audit cycle queue items.",
            "Inventory deep link to Kronos (https://dcf.corp.dig.com/vista/kronos/)—not embedded inventory UI.",
        ],
    )

    add_heading(doc, "7.7 Evidence & auditor view", 2)
    add_bullets(
        doc,
        [
            "Evidence export action packages posture metrics, open exceptions, and policy mappings for a selected scope and time range.",
            "Auditor view: read-only Portal mode hiding mutating actions; optimized for attestation walkthroughs.",
            "Export formats: PDF and/or structured download (CSV/JSON) per audit program requirements.",
            "Phase 3: scheduled evidence packages and automated refresh tied to catalog asset IDs.",
        ],
    )

    add_heading(doc, "7.8 Search & global navigation", 2)
    add_bullets(
        doc,
        [
            "Global search across governance objects (policies, classifications, queue items) and catalog assets where entitled.",
            "Search results respect RBAC; route to Portal workflow or BDC asset detail appropriately.",
            "Breadcrumbs reflect dimension → function → detail hierarchy.",
        ],
    )

    # --- 8. User stories ---
    add_heading(doc, "8. User stories", 1)
    add_bullets(
        doc,
        [
            "As a privacy officer, I land on Governance Overview and see PCI and open Privacy queue items so I know what needs action this week.",
            "As a privacy officer, I open a DSR SLA breach from the queue and hand off to MAUI with case context preserved.",
            "As a data steward, I review pending classifications in Platform, approve or reject with rationale, and open the linked catalog asset in Business Data Catalog.",
            "As a governance program lead, I export an evidence package for an internal audit without assembling spreadsheets from five tools.",
            "As an executive, I view Platform / Security / Regulatory section indices on Overview and drill into exceptions by severity.",
            "As an access manager, I see overdue access attestations in Security queue and route to Access Manager for completion.",
            "As an auditor, I use read-only auditor view to walk posture metrics and exception history for a defined scope.",
            "As a records owner, I follow Regulatory → Inventory to Kronos for business inventory maintenance.",
        ],
    )

    # --- 9. UI / UX requirements ---
    add_heading(doc, "9. UI / UX requirements", 1)
    add_bullets(
        doc,
        [
            "Enterprise UX alignment: Sage-style patterns, Dave Hoffman / enterprise design review where applicable.",
            "Container-first layout: Suman guidance—validate dimension hub layout and KPI containers before final metric definitions.",
            "Accessibility: WCAG-oriented patterns for core flows (exact tier set by program).",
            "Responsive: desktop-first for program operators; read-only auditor view acceptable on tablet.",
            "Empty states: explain when live feeds are not yet connected (Phase 1 prototypes vs Phase 2 live data).",
            "Stale data indicators: last successful sync timestamp when KPIs are federated from downstream systems.",
            "Destructive actions (reject classification, close exception) require confirmation and audit logging.",
            "Prototype references: data-navigator-governance/index.html (UDGE demo); navigator-version-2/index.html (v2 shell).",
        ],
    )

    # --- 10. Architecture & integrations ---
    add_heading(doc, "10. Architecture & integrations", 1)
    add_para(
        doc,
        "The Portal follows the Governance Platform capability architecture: Governance Console & "
        "Applications above Gov Metadata Management Services, Governance Metadata Repository, OMS, "
        "and Compliance & monitoring services.",
    )
    add_table(
        doc,
        ["Layer", "What it is", "Portal connection"],
        [
            (
                "Governance Console & Applications",
                "Portal, workflow UIs, management consoles",
                "Primary UI surface; read/update via Gov Application Services",
            ),
            (
                "Gov Metadata Management Services",
                "Gov app backend, evidence capture, enterprise integrations",
                "APIs for Portal workflow modules; no direct browser writes to Postgres",
            ),
            (
                "Governance Metadata Repository",
                "Postgres (PULS) — classifications, policies, taxonomies, metrics definitions",
                "System of record for governance definitions",
            ),
            (
                "OMS + platforms",
                "Operational metadata, Horizon catalog, tagging enforcement",
                "Navigator/BDC read OMS; defined vs observed compares repo definitions to OMS state",
            ),
            (
                "Compliance & monitoring services",
                "Scanning, DSR orchestration, anonymization, conflict detection",
                "Portal aggregates signals; specialists work in downstream tools",
            ),
            (
                "Policy engines (SAME, DSR, RIM)",
                "Automated enforcement",
                "Read classifications, policies, and rules from governance services",
            ),
        ],
    )
    add_para(doc, "Integration principles:")
    add_bullets(
        doc,
        [
            "One metadata spine: 626/MCI supplies → OMS stores → governance repo defines → Portal orients.",
            "API boundaries: Portal UIs call Gov Application Services and OMS APIs; browsers do not write directly to Postgres.",
            "OMS projection: OMS reads/syncs governance definitions onto catalog assets; governance repo stays tool-agnostic.",
            "Open design question: confirm OMS reads governance definitions directly from Postgres or via Gov App API— affects classification display and schema ownership.",
        ],
    )
    add_heading(doc, "10.1 Build ownership", 2)
    add_table(
        doc,
        ["Component", "Owning team / partner", "Notes"],
        [
            ("Governance Portal UX", "Data Governance (DEET)", "Product, UX, program delivery"),
            ("Gov Application Services & workflow backend", "SISU partnership + Data Governance", "APIs, audit logging, integrations"),
            ("Governance Metadata Repository (PULS)", "SISU partnership + Data Governance program", "Classifications, policies, metrics definitions"),
            ("OMS store & catalog APIs", "OMS Foundations (DFP)", "Operational metadata hub"),
            ("Access Manager", "Platform / access program", "Bidirectional read/update with classifications"),
            ("Business inventory (Kronos)", "Inventory program on SISU Postgres", "Deep link from Regulatory dimension"),
            ("Privacy pillar tools", "Privacy engineering & operations", "Portal aggregates KPIs; does not replace"),
            ("Metadata supply", "Data 626 / MCI + metadata WG", "Feeds OMS; Portal surfaces outcomes"),
        ],
    )

    # --- 11. Phasing ---
    add_heading(doc, "11. Phasing & MVP", 1)
    add_para(
        doc,
        "Portal delivery is sequenced after Business Data Catalog MVP to respect Alation exit "
        "priorities and Gov App backend readiness.",
    )
    add_heading(doc, "Phase 1 — Foundation (parallel to BDC MVP)", 2)
    add_bullets(
        doc,
        [
            "Portal shell & IA: four dimension buckets, Governance Overview layout, Overview + Work queue tab pattern.",
            "Workstream prototypes with representative KPI cards and charts (UDGE demo).",
            "Work queue model and golden-path handoff seeds to Navigator / downstream tools.",
            "Stakeholder reviews with privacy, legal, and program leads on container layout.",
            "Clear Portal vs BDC boundary in intake and demo materials.",
        ],
    )
    add_heading(doc, "Phase 2 — Live hub (Governance Portal GA target)", 2)
    add_bullets(
        doc,
        [
            "Live KPI federation from consent, DSR, tracker, access, GIS, and data-use systems.",
            "Gov App backend + PULS read paths for classifications and policy workflows.",
            "Role-based landing pages and entitlement-aware queue visibility.",
            "Evidence export v1 and auditor view.",
            "v2 governance-first IA recommended for Portal launch.",
        ],
    )
    add_heading(doc, "Phase 3 — Scale & automation", 2)
    add_bullets(
        doc,
        [
            "Metrics governance UI and executive reporting automation.",
            "Resolved OMS↔governance read model for defined vs observed.",
            "Scheduled evidence packages and hardened audit exports.",
            "Enterprise expansion: additional segments inherit dimension model where policy aligns.",
        ],
    )

    # --- 12. Success criteria ---
    add_heading(doc, "12. Success criteria", 1)
    add_bullets(
        doc,
        [
            "Posture indices: PCI and Platform / Security / Regulatory section indices trended month-over-month with target thresholds.",
            "Exception SLA: % of work-queue items closed within SLA; count of breaches by dimension.",
            "Portal adoption: weekly active users among privacy, legal, compliance, and stewardship roles.",
            "Handoff completion: % of Portal-routed items actioned in downstream tools within agreed windows.",
            "Evidence readiness: time to produce audit/regulatory evidence package from Portal export.",
            "Navigation coherence: stewards move Platform → Classifications → BDC asset without dead links.",
            "Stakeholder satisfaction: privacy council feedback on single unified experience vs prior dashboard sprawl.",
        ],
    )

    # --- 13. Non-functional requirements ---
    add_heading(doc, "13. Non-functional requirements", 1)
    add_bullets(
        doc,
        [
            "Security: SSO via corporate IdP; RBAC on every read/write; auditor role is read-only.",
            "Audit: append-only logging for classification/policy mutations and queue disposition changes.",
            "Performance: Governance Overview P95 load target TBD with program; queue pagination for large backlogs.",
            "Reliability: graceful degradation when a downstream KPI feed is unavailable—show stale state, not silent failure.",
            "Privacy: no individual user activity in public leaderboards; aggregate KPIs only.",
            "Availability: align with enterprise governance program SLA (target TBD with SISU/ops).",
        ],
    )

    # --- 14. Dependencies & risks ---
    add_heading(doc, "14. Dependencies & risks", 1)
    add_bullets(
        doc,
        [
            "Depends on: Gov App backend and PULS schema stability; downstream privacy/access system APIs for KPI feeds.",
            "Depends on: OMS catalog asset IDs for defined vs observed and deep links to Business Data Catalog.",
            "Depends on: enterprise design review and privacy/legal sign-off on Overview metrics.",
            "Risk: Portal launches with demo-grade data—mitigate by labeling beta until live feeds verified.",
            "Risk: scope creep absorbs BDC or specialist tools—mitigate with RACI and architecture diagram.",
            "Risk: Foundations/OMS UI collision—escalate via metadata WG when Navigator and Foundations scope overlap.",
            "Risk: incident/problem management integration undefined—track with Darren/Lance if Portal proceeds.",
        ],
    )

    # --- 15. Open questions ---
    add_heading(doc, "15. Open questions", 1)
    add_bullets(
        doc,
        [
            "OMS reads governance repo directly from Postgres or exclusively via Gov App API?",
            "Which KPI feeds are available for Phase 2 GA vs deferred to Phase 3?",
            "Final PCI formula and section index weightings—privacy council approval path?",
            "v1 two-tab vs v2 governance-first IA for combined BDC + Portal shell at GA?",
            "Incident management integration scope with enterprise problem management?",
            "Entitlement model: which queue items visible to which IdP groups per dimension?",
        ],
    )

    # --- 16. References ---
    add_heading(doc, "16. References", 1)
    add_bullets(
        doc,
        [
            "UDGE demo wireframes: data-navigator-governance/index.html",
            "v2 shell prototype: navigator-version-2/index.html",
            "Architecture Q&A memo: requirements-gathering/slides/Data-Gov-Portal-Architecture-QA-Memo.docx",
            "Value deck: requirements-gathering/slides/data-governance-portal-value-deck-3-slides.html",
            "Navigator umbrella PRD: requirements-gathering/prd/word/PRD-Navigator-Data-Governance-Umbrella.docx",
            "Architecture diagram: data-navigator-governance/governance-product-systems-architecture-lucidchart-import.md",
            "OMS alignment deck: requirements-gathering/slides/navigator-oms-alignment-deck-3-slides.html",
            "Intake report: requirements-gathering/meetings/data-platforms-intake-2026-05-28-navigator-report.md",
        ],
    )

    doc.save(OUT)
    return OUT


if __name__ == "__main__":
    path = build()
    print(f"Wrote {path}")
