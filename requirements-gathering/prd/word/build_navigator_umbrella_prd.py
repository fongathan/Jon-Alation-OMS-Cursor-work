#!/usr/bin/env python3
"""Build Navigator umbrella product PRD (.docx)."""
from __future__ import annotations

from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from docx.shared import Pt

OUT = Path(__file__).resolve().parent / "PRD-Navigator-Data-Governance-Umbrella.docx"


def add_heading(doc: Document, text: str, level: int = 1) -> None:
    doc.add_heading(text, level=level)


def add_para(doc: Document, text: str) -> None:
    doc.add_paragraph(text)


def add_bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        doc.add_paragraph(item, style="List Bullet")


def build() -> Path:
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    title = doc.add_paragraph()
    title.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    tr = title.add_run("Product Requirements Document")
    tr.bold = True
    tr.font.size = Pt(16)

    sub = doc.add_paragraph()
    sub.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    sr = sub.add_run("Navigator — Data Governance Product Family")
    sr.bold = True
    sr.font.size = Pt(14)

    meta = doc.add_paragraph()
    meta.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    meta.add_run(f"Version 0.3 · Draft · {date.today().isoformat()}")

    doc.add_paragraph()

    # --- 1. Executive summary ---
    add_heading(doc, "1. Executive summary", 1)
    add_para(
        doc,
        "Navigator is the proposed umbrella data governance product family for Disney "
        "Entertainment & ESPN Technology (DEEPT). Rather than treating catalog discovery, "
        "policy management, program operations, and business inventory as unrelated tools, "
        "Navigator reframes them as coordinated parts of one program—with a shared portal "
        "and explicit boundaries between what Data Governance builds and what platform teams "
        "(OMS, SISU) operate.",
    )
    add_para(
        doc,
        "What was previously called the Governance Portal becomes Navigator: the governance-facing "
        "home that orients users across Privacy, Platform, Security, and Regulatory dimensions. "
        "What was called Business Data Navigator becomes the Business Data Catalog (BDC): the "
        "business-first experience for finding and curating metadata on top of OMS. Behind Navigator "
        "sits the Policy Drive Governance Framework (PULS, Access Manager, Gov App backend). "
        "Kronos remains the system of record for business inventory, linked from Navigator rather "
        "than duplicated inside it.",
    )
    add_para(
        doc,
        "This PRD establishes the product story executives and engineers can align on: why the "
        "family exists, how the pieces connect, who owns what, and what must ship in what order "
        "given Alation exit pressure and OMS platform constraints. Wireframes and the systems "
        "architecture diagram are the visual companions to this document.",
    )

    # --- 2. Problem statement ---
    add_heading(doc, "2. Problem statement", 1)
    add_para(
        doc,
        "DEEPT is exiting Alation while maturing OMS as the operational metadata store. Alation "
        "succeeded because it bundled discovery, glossary, stewardship, and lightweight governance "
        "linkage into one place stewards and analysts actually opened. Replacing Alation with OMS "
        "Foundations' technical catalog alone would leave business users and governance operators "
        "without the workflows they depend on.",
    )
    add_para(
        doc,
        "Governance operations already live in multiple systems. Classifications and policies sit "
        "in PULS. Access decisions flow through Access Manager. Business inventory is maintained in "
        "Kronos on SISU Postgres. Without a unifying product narrative, each system feels like a "
        "separate initiative; stewards struggle to see how a classification in PULS relates to a "
        "table in the catalog, and program leads cannot tell a single posture story to leadership.",
    )
    add_para(
        doc,
        "OMS Foundations is building the store, ingestion, graph, and APIs. Data Governance must "
        "deliver business-facing value on those APIs without forking the platform or re-implementing "
        "ingestion. Navigator exists to draw a clear line—OMS is the engine; Navigator and BDC are "
        "the governance program's vehicles for consuming it.",
    )
    add_para(doc, "The core problems this product family must solve are:")
    add_bullets(
        doc,
        [
            "Continuity: preserve high-value Alation patterns (search, browse, glossary, ownership, "
            "tags) on OMS for in-scope domains before contract end.",
            "Coherence: give governance operators one portal that connects privacy, platform, "
            "security, and regulatory work without re-platforming Kronos or PULS.",
            "Clarity: document and enforce build boundaries so DG owns shells and workflows while "
            "OMS and SISU own stores and backends.",
        ],
    )

    # --- 3. Product vision ---
    add_heading(doc, "3. Product vision", 1)
    add_para(
        doc,
        "The vision is a single Navigator experience that answers two questions for every user: "
        "\"Where is my program work?\" and \"Where is the governed data?\" Program operators "
        "land on Governance Overview, see posture across dimensions, and drill into the workflow "
        "appropriate to their role—consent queues, classification review, access attestations, "
        "regulatory incidents, and so on. Stewards and analysts who need tables, columns, "
        "definitions, and lineage move into Business Data Catalog without leaving the broader "
        "Navigator program.",
    )
    add_para(
        doc,
        "Navigator is not trying to become a monolith that absorbs Kronos or PULS. Inventory "
        "stays in Kronos; policy truth stays in PULS. Navigator's job is orientation, navigation, "
        "workflow orchestration, and evidence packaging—not re-hosting every datastore.",
    )
    add_para(
        doc,
        "Over time, Navigator should feel like the governance program's operating system: the "
        "place where metrics are interpreted, policies are linked to assets, auditors receive "
        "exports, and executives see a defensible posture narrative.",
    )

    # --- 4. Architecture ---
    add_heading(doc, "4. Architecture overview", 1)
    add_para(
        doc,
        "The architecture separates three platform zones, each with a distinct mandate.",
    )
    add_para(
        doc,
        "OMS (Operational Metadata Store) is the technical metadata hub. It ingests connector "
        "metadata, maintains graph relationships (Neptune), and exposes catalog APIs consumed "
        "by downstream experiences. Business Data Catalog reads from OMS—it does not own "
        "ingestion pipelines or the graph store.",
    )
    add_para(
        doc,
        "The SISU relational stack holds business inventory in Postgres. Kronos is the UI for "
        "that inventory. Navigator surfaces Kronos via an Inventory link in the Regulatory "
        "dimension rather than rebuilding inventory screens inside the portal.",
    )
    add_para(
        doc,
        "The governance metadata stack centers on PULS Postgres: classifications and tag "
        "definitions, policies, and metrics. Access Manager maintains bidirectional read/update "
        "with classifications for enforcement alignment. The Gov App backend read/updates policies "
        "and metrics and powers Navigator workflow UIs.",
    )
    add_para(doc, "At a glance, the zones and their primary consumers:")
    add_bullets(
        doc,
        [
            "OMS — technical metadata store and APIs; consumed by BDC and (optionally) governance read paths.",
            "SISU / Kronos — business inventory; consumed by inventory operators; linked from Navigator.",
            "PULS + Access Manager + Gov App backend — policy and control metadata; consumed by Navigator workflows.",
        ],
    )
    add_para(
        doc,
        "One open design question remains: the \"Read ?\" link from OMS to the Governance Metadata "
        "Repository. Confirm whether OMS reads classifications, policies, and metrics directly from "
        "PULS Postgres or exclusively through the Gov App backend API. The answer affects latency, "
        "coupling, and who owns schema evolution when PULS changes.",
    )

    # --- 5. Product components ---
    add_heading(doc, "5. Product components", 1)
    add_para(
        doc,
        "Four named components make up the Navigator family. Each has a clear user outcome, "
        "a system-of-record boundary, and a build owner.",
    )

    add_heading(doc, "5.1 Navigator (governance portal & Gov App backend)", 2)
    add_para(
        doc,
        "Navigator is the main data governance portal—the experience formerly branded Governance "
        "Portal—renamed to match the umbrella product. Its primary users are privacy operators, "
        "security and access reviewers, regulatory program managers, governance leads, and auditors.",
    )
    add_para(
        doc,
        "The portal is structured around four dimension buckets that mirror how DEEPT runs "
        "governance programs in practice. Privacy covers consent, data-subject requests, and "
        "tracker remediation. Platform connects stewards to Business Data Catalog and houses "
        "classifications and policies—the bridge between \"program\" and \"catalog.\" Security "
        "covers access, GIS compliance, and data use. Regulatory covers laws and regulations, "
        "incidents, and inventory (via Kronos).",
    )
    add_para(
        doc,
        "Workflow UIs inside Navigator are the human interface to the Gov App backend, which "
        "read/updates PULS. Evidence export and auditor view support the need to package control "
        "narratives for attestation cycles.",
    )
    add_para(doc, "Navigator delivers:")
    add_bullets(
        doc,
        [
            "Governance Overview with program metrics and trend views across Privacy, Platform, Security, and Regulatory.",
            "Dimension buckets with deep links into in-portal workflows and external systems (BDC, Kronos).",
            "Gov App backend integration for classifications, policies, and metrics (phased).",
            "Auditor-oriented export and read-only views for attestation cycles.",
        ],
    )
    add_para(
        doc,
        "Build ownership: Data Governance owns the Navigator shell, information architecture, "
        "and workflow UX. SISU partners on Gov App backend services and PULS connectivity.",
    )

    add_heading(doc, "5.2 Policy Drive Governance Framework", 2)
    add_para(
        doc,
        "The Policy Drive Governance Framework is the policy and control metadata layer behind "
        "Navigator. With it, classifications applied in the catalog trace to PULS definitions, "
        "policies map to assets, and metrics reflect measurable program posture rather than static copy.",
    )
    add_para(
        doc,
        "PULS (Postgres) is the relational system of record. Classifications and tag definitions "
        "encode sensitivity, domain, and control labels. Policies encode obligations and control "
        "mappings. Metrics encode governance KPIs that feed Governance Overview and executive "
        "reporting.",
    )
    add_para(
        doc,
        "Access Manager sits adjacent to classifications because enforcement is where governance "
        "meets reality: a label without access alignment is documentation only. The bidirectional "
        "read/update path between Access Manager and PULS classifications is a deliberate "
        "integration point.",
    )
    add_para(
        doc,
        "The Gov App backend mediates Navigator's write paths to policies and metrics. Front-end "
        "workflow UIs should not write directly to Postgres from the browser; the backend provides "
        "validation, audit logging, and a stable API surface as PULS schema evolves.",
    )
    add_para(doc, "Framework elements:")
    add_bullets(
        doc,
        [
            "PULS Postgres — classifications/tags, policies, metrics (system of record).",
            "Access Manager — classification enforcement and access-control alignment.",
            "Gov App backend — API layer for Navigator workflow UIs.",
        ],
    )

    add_heading(doc, "5.3 Business Data Catalog (formerly Business Data Navigator)", 2)
    add_para(
        doc,
        "Business Data Catalog (BDC) is the business-first UI on top of OMS APIs and the primary "
        "Alation migration target for stewards, analysts, and engineers who need to answer "
        "\"what data exists, what does it mean, who owns it, and can I use it?\"",
    )
    add_para(
        doc,
        "BDC is built as a consumer of OMS, not a competitor to it. OMS Foundations optimizes "
        "ingestion, graph completeness, and API stability. BDC optimizes discoverability, business "
        "language (glossary, descriptions, MCI-style context), and stewardship workflows that "
        "Alation users expect.",
    )
    add_para(
        doc,
        "Users open BDC from the left navigation (Business Data Catalog) or from Navigator's "
        "Platform dimension via Stewards. MVP scope favors proven Alation workflows: Data Sources "
        "browse, catalog search, glossary, asset detail, ownership, tags/classifications display, "
        "and lineage views where OMS graph data is available.",
    )
    add_para(doc, "BDC scope highlights:")
    add_bullets(
        doc,
        [
            "Data Sources, Catalog, Glossary, asset detail, lineage, and stewardship flows.",
            "Reads OMS REST/graph APIs; no duplicate ingestion or graph store.",
            "Surfaces classifications/tags delivered via governance integration (path TBD per Read ?).",
            "Phased parity with Alation high-use features; explicit out-of-scope list per release.",
        ],
    )
    add_para(
        doc,
        "Build ownership: DG owns BDC product and UX; OMS team owns operational metadata store, "
        "connectors, and APIs.",
    )

    add_heading(doc, "5.4 Kronos (business inventory)", 2)
    add_para(
        doc,
        "Kronos is where business inventory is accessed and stored. It runs against SISU Postgres "
        "(Business Inventory) and serves teams who maintain the authoritative record of what "
        "business data exists, for what purpose, and under what obligations—independent of "
        "technical table metadata in OMS.",
    )
    add_para(
        doc,
        "Duplicating Kronos inside Navigator would fork truth and double maintenance. The wireframes "
        "link Regulatory → Inventory to Kronos (https://dcf.corp.dig.com/vista/kronos/) as an "
        "external application.",
    )
    add_para(doc, "Kronos relationship to Navigator:")
    add_bullets(
        doc,
        [
            "System of record for business inventory on SISU Postgres.",
            "Deep-linked from Navigator Regulatory dimension; not embedded in Phase 1.",
            "Complements OMS technical metadata—business context vs. warehouse objects.",
        ],
    )

    # --- 6. Personas ---
    add_heading(doc, "6. User personas", 1)
    add_para(
        doc,
        "Navigator serves multiple personas with different success criteria. A single UI shell "
        "must not force analysts through compliance queues or auditors through glossary edit "
        "screens.",
    )
    add_para(doc, "Primary roles:")
    table = doc.add_table(rows=5, cols=2)
    table.style = "Table Grid"
    personas = [
        (
            "Data Steward / Owner",
            "Curates catalog metadata in BDC; applies and reviews classifications via Navigator/PULS.",
        ),
        (
            "Privacy / Compliance Operator",
            "Runs consent, DSR, tracker, and regulatory programs in Navigator; needs queues, SLAs, "
            "and exportable evidence.",
        ),
        (
            "Data Analyst / BI",
            "Discovers governed tables and reports in BDC; trusts glossary and ownership.",
        ),
        (
            "Governance Program Lead",
            "Uses Governance Overview and metrics for posture reporting and escalations.",
        ),
        (
            "Platform Engineer",
            "Integrates via OMS APIs; expects BDC to reflect ingestion outcomes without a parallel catalog store.",
        ),
    ]
    for i, (role, need) in enumerate(personas):
        table.rows[i].cells[0].text = role
        table.rows[i].cells[1].text = need

    # --- 7. Phasing ---
    add_heading(doc, "7. Phasing & priorities", 1)
    add_para(
        doc,
        "Delivery is constrained by Alation exit timelines, OMS API maturity, and finite DG/SISU "
        "capacity. Phasing front-loads catalog value (BDC on OMS) while sequencing Gov App backend "
        "and PULS integration behind a stable Navigator shell.",
    )
    add_para(
        doc,
        "Phase 1 establishes the product story and migration beachhead. Business Data Catalog MVP "
        "consumes OMS APIs for in-scope domains. Navigator wireframes reflect renamed components, "
        "dimension buckets, Governance Overview layout, and the Kronos inventory link.",
    )
    add_para(
        doc,
        "Phase 2 connects Navigator to authoritative policy data. Gov App backend and PULS read "
        "paths power classifications and policy workflows inside the portal.",
    )
    add_para(
        doc,
        "Phase 3 deepens automation and executive readiness: metrics governance UI, resolution of "
        "the OMS↔PULS Read ? pattern, richer Governance Overview automation, and hardened evidence "
        "export for audit cycles.",
    )
    add_para(doc, "Phase summary:")
    add_bullets(
        doc,
        [
            "Phase 1 — BDC MVP on OMS APIs; Navigator IA and wireframes; Kronos deep link.",
            "Phase 2 — Gov App backend + PULS; classification and policy workflows live in Navigator.",
            "Phase 3 — Metrics UI, OMS↔PULS read model, executive reporting automation.",
        ],
    )

    # --- 8. Success criteria ---
    add_heading(doc, "8. Success criteria (high level)", 1)
    add_para(
        doc,
        "Success is measured by adoption and continuity of governed work—not by shipping every "
        "workflow in one release.",
    )
    add_para(
        doc,
        "For BDC, the bar is practical replacement: an analyst in an in-scope domain completes "
        "primary discovery without opening Alation. For Navigator, the bar is coherent program "
        "operations: a steward moves from Governance Overview to classifications or policies to a "
        "catalog asset without hitting a dead link.",
    )
    add_para(doc, "Program-level indicators:")
    add_bullets(
        doc,
        [
            "Catalog continuity: in-scope Alation workflows available in BDC before contract end.",
            "Navigation coherence: Platform and Regulatory deep links (BDC, Kronos) work from Navigator.",
            "Architecture clarity: documented ownership for OMS, PULS, SISU, and DG components.",
            "Posture reporting: Governance Overview metrics traceable to PULS/metrics backend (Phase 2+).",
        ],
    )

    # --- 9. Dependencies & risks ---
    add_heading(doc, "9. Dependencies & risks", 1)
    add_para(
        doc,
        "Navigator is a coordination product. Its schedule depends on OMS API completeness and "
        "operational SLA for BDC—search, browse, asset detail, lineage, and stewardship writes "
        "must be versioned and monitored.",
    )
    add_para(
        doc,
        "PULS and Access Manager sequencing is the second critical path. If Navigator workflow "
        "UIs launch before Gov App backend read paths exist, the portal risks demo-grade data "
        "that undermines trust with compliance stakeholders. The mitigation: ship BDC on OMS while "
        "labeling governance workflows as beta until PULS integration is verified.",
    )
    add_para(
        doc,
        "Organizational clarity remains a risk. If stakeholders conflate BDC with OMS Foundations' "
        "catalog application, teams will argue about duplicate features instead of complementary roles.",
    )
    add_para(doc, "Key risks and mitigations:")
    add_bullets(
        doc,
        [
            "OMS API gaps — publish BDC-required API matrix; joint SLA with OMS Foundations.",
            "PULS integration delay — phase Navigator workflows; show real data paths before GA.",
            "Scope creep / duplicate catalog — enforce ownership boundaries in PRD and architecture diagram.",
            "Alation exit pressure — prioritize proven workflows; defer low-adoption Alation features.",
        ],
    )

    # --- 10. References ---
    add_heading(doc, "10. References", 1)
    add_para(
        doc,
        "Supporting artifacts in the program workspace:",
    )
    add_bullets(
        doc,
        [
            "Wireframes: data-navigator-governance/index.html",
            "Architecture: data-navigator-governance/governance-product-systems-architecture-lucidchart-import.md",
            "Migration brief: Product-Brief-Alation-to-OMS-Migration.md",
            "Program plan: Alation-to-OMS-Migration-Project-Plan.md",
            "Overview deck: requirements-gathering/slides/navigator-umbrella-overview-deck.html",
        ],
    )

    doc.save(OUT)
    return OUT


if __name__ == "__main__":
    path = build()
    print(f"Wrote {path}")
