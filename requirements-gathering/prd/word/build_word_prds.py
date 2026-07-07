#!/usr/bin/env python3
"""
Build one Word (.docx) PRD per feature (1–35) from the DEEPT Feature Inventory Matrix.
Output: PRD-F{nn}-{slug}.docx in this directory.

Re-run after matrix changes: python3 build_word_prds.py
"""
from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from docx.shared import Pt

OUT_DIR = Path(__file__).resolve().parent

# Source: Google Doc Feature Inventory Matrix (DEEPT, Mar 19, 2026)
ROWS: list[tuple[int, str, str, str, str, str]] = [
    (
        1,
        "Search & Discovery",
        "Keyword and semantic search across catalog for tables, schemas, columns, articles, glossary terms, AI contextually predicted results, AI natural language search for non-SQL users",
        "High",
        "Find data assets",
        "What do users search for most?",
    ),
    (
        2,
        "Business Glossary",
        "Terms, definitions, synonyms, stewardship, linked to catalog objects",
        "High",
        "Standardize terminology",
        "How many terms? Who maintains?",
    ),
    (
        3,
        "Catalog Browse",
        "Browse datasources, schemas, tables by hierarchy",
        "High",
        "Explore by source/schema",
        "Primary navigation pattern?",
    ),
    (
        4,
        "Table / Asset Detail",
        "View metadata for tables, columns: description, owner, tags, lineage",
        "High",
        "Understand asset context",
        "Which fields are most viewed?",
    ),
    (
        5,
        "Data Lineage",
        "Upstream/downstream lineage (table, column, dataflow level)",
        "High",
        "Impact analysis",
        "Table-level vs. column-level?",
    ),
    (
        6,
        "Tags & Classification",
        "Apply tags (PII, domain, custom) to assets",
        "High",
        "Classification, compliance",
        "Which tags are critical?",
    ),
    (
        7,
        "Custom Fields",
        "User-defined metadata fields on assets",
        "Medium",
        "Domain-specific metadata",
        "List custom fields in use",
    ),
    (
        8,
        "Documentation / Articles",
        "Rich-text documentation linked to catalog objects",
        "High",
        "How-to guides, data dictionaries",
        "Article count, structure",
    ),
    (
        9,
        "Ownership / Stewardship",
        "Assign owners, stewards to assets",
        "High",
        "Accountability, escalation",
        "Steward assignment coverage",
    ),
    (
        10,
        "Usage / Access Metrics",
        "Last access, unique users, top users, query history",
        "Medium",
        "Usage reporting",
        "Who consumes these reports?",
    ),
    (
        11,
        "Data Sources & Connectors",
        "Connected datasources (Snowflake, etc.)",
        "High",
        "Technical metadata ingestion",
        "Which sources are critical?",
    ),
    (
        12,
        "Policy Center",
        "Policies, risk classification, compliance rules",
        "Medium",
        "Governance, compliance",
        "Policies in use",
    ),
    (
        13,
        "Compilation / Data Quality",
        "Data quality rules, profiling",
        "Low–Medium",
        "Quality monitoring",
        "Extent of use",
    ),
    (
        14,
        "API & Integrations",
        "REST API for bulk ops, custom metadata, BI tool integration",
        "Medium",
        "Automation, tool integration",
        "Which integrations?",
    ),
    (
        15,
        "Version Control",
        "Asset versioning, change history",
        "Low–Medium",
        "Audit trail, rollback",
        "How often used?",
    ),
    (
        16,
        "Recommendations",
        "AI-suggested metadata, stewardship actions",
        "Low",
        "Metadata improvement",
        "Adoption level",
    ),
    (
        17,
        "Saved Queries (Compose) / Favorites",
        "Save searches, bookmark assets",
        "Medium",
        "Quick access to frequent assets",
        "Usage patterns",
    ),
    (
        18,
        "Comments / Discussions",
        "Comments on assets, Q&A",
        "Low–Medium",
        "Collaboration, questions",
        "Active use?",
    ),
    (
        19,
        "Export / Reports",
        "Export catalog, usage reports, compliance reports",
        "Medium",
        "Leadership reporting, audits",
        "Report types",
    ),
    (
        20,
        "Authentication / SSO",
        "Single sign-on, role-based access",
        "High",
        "Access control",
        "SSO provider",
    ),
    (
        21,
        "Flags (Endorsements, Warnings, Deprecations)",
        "Data quality signals on catalog objects (trusted, deprecated, etc.)",
        "High",
        "Data quality signals",
        "Examples: trusted endorsement, deprecated warning",
    ),
    (
        22,
        "Domains",
        "Business area groupings (e.g., Ad Platforms, DEET Data, Atlas)",
        "High",
        "Business area groupings",
        "Used to scope search, assign stewards, and apply domain policies",
    ),
    (
        23,
        "Folders",
        "Article/content hierarchy and organization",
        "High",
        "Content hierarchy",
        "Examples: How-to folder, project-specific article collections",
    ),
    (
        24,
        "Dataflows",
        "ETL/transformation jobs, stored procedures, dbt models, Airflow DAGs. Distinct from lineage: pipeline definitions/executables.",
        "High",
        "ETL / pipelines",
        "Examples: dbt model 'sales_transform', Airflow DAG 'daily_load'.",
    ),
    (
        25,
        "Admin Analytics dashboard",
        "Collects telemetry (searches, views, queries, top assets, user activity, e.g. in PostgreSQL); pre-built dashboard plus query interface for catalog analytics; filters by time, asset type, user, team",
        "High",
        "Usage reporting, analytics",
        "Admin views MAU, top-searched datasets, declining usage queries; exports CSV; admin role",
    ),
    (
        26,
        "Mass/Bulk edit",
        "Select multiple assets and apply metadata changes (owners, tags, descriptions) in one operation; CSV upload with audit logging",
        "High",
        "Bulk metadata updates",
        "Update owner + compliance tag on 200 tables; import CSV mapping",
    ),
    (
        27,
        "Mass set rules / catalog sets",
        "Define rules (pattern, metadata condition) on asset sets; rule engine matches assets and applies tags/policies or ownership at scale",
        "Medium",
        "Governance automation",
        'Tables with "pii" in column names → retention policy + steward',
    ),
    (
        28,
        "AI chat assistant in app",
        "LLM-backed interface using catalog metadata to answer questions, locate datasets, suggest SQL, summarize definitions",
        "Medium",
        "Natural language search, discovery",
        'e.g. "Where is revenue defined?" → dataset, column, sample SQL, glossary',
    ),
    (
        29,
        "Data Dictionary download & uploads",
        "Download catalog metadata (titles, descriptions, custom fields) for sources and children; upload CSV/Excel to bulk-update many assets",
        "Medium",
        "Bulk metadata export/import, documentation",
        "Steward assignments, descriptions, custom fields at scale",
    ),
    (
        30,
        "Expanded Data Sources",
        "Vendor data sources outside of Snowflake (e.g. high-touch systems)",
        "High",
        "Technical metadata ingestion",
        "Consider Data 626’s existing connectors",
    ),
    (
        31,
        "Details icon",
        '“I” icon or “What’s this” to explain each feature in context (metadata for the metadata)',
        "Medium",
        "Help users understand UI capabilities",
        "Which features need inline help?",
    ),
    (
        32,
        "Persona – view filtering",
        "Filter or tailor the catalog experience by persona (e.g. Data engineer vs. Sales)",
        "Medium",
        "Customize experience; reduce noise",
        "Which personas and which views?",
    ),
    (
        33,
        "Product Pages",
        "Per-product documentation: impacted metrics, links, contacts, data domain, stewards—optionally AI-assisted from wiki and internal sources",
        "Medium",
        "Home for products; links to metrics and catalog assets",
        "Interconnect catalog with product truth",
    ),
    (
        34,
        "AI – “How to”",
        "Ask AI how to use the catalog or where to go for a need (in-app help)",
        "Low",
        "Help center for users",
        "Governance for answers and PII",
    ),
    (
        35,
        "Overview section",
        "Sub-home for Business Data Catalog: snapshot of catalog usage—social feeds, popular tables/searches, active searchers (aggregated, not individual PII)",
        "Medium",
        "Discovery via community usage patterns",
        "Must not expose individual activity; partnership + enforcement note from matrix",
    ),
]


def slugify(title: str) -> str:
    s = title.lower()
    s = re.sub(r"[^a-z0-9]+", "-", s).strip("-")
    s = re.sub(r"-+", "-", s)
    return s[:80] or "feature"


def add_heading(doc: Document, text: str, level: int = 1) -> None:
    doc.add_heading(text, level=level)


def add_para(doc: Document, text: str, bold: bool = False) -> None:
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = bold
    run.font.size = Pt(11)


def add_bullets(doc: Document, items: list[str]) -> None:
    for it in items:
        doc.add_paragraph(it, style="List Bullet")


def component_library(
    num: int, title: str, desc: str, primary: str, notes: str, priority: str
) -> list[tuple[str, list[str]]]:
    """Return (component_name, bullet_points) for piece-by-piece breakdown."""
    common_close = [
        "Telemetry: key events for adoption reporting (where privacy allows).",
        "Documentation: in-app help links (see also Feature 31, 34).",
    ]
    t = title.lower()

    def pack(name: str, *bullets: str) -> tuple[str, list[str]]:
        return (name, list(bullets))

    # Feature-specific packs
    if num == 1:
        return [
            pack(
                "Search index & coverage",
                "Ingest object types: tables, columns, schemas, articles, glossary terms, policies as in scope.",
                "Include business metadata fields (owner, domain, tags) in index.",
                "Optional: semantic / vector index for similarity and NL queries.",
                "AI contextual predictions: surface likely next assets from behavior (policy-gated).",
            ),
            pack(
                "Query experience",
                "Global search entry; scope toggle (all sources vs. current context).",
                "Natural language mode for non-SQL users with grounded answers and citations.",
                "Query suggestions, spell-tolerance, synonym expansion from glossary.",
            ),
            pack(
                "Results & ranking",
                "Rank by text match, popularity, trust flags, recency, domain affinity.",
                "Snippets must respect RBAC (no forbidden object names in preview if denied).",
            ),
            pack(
                "Filters & facets",
                "Post-search facets: domain, environment, sensitivity, owner, object type, trust state.",
            ),
            pack(
                "Security & compliance",
                "Enforce SSO identity; row/scope rules for catalog visibility.",
                "Audit sensitive searches if required by policy.",
            ),
        ]
    if num == 31:
        return [
            pack(
                "Trigger & placement",
                "Consistent “i” or “What’s this?” affordance next to labels, columns, and toolbar actions.",
                "Keyboard accessible; screen-reader friendly description.",
            ),
            pack(
                "Content model",
                "Short summary + “Learn more” deep link to docs or Product Pages (Feature 33).",
                "Versioned copy per release; owners for each help topic.",
            ),
            pack(
                "Localization & governance",
                "Optional locale strings; legal/compliance review for sensitive areas.",
            ),
        ]
    if num == 32:
        return [
            pack(
                "Persona definitions",
                "Configurable personas (e.g. Engineer, Analyst, Steward, Sales) mapped to IdP groups or self-selection with guardrails.",
            ),
            pack(
                "View presets",
                "Default columns, filters, nav density, and optional tabs per persona.",
                "Allow override: users can switch persona or reset to org default.",
            ),
            pack(
                "Persistence",
                "Save last-used persona; optional admin-locked defaults per org unit.",
            ),
        ]
    if num == 33:
        return [
            pack(
                "Product record",
                "Stable product id; name; executive summary; domain; stewards; contacts.",
                "Links to metrics, dashboards, key tables, glossary entries.",
            ),
            pack(
                "AI ingestion pipeline",
                "Connectors to wiki / Confluence / docs; extraction, dedup, human review queue before publish.",
                "Attribution: every AI-derived block shows source links and confidence.",
            ),
            pack(
                "Publishing workflow",
                "Draft → review → published; scheduled refresh with diff alerts to stewards.",
            ),
        ]
    if num == 34:
        return [
            pack(
                "Help scope",
                "Catalog navigation, permissions, how to request access, glossary usage, export patterns.",
            ),
            pack(
                "Safety",
                "Grounded on approved help corpuses; refuse unsupported how-tos; no PII in logs.",
            ),
            pack(
                "Surfacing",
                "Persistent entry (e.g. help panel) + contextual launch from Feature 31 tooltips.",
            ),
        ]
    if num == 35:
        return [
            pack(
                "Widgets & KPIs",
                "Popular tables (aggregated), trending searches (hashed or bucketed), quality/completeness rollups.",
                "Optional social feed: announcements, curated highlights—not raw user timelines.",
            ),
            pack(
                "Privacy",
                "No per-user public leaderboards; minimum aggregation thresholds; redaction rules.",
            ),
            pack(
                "Data pipeline",
                "Telemetry warehouse (e.g. PostgreSQL) feeds widgets with SLAs and freshness labels.",
            ),
        ]
    if num == 30:
        return [
            pack(
                "Connector strategy",
                "Prioritize high-touch vendor systems beyond Snowflake; reuse internal connector investments (e.g. Data 626) where possible.",
            ),
            pack(
                "Metadata model",
                "Normalize external source objects into OMS/BDC asset types; lineage hooks where available.",
            ),
            pack(
                "Operations",
                "Credential vault integration, sync schedules, failure surfacing, ownership per connector.",
            ),
        ]
    if "glossary" in t:
        return [
            pack("Term lifecycle", "Draft, published, deprecated; owners; review dates."),
            pack("Rich definitions", "Synonyms, abbreviations, related terms, policy callouts."),
            pack("Mappings", "Link terms to tables/columns/metrics; cardinality and confidence."),
            pack("Consumption", "Inline on asset pages, search snippets, export columns."),
        ]
    if "lineage" in t:
        return [
            pack("Graph & paths", "Upstream/downstream at table and column granularity where supported."),
            pack("Pipeline linkage", "Integrate dataflow objects (Feature 24) as lineage nodes."),
            pack("Impact analysis", "Export/subgraph for change planning; depth and RBAC limits."),
        ]
    if "bulk" in t or num == 26:
        return [
            pack("Selection", "Checkbox selection, select-all in view, select-by-filter with safeguards."),
            pack("Batch editor", "Apply tags, owners, descriptions; preview diff before commit."),
            pack("CSV pipeline", "Template download, validation report, dry-run, async job for large sets."),
            pack("Audit", "Actor, timestamp, before/after per field; downloadable job log."),
        ]
    if "dictionary" in t or num == 29:
        return [
            pack("Export", "Column picker, stable keys (FQN), versioned template headers."),
            pack("Import", "Row-level validation; permission checks; partial success semantics."),
            pack("Round-trip", "Re-import unchanged rows is idempotent."),
        ]
    if "sso" in t or "authentication" in t:
        return [
            pack("Identity", "SAML/OIDC with corporate IdP; session timeout aligned to policy."),
            pack("Authorization", "Roles: admin, steward, editor, consumer, auditor; optional domain scope."),
            pack("Service access", "PATs or service principals for APIs with narrow scopes."),
        ]

    # Default rich breakdown for remaining features
    return [
        pack(
            "Capability overview",
            f"Addresses matrix priority **{priority}** and primary use case: **{primary}**.",
            f"Matrix description: {desc}",
        ),
        pack(
            "User-facing experience",
            "Navigation entry points, empty states, loading and error handling.",
            "Consistency with OMS integrated vs. BDC on-top patterns where applicable.",
        ),
        pack(
            "Data & integrations",
            "Read/write paths to OMS metadata services; connector or warehouse dependencies as relevant.",
        ),
        pack(
            "Permissions & audit",
            "Role checks on every mutation; append-only audit for compliance-sensitive fields.",
        ),
        pack(
            "Observability & quality",
            "Metrics for feature usage, failure rates, and data freshness surfaced to admins.",
        ),
        pack(
            "Rollout",
            "Feature flags, steward training assets, migration notes from Alation parity.",
        ),
    ]


def build_document(
    num: int, title: str, desc: str, priority: str, primary: str, notes: str
) -> Document:
    doc = Document()
    title_text = f"PRD — Feature {num:02d}: {title}"
    h = doc.add_heading(title_text, 0)
    h.alignment = WD_PARAGRAPH_ALIGNMENT.LEFT

    add_para(doc, "Product Requirements Document", True)
    add_para(
        doc,
        "Source: Feature Inventory Matrix — Alation to Internal Solution Migration (DEEPT). "
        "Document control: link this file from the matrix column **PRD page link**.",
    )
    add_para(doc, f"Priority (matrix): {priority}")
    add_para(doc, f"Primary use case (matrix): {primary}")
    add_para(doc, f"Validation notes (matrix): {notes}")

    add_heading(doc, "1. Summary", 1)
    add_para(
        doc,
        f"This PRD specifies **{title}** for the in-house data catalog / OMS program. "
        f"It expands the single matrix row into implementable components, UX expectations, and acceptance criteria.",
    )

    add_heading(doc, "2. Problem & outcomes", 1)
    add_para(
        doc,
        "Users and stewards need predictable behavior, clear accountability, and parity with "
        "current Alation workflows where this capability is in use. Success is measured by adoption, "
        "reduced time-to-answer, and governance coverage—not feature count alone.",
    )

    add_heading(doc, "3. Scope", 1)
    add_heading(doc, "3.1 In scope", 2)
    add_bullets(
        doc,
        [
            f"End-to-end behavior for: {title}.",
            "DEEPT / enterprise personas: steward, analyst, engineer, governance, admin.",
            "Accessibility: WCAG-oriented patterns for core flows (exact tier set by program).",
        ],
    )
    add_heading(doc, "3.2 Out of scope (unless added by change request)", 2)
    add_bullets(
        doc,
        [
            "Duplicating full warehouse compute consoles unless explicitly linked as deep integration.",
            "Non-catalog products (e.g. full ITSM) except via links or thin embeds.",
        ],
    )

    add_heading(doc, "4. Personas", 1)
    table = doc.add_table(rows=5, cols=2)
    table.style = "Table Grid"
    personas = [
        ("Data Steward / Owner", "Creates and curates metadata; accountable for domain quality."),
        ("Data Analyst / BI", "Discovers and consumes datasets; needs trust and clarity."),
        ("Data Engineer", "Needs lineage, technical metadata, automation via API."),
        ("Governance / Compliance", "Needs policy coverage, exports, and attestations."),
        ("Platform Admin", "Operates connectors, roles, and platform health."),
    ]
    for i, (role, need) in enumerate(personas):
        table.rows[i].cells[0].text = role
        table.rows[i].cells[1].text = need

    add_heading(doc, "5. Component breakdown (piece by piece)", 1)
    for comp_name, bullets in component_library(num, title, desc, primary, notes, priority):
        add_heading(doc, comp_name, 2)
        add_bullets(doc, bullets)

    add_heading(doc, "6. User stories", 1)
    stories = [
        f"As a steward, I can perform core curation tasks for **{title}** without leaving the catalog shell.",
        f"As an analyst, I can rely on **{title}** to decide whether a dataset is appropriate for my use case.",
        f"As an engineer, I can integrate or automate **{title}** via documented APIs where applicable.",
    ]
    add_bullets(doc, stories)

    add_heading(doc, "7. Use cases & examples", 1)
    add_bullets(
        doc,
        [
            f"Happy path: user completes the primary job ({primary}) in under the program’s target time with no support ticket.",
            f"Edge case: partial metadata from source systems; UI states 'unknown' or 'stale' with last sync time.",
            f"Governance case: auditor requests evidence of who changed business metadata related to **{title}**.",
        ],
    )

    add_heading(doc, "8. UI / UX notes", 1)
    add_bullets(
        doc,
        [
            "Use clear hierarchy: primary action visible, destructive actions confirmed.",
            "Tables: pagination, sticky headers, column resize where appropriate.",
            "Mobile: read-mostly experience acceptable unless program mandates full edit on mobile.",
            "Integrated OMS vs. dedicated BDC shell: follow program IA decision; deep-link between shells.",
        ],
    )

    add_heading(doc, "9. Data model & APIs", 1)
    add_bullets(
        doc,
        [
            "Prefer OMS as operational system of record for technical metadata.",
            "Business metadata via versioned REST (or program-standard) contracts; document error codes.",
            "Bulk operations return job ids for large writes; support dry-run where risk is high.",
        ],
    )

    add_heading(doc, "10. Non-functional requirements", 1)
    add_bullets(
        doc,
        [
            "Performance: define P95 targets per surface (search, browse, detail).",
            "Reliability: graceful degradation when telemetry or optional AI services are down.",
            "Security: RBAC on every read/write; no PII in analytics where matrix calls for aggregation (esp. Feature 35).",
        ],
    )

    add_heading(doc, "11. Acceptance criteria (samples)", 1)
    add_bullets(
        doc,
        [
            "Given an authenticated user with appropriate role, when they use the primary flow then the matrix **Primary use case** is achievable without error.",
            "Given insufficient permissions, when the user attempts a restricted action then the system denies with an explanatory message (not silent failure).",
            "Given stale connector data, when the user views dependent UI then stale state is visible with last successful sync timestamp.",
        ],
    )

    add_heading(doc, "12. Dependencies & risks", 1)
    add_bullets(
        doc,
        [
            "Depends on: identity (SSO), connector health, glossary/tags where referenced.",
            "Risks: scope creep between OMS integrated and BDC-on-top; mitigate with shared API and UX contract.",
            f"Matrix note to resolve in interviews: {notes}",
        ],
    )

    add_heading(doc, "13. Open questions", 1)
    add_bullets(
        doc,
        [
            "MVP vs. phase-2 cut line for this feature?",
            "Which environments (prod / staging) ship first?",
            "Executive reporting: which KPIs prove value for this capability?",
        ],
    )

    add_heading(doc, "14. References", 1)
    add_bullets(
        doc,
        [
            "Feature Inventory Matrix (Google Doc): Feature row " + str(num),
            "Workspace: `requirements-gathering/1-Feature-Inventory-Matrix.md` (if kept in sync)",
            "BDC / OMS UX: `business-data-catalog/BDC-OMS-integrated-vs-on-top-decision.md`",
        ],
    )

    return doc


def main() -> None:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    for num, title, desc, priority, primary, notes in ROWS:
        slug = slugify(title)
        fn = OUT_DIR / f"PRD-F{num:02d}-{slug}.docx"
        doc = build_document(num, title, desc, priority, primary, notes)
        doc.save(fn)
        print(fn)
    print(f"Done: {len(ROWS)} Word documents in {OUT_DIR}")


if __name__ == "__main__":
    main()
