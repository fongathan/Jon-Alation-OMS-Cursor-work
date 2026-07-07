#!/usr/bin/env python3
"""Build one-page Executive Summary (.docx) for Phillip — July 2026 Gov/DFP workshop."""
from __future__ import annotations

from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from docx.shared import Inches, Pt, RGBColor

OUT = Path(__file__).resolve().parent / "Workshop-Executive-Summary-Phillip.docx"

NAVY = RGBColor(0x0F, 0x27, 0x44)


def set_tight_spacing(paragraph, before: int = 0, after: int = 3, line: float = 1.08) -> None:
    pf = paragraph.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = line


def add_heading(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    set_tight_spacing(p, before=6, after=2)
    run = p.add_run(text.upper())
    run.bold = True
    run.font.size = Pt(10)
    run.font.color.rgb = NAVY


def add_body(doc: Document, text: str, bold: bool = False) -> None:
    p = doc.add_paragraph()
    set_tight_spacing(p)
    run = p.add_run(text)
    run.font.size = Pt(10)
    run.bold = bold


def add_bullets(doc: Document, items: list[str], size: int = 10) -> None:
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        set_tight_spacing(p, after=1)
        p.paragraph_format.left_indent = Inches(0.2)
        run = p.add_run(item)
        run.font.size = Pt(size)


def build() -> Path:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.55)
    section.bottom_margin = Inches(0.5)
    section.left_margin = Inches(0.65)
    section.right_margin = Inches(0.65)

    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(10)

    title = doc.add_paragraph()
    title.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    set_tight_spacing(title, before=0, after=2)
    tr = title.add_run("Executive Summary")
    tr.bold = True
    tr.font.size = Pt(16)
    tr.font.color.rgb = NAVY

    sub = doc.add_paragraph()
    sub.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    set_tight_spacing(sub, after=2)
    sr = sub.add_run(
        "Data Governance & Data Foundations — Complementary Capabilities Workshop"
    )
    sr.bold = True
    sr.font.size = Pt(11)

    meta = doc.add_paragraph()
    meta.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    set_tight_spacing(meta, after=6)
    meta.add_run(
        f"For Phillip · July 21–23, 2026 · Seattle · Draft {date.today().strftime('%B %d, %Y')}"
    ).font.size = Pt(9)

    add_heading(doc, "The ask")
    add_body(
        doc,
        "Approve a 2.5-day cross-team workshop (July 21–23, Disney Seattle) with Data "
        "Governance and Data Foundations Platform (DFP) leaders. Estimated travel: ~10 "
        "attendees from Santa Monica, Glendale, Texas, and San Francisco (cost TBD — Finance).",
    )
    add_body(doc, "Deliverables to you by July 29, 2026:", bold=True)
    add_bullets(
        doc,
        [
            "Alignment memo — portal roles, overlap resolution, escalation path",
            "RACI by capability — glossary, dictionary, classification, policy, lineage, stewardship",
            "Joint roadmap slice — Q4 FY26 / Q1 FY27 aligned to DFP phased delivery",
            "Decision log — open items flagged for VP ratification where needed",
        ],
    )
    add_body(doc, "Optional: 30-minute readout the week of July 28.")

    add_heading(doc, "Why now")
    add_body(
        doc,
        "You now lead both organizations. Without explicit alignment before build cycles accelerate, we risk:",
    )
    add_bullets(
        doc,
        [
            "Duplicate UI/metadata work across OMS unification, Navigator/Portal, and DFP agent shell",
            "Competing governance definitions — DFP Process Repository vs Gov policy/classification programs",
            "Blurred Alation-exit scope between OMS technical catalog and business glossary/dictionary",
            "Agent-first DFP requiring machine-readable rules — unclear content vs orchestration ownership",
        ],
    )
    add_body(
        doc,
        "This is structured alignment under unified leadership — not a turf negotiation — to ship "
        "one coherent experience for stewards, privacy leads, and data practitioners.",
    )

    add_heading(doc, "Organizing principle")
    add_body(
        doc,
        "One org, two jobs-to-be-done, one metadata spine:",
    )
    add_bullets(
        doc,
        [
            "Gov Portal — Are we in policy? What needs a human this week?",
            "Navigator — Discover and steward data assets",
            "DFP / Platform UIs + Agent — Operate the platform; execute governed actions",
        ],
    )
    add_body(
        doc,
        "Shared spine: Governance Metadata Repository ↔ OMS Semantic Registry ↔ Process Repository. "
        "Logically separate experiences; one product family (shared design language, deep links, same asset IDs).",
    )

    add_heading(doc, "Decisions we will bring back for your ratification")
    add_bullets(
        doc,
        [
            "Portal IA: two products in one shared shell vs merged experience",
            "System of record: business metrics/definitions (Gov) vs technical semantics (OMS)",
            "Engineering ownership: which squad builds and supports each in-scope capability",
            "Escalation rule if teams deadlock (Gov owns business meaning; DFP owns platform orchestration)",
        ],
    )

    add_heading(doc, "Co-sponsors")
    add_body(
        doc,
        "Data Governance: Niru Sharma, Jonathan Fong · "
        "Data Foundations: Ryan Malec, Renat Gilfanov",
    )
    add_body(
        doc,
        "Full workshop brief: https://docs.google.com/document/d/1tvA_XJnfHsVTteVUY32gnWrtueK81t630zC9VZiPtZA/edit",
    )

    doc.save(OUT)
    return OUT


if __name__ == "__main__":
    path = build()
    print(path)
