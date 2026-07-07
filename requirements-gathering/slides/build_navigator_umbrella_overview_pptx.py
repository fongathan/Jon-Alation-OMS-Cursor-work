#!/usr/bin/env python3
"""Build PowerPoint deck mirroring navigator-umbrella-overview-deck.html."""
from __future__ import annotations

import re
from datetime import date
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

OUT = Path(__file__).resolve().parent / "Navigator-Umbrella-Overview-Deck.pptx"

# Match HTML deck palette
NAVY = RGBColor(0x0F, 0x27, 0x44)
HEADER_BLUE = RGBColor(0x1E, 0x5A, 0x96)
SIDEBAR_BG = RGBColor(0xE4, 0xEE, 0xF8)
FOOTER_BG = RGBColor(0x3D, 0x51, 0x66)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
TEXT = RGBColor(0x1A, 0x1A, 0x1A)
MUTED_FOOTER = RGBColor(0xC5, 0xD4, 0xE3)
LOGO_FOOTER = RGBColor(0x9E, 0xB0, 0xC4)
CALLOUT_BG = RGBColor(0xE8, 0xF1, 0xF8)
CALLOUT_STRONG_BG = RGBColor(0xED, 0xF0, 0xF4)
BORDER = RGBColor(0xD0, 0xD8, 0xE2)

SW = Inches(13.333)
SH = Inches(7.5)
HDR_H = Inches(0.68)
FTR_H = Inches(0.4)
BODY_TOP = HDR_H
BODY_H = SH - HDR_H - FTR_H
ASIDE_W = Inches(5.05)
MAIN_LEFT = ASIDE_W
MAIN_W = SW - ASIDE_W
COL3_W = SW / 3
PAD = Inches(0.2)
FONT = "Calibri"

# Typography — tuned for presentation readability
SZ_SLIDE_TITLE = 22
SZ_SLIDE_NUM = 9
SZ_FOOTER = 9
SZ_FOOTER_BADGE = 8
SZ_HEADING = 12
SZ_BODY = 12
SZ_BULLET = 12
SZ_CALLOUT = 11
SZ_CALLOUT_STRONG = 12
SZ_CARD_TITLE = 12
SZ_CARD_BODY = 11


def fill(shape, color: RGBColor) -> None:
    shape.fill.solid()
    shape.fill.fore_color.rgb = color


def no_line(shape) -> None:
    shape.line.fill.background()


def tb(slide, left, top, width, height):
    return slide.shapes.add_textbox(left, top, width, height)


def run_style(run, *, size=11, bold=False, color=TEXT, italic=False):
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color


def add_rich_para(text_frame, text: str, *, size=SZ_BODY, color=TEXT, space_after=5, level=0):
    """Parse **bold** markers into runs."""
    p = text_frame.paragraphs[0] if not text_frame.text else text_frame.add_paragraph()
    p.level = level
    p.space_after = Pt(space_after)
    p.line_spacing = 1.2
    parts = re.split(r"(\*\*[^*]+\*\*)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            r = p.add_run()
            r.text = part[2:-2]
            run_style(r, size=size, bold=True, color=color)
        else:
            r = p.add_run()
            r.text = part
            run_style(r, size=size, color=color)
    return p


def add_bullets(slide, left, top, width, height, items: list[str], *, size=SZ_BULLET):
    box = tb(slide, left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    tf.clear()
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(6)
        p.line_spacing = 1.18
        lead = p.add_run()
        lead.text = "• "
        run_style(lead, size=size, color=TEXT)
        parts = re.split(r"(\*\*[^*]+\*\*)", item)
        for part in parts:
            if not part:
                continue
            r = p.add_run()
            if part.startswith("**") and part.endswith("**"):
                r.text = part[2:-2]
                run_style(r, size=size, bold=True, color=TEXT)
            else:
                r.text = part
                run_style(r, size=size, color=TEXT)
    return box


def add_heading(slide, left, top, width, text: str, *, main=False):
    box = tb(slide, left, top, width, Inches(0.32))
    tf = box.text_frame
    p = tf.paragraphs[0]
    p.text = text.upper()
    run_style(p.runs[0], size=SZ_HEADING, bold=True, color=HEADER_BLUE if main else NAVY)
    return box


def add_para(slide, left, top, width, height, text: str, *, size=SZ_BODY):
    box = tb(slide, left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    add_rich_para(tf, text, size=size)
    return box


def add_callout(slide, left, top, width, height, text: str, *, strong=False):
    bg = CALLOUT_STRONG_BG if strong else CALLOUT_BG
    border = NAVY if strong else HEADER_BLUE
    rect = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    fill(rect, bg)
    no_line(rect)
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, Inches(0.06), height)
    fill(bar, border)
    no_line(bar)
    add_para(
        slide,
        left + Inches(0.14),
        top + Inches(0.06),
        width - Inches(0.2),
        height - Inches(0.1),
        text,
        size=SZ_CALLOUT_STRONG if strong else SZ_CALLOUT,
    )


def add_vocab_card(slide, left, top, width, title: str, body: str):
    h = Inches(0.9)
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, h)
    fill(card, WHITE)
    card.line.color.rgb = BORDER
    card.line.width = Pt(0.75)
    add_para(slide, left + Inches(0.1), top + Inches(0.06), width - Inches(0.15), Inches(0.24), f"**{title}**", size=SZ_CARD_TITLE)
    add_para(slide, left + Inches(0.1), top + Inches(0.3), width - Inches(0.15), Inches(0.52), body, size=SZ_CARD_BODY)
    return h


def slide_frame(slide, title: str, num: int, total: int = 6):
    # white slide base
    base = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, BODY_TOP, SW, BODY_H)
    fill(base, WHITE)
    no_line(base)

    # header
    hdr = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SW, HDR_H)
    fill(hdr, HEADER_BLUE)
    no_line(hdr)
    # navy overlay left portion for gradient feel
    grad = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SW * 0.45, HDR_H)
    fill(grad, NAVY)
    no_line(grad)

    tbox = tb(slide, Inches(0.35), Inches(0.1), Inches(10.8), Inches(0.45))
    p = tbox.text_frame.paragraphs[0]
    p.text = title
    run_style(p.runs[0], size=SZ_SLIDE_TITLE, bold=True, color=WHITE)

    nbox = tb(slide, Inches(11.2), Inches(0.18), Inches(1.9), Inches(0.3))
    p = nbox.text_frame.paragraphs[0]
    p.text = f"SLIDE {num} OF {total}"
    p.alignment = PP_ALIGN.RIGHT
    run_style(p.runs[0], size=SZ_SLIDE_NUM, bold=True, color=WHITE)

    # footer
    ftr = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, SH - FTR_H, SW, FTR_H)
    fill(ftr, FOOTER_BG)
    no_line(ftr)
    lbox = tb(slide, Inches(0.35), SH - Inches(0.32), Inches(7), Inches(0.25))
    p = lbox.text_frame.paragraphs[0]
    p.text = "Data Governance · Navigator product family"
    run_style(p.runs[0], size=SZ_FOOTER, color=MUTED_FOOTER)
    rbox = tb(slide, Inches(9.5), SH - Inches(0.32), Inches(3.5), Inches(0.25))
    p = rbox.text_frame.paragraphs[0]
    p.text = "INTERNAL — PROGRAM & LEADERSHIP"
    p.alignment = PP_ALIGN.RIGHT
    run_style(p.runs[0], size=SZ_FOOTER_BADGE, bold=True, color=LOGO_FOOTER)


def aside_panel(slide):
    panel = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, BODY_TOP, ASIDE_W, BODY_H)
    fill(panel, SIDEBAR_BG)
    no_line(panel)
    divider = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, ASIDE_W - Inches(0.01), BODY_TOP, Inches(0.01), BODY_H)
    fill(divider, BORDER)
    no_line(divider)


def col3_panel(slide, index: int):
    left = COL3_W * index
    if index < 2:
        div = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left + COL3_W - Inches(0.01), BODY_TOP, Inches(0.01), BODY_H)
        fill(div, BORDER)
        no_line(div)
    if index == 0:
        panel = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, BODY_TOP, COL3_W, BODY_H)
        fill(panel, SIDEBAR_BG)
        no_line(panel)


def slide_1(slide):
    slide_frame(slide, "Navigator — the collective data governance product family", 1)
    aside_panel(slide)
    x0, x1 = PAD, ASIDE_W + PAD
    y = BODY_TOP + PAD
    w0, w1 = ASIDE_W - PAD * 2, MAIN_W - PAD * 2

    add_heading(slide, x0, y, w0, "What Navigator is")
    y += Inches(0.32)
    add_callout(
        slide, x0, y, w0, Inches(0.78),
        "**Navigator** is the umbrella name for DEEPT's data governance products, programs, catalog experience, and operations — one coherent story instead of disconnected tools.",
        strong=True,
    )
    y += Inches(0.86)
    add_para(slide, x0, y, w0, Inches(0.32), "It reframes what we previously called separate initiatives:")
    y += Inches(0.34)
    add_bullets(
        slide, x0, y, w0, Inches(1.25),
        [
            "**Governance Portal** → **Navigator** (program command center)",
            "**Business Data Navigator** → **Business Data Catalog** (catalog on OMS)",
            "**PULS + Access Manager + Gov backend** → **Policy Drive Governance Framework**",
            "**Kronos** → business inventory (linked, not duplicated)",
        ],
    )
    y += Inches(1.28)
    add_callout(
        slide, x0, y, w0, Inches(0.68),
        "**Design principle:** Unified experience, federated execution — one portal for orientation and posture; authoritative systems stay authoritative.",
    )

    add_heading(slide, x1, BODY_TOP + PAD, w1, "What it does for the organization", main=True)
    add_bullets(
        slide, x1, BODY_TOP + PAD + Inches(0.38), w1, BODY_H - Inches(0.52),
        [
            "**Answers two questions** — \"Where is my program work?\" (Navigator portal) and \"Where is the governed data?\" (Business Data Catalog).",
            "**Governance Overview** — Privacy, Platform, Security, and Regulatory posture in one leadership-ready view.",
            "**Dimension routing** — Consent, DSRs, classifications, policies, access, GIS, regulations, and inventory each have a clear home — with deep links where work lives elsewhere.",
            "**Catalog continuity** — Stewards and analysts discover and curate metadata on OMS through BDC as Alation exits.",
            "**Policy truth in PULS** — Classifications, policies, and metrics stay in a governance repository — not scattered across spreadsheets.",
            "**Inventory via Kronos** — Regulatory dimension links to Kronos for business inventory; Navigator does not fork that system of record.",
            "**Clear platform boundaries** — OMS is the engine; DG owns business-facing shells; SISU partners on governance backend and inventory.",
        ],
    )


def slide_2(slide):
    slide_frame(slide, "Why Navigator — continuity, coherence, and clarity", 2)
    aside_panel(slide)
    x0, x1 = PAD, ASIDE_W + PAD
    y = BODY_TOP + PAD
    w0, w1 = ASIDE_W - PAD * 2, MAIN_W - PAD * 2

    add_heading(slide, x0, y, w0, "Why it matters now")
    y += Inches(0.32)
    add_bullets(
        slide, x0, y, w0, Inches(1.65),
        [
            "**Alation exit (Oct 2026)** — Teams need search, glossary, ownership, and tags on OMS — not ad hoc metadata in Confluence and Slack.",
            "**Fragmented governance** — PULS, Access Manager, Kronos, consent, DSR, and security tools work — but no single product narrative connects them.",
            "**OMS vs DG tension** — Foundations builds the store and APIs; DG must ship business value without building \"another OMS.\"",
            "**Dashboard sprawl** — Privacy and compliance stakeholders asked for a holistic unified experience — without rebuilding every backend.",
        ],
    )
    y += Inches(1.85)
    add_heading(slide, x0, y, w0, "What success requires")
    y += Inches(0.34)
    add_bullets(
        slide, x0, y, w0, Inches(1.05),
        [
            "**Continuity** — High-value Alation workflows on OMS for in-scope domains.",
            "**Coherence** — One portal connecting program dimensions and catalog assets.",
            "**Clarity** — Documented ownership: who builds shells vs who runs stores.",
        ],
    )

    add_heading(slide, x1, BODY_TOP + PAD, w1, "Stakeholder benefits", main=True)
    add_bullets(
        slide, x1, BODY_TOP + PAD + Inches(0.38), w1, Inches(3.35),
        [
            "**Privacy, legal & compliance** — One map of consent, DSR, tracker, regulatory, and incident posture — with drill-down to operational tools.",
            "**Stewards & DPMs** — Platform dimension links to BDC; classifications and policies tie to catalog assets in PULS.",
            "**Analysts & BI** — Business Data Catalog as daily discovery surface with glossary, ownership, and trust signals.",
            "**Executives & program leads** — Governance Overview indices, trends, and exportable evidence for council and audit requests.",
            "**Engineering & platform** — OMS remains technical metadata hub; Navigator consumes APIs — no duplicate ingestion in DG UI.",
            "**Inventory operators** — Kronos stays authoritative; Navigator provides one-click access from Regulatory → Inventory.",
        ],
    )
    add_callout(
        slide, x1, BODY_TOP + Inches(4.62), w1, Inches(0.72),
        "**FY27 framing:** Navigator is the **operating layer for data governance** — programs, catalog, policy, and inventory — not a replacement for every specialist system, but the shell that makes the whole program legible.",
    )


def slide_3(slide):
    slide_frame(slide, "Four parts under the Navigator umbrella", 3)
    aside_panel(slide)
    x0, x1 = PAD, ASIDE_W + PAD
    w0, w1 = ASIDE_W - PAD * 2, MAIN_W - PAD * 2

    add_heading(slide, x0, BODY_TOP + PAD, w0, "The umbrella model")
    add_para(
        slide, x0, BODY_TOP + PAD + Inches(0.35), w0, Inches(1.1),
        "**Navigator** (the name) is both the governance portal experience and the umbrella for the whole family. The four named parts are how we explain build scope, ownership, and architecture to executives and engineers.",
    )
    add_callout(
        slide, x0, BODY_TOP + Inches(1.72), w0, Inches(0.58),
        "Navigator orchestrates and connects — it does **not** collapse Kronos, PULS, and OMS into one database.",
    )

    add_heading(slide, x1, BODY_TOP + PAD, w1, "Named components", main=True)
    cards = [
        ("1. Navigator (portal + Gov App backend)", "Governance command center: Overview, Privacy / Platform / Security / Regulatory dimensions, workflow UIs, evidence export. Gov App backend read/updates policies and metrics in PULS."),
        ("2. Policy Drive Governance Framework", "PULS Postgres (classifications, policies, metrics) + Access Manager + Gov App backend. System of record for control metadata behind Navigator workflows."),
        ("3. Business Data Catalog (BDC)", "Business-first UI on OMS APIs — Data Sources, Catalog, Glossary, stewardship. Alation migration target; opened from left nav or Platform → Stewards."),
        ("4. Kronos (business inventory)", "SISU Postgres inventory + Kronos UI. Regulatory → Inventory deep link. Business context complements OMS technical metadata."),
    ]
    cy = BODY_TOP + PAD + Inches(0.38)
    for title, body in cards:
        h = add_vocab_card(slide, x1, cy, w1, title, body)
        cy += h + Inches(0.06)


def slide_4(slide):
    slide_frame(slide, "Navigator portal & Business Data Catalog", 4)
    aside_panel(slide)
    x0, x1 = PAD, ASIDE_W + PAD
    w0, w1 = ASIDE_W - PAD * 2, MAIN_W - PAD * 2

    add_heading(slide, x0, BODY_TOP + PAD, w0, "Navigator portal")
    add_para(slide, x0, BODY_TOP + PAD + Inches(0.34), w0, Inches(0.38), "Program operators land here. Four dimension buckets mirror how DEEPT runs governance:")
    add_bullets(
        slide, x0, BODY_TOP + PAD + Inches(0.72), w0, Inches(1.15),
        [
            "**Privacy** — Consent · DSRs · Tracker remediation",
            "**Platform** — Stewards → BDC · Classifications · Policies",
            "**Security** — Access · GIS · Data use",
            "**Regulatory** — Laws & regulations · Incidents · **Inventory → Kronos**",
        ],
    )
    add_para(
        slide, x0, BODY_TOP + Inches(1.95), w0, Inches(1.0),
        "Governance Overview aggregates PCI-style and section indices for leadership. Workflow UIs connect to Gov App backend; auditor view and evidence export support attestation cycles.",
    )

    add_heading(slide, x1, BODY_TOP + PAD, w1, "Business Data Catalog", main=True)
    add_para(slide, x1, BODY_TOP + PAD + Inches(0.34), w1, Inches(0.45), "The daily workspace for metadata consumers and stewards. Built **on top of OMS APIs** — not a second operational metadata store.")
    add_bullets(
        slide, x1, BODY_TOP + PAD + Inches(0.82), w1, Inches(1.45),
        [
            "**Discovery** — Data Sources browse, catalog search, glossary, asset detail, lineage (where OMS graph supports).",
            "**Stewardship** — Ownership, descriptions, tags/classifications display, suggest-edit flows.",
            "**Entry points** — Left nav \"Business Data Catalog\" or Navigator Platform → Stewards.",
            "**MVP focus** — Proven Alation workflow parity for in-scope domains before contract end.",
        ],
    )
    add_callout(
        slide, x1, BODY_TOP + Inches(2.38), w1, Inches(0.72),
        "**Ownership:** Data Governance owns BDC product & UX · OMS Foundations owns store, connectors & APIs · OMS is the engine; BDC is how the governance program experiences catalog value.",
    )


def slide_5(slide):
    slide_frame(slide, "Policy framework, Kronos & architecture", 5)
    aside_panel(slide)
    x0, x1 = PAD, ASIDE_W + PAD
    w0, w1 = ASIDE_W - PAD * 2, MAIN_W - PAD * 2

    add_heading(slide, x0, BODY_TOP + PAD, w0, "Policy Drive Governance Framework")
    add_para(slide, x0, BODY_TOP + PAD + Inches(0.34), w0, Inches(0.28), "Authoritative layer behind Navigator workflows:")
    add_bullets(
        slide, x0, BODY_TOP + PAD + Inches(0.62), w0, Inches(1.1),
        [
            "**PULS Postgres** — Classifications & tags, policies, metrics (SF relational SoR).",
            "**Access Manager** — Bidirectional read/update with classifications; enforcement alignment.",
            "**Gov App backend** — API for Navigator UIs; audit logging; no direct browser writes to PULS.",
        ],
    )
    add_heading(slide, x0, BODY_TOP + Inches(1.82), w0, "Kronos")
    add_bullets(
        slide, x0, BODY_TOP + Inches(2.14), w0, Inches(0.95),
        [
            "Business inventory on **SISU Postgres**.",
            "Edited in **Kronos** — linked from Navigator Regulatory dimension.",
            "Complements OMS tables/columns with business purpose and obligations.",
        ],
    )

    add_heading(slide, x1, BODY_TOP + PAD, w1, "Three platform zones", main=True)
    add_bullets(
        slide, x1, BODY_TOP + PAD + Inches(0.38), w1, Inches(1.45),
        [
            "**OMS** — Technical metadata hub: ingestion, Neptune graph, Snowflake/Databricks relationships, catalog APIs. BDC reads from OMS.",
            "**Governance stack** — PULS + Access Manager + Gov App backend. Navigator portal UIs read/write via backend.",
            "**SISU / Kronos** — Business inventory relational store + Kronos UI.",
        ],
    )
    add_callout(
        slide, x1, BODY_TOP + Inches(1.95), w1, Inches(0.82),
        "**Open design question — \"Read ?\":** Confirm whether OMS reads classifications, policies, and metrics directly from PULS Postgres or via Gov App backend API. Answer affects BDC classification display and schema ownership.",
    )
    add_callout(
        slide, x1, BODY_TOP + Inches(2.82), w1, Inches(0.48),
        "**Architecture diagram:** data-navigator-governance/governance-product-systems-architecture.drawio",
    )


def slide_6(slide):
    slide_frame(slide, "Who builds what — phasing — and how we measure success", 6)
    for i in range(3):
        col3_panel(slide, i)

    cols = [
        (
            "Who builds what",
            [
                "**Navigator portal & BDC** — Data Governance (product & UX)",
                "**OMS store & APIs** — OMS Foundations",
                "**PULS & Gov backend** — SISU partnership + DG program",
                "**Kronos inventory** — SISU Postgres + inventory program",
            ],
            "**Clear RACI** prevents Navigator becoming \"the OMS app\" and protects Alation exit scope.",
            False,
        ),
        (
            "Phased delivery",
            [
                "**Phase 1 (now)** — BDC MVP on OMS APIs; Navigator IA & wireframes; Kronos inventory link.",
                "**Phase 2** — Gov App backend + PULS paths; classifications & policy workflows in Navigator.",
                "**Phase 3** — Metrics governance UI; resolve OMS↔PULS read model; evidence automation.",
            ],
            "Front-load catalog continuity; sequence PULS without blocking BDC on OMS.",
            False,
        ),
        (
            "Success indicators",
            [
                "**Catalog continuity** — In-scope Alation workflows in BDC before contract end.",
                "**Navigation coherence** — Platform → BDC and Regulatory → Kronos work in one click.",
                "**Posture reporting** — Governance Overview metrics traceable to program backends.",
                "**Adoption** — Weekly active stewards, privacy, and compliance roles in Navigator family UIs.",
                "**Architecture clarity** — Single alignment story for Foundations, OMS, metadata WG, and DG.",
            ],
            "**Artifacts:** Wireframes · umbrella PRD · this deck · OMS alignment deck — keep naming and links in sync.",
            True,
        ),
    ]

    for i, (heading, bullets, callout, main) in enumerate(cols):
        left = COL3_W * i + PAD
        w = COL3_W - PAD * 2
        add_heading(slide, left, BODY_TOP + PAD, w, heading, main=main)
        add_bullets(slide, left, BODY_TOP + PAD + Inches(0.38), w, Inches(2.45), bullets)
        add_callout(slide, left, BODY_TOP + Inches(2.92), w, Inches(0.68), callout)


def build() -> Path:
    prs = Presentation()
    prs.slide_width = SW
    prs.slide_height = SH
    blank = prs.slide_layouts[6]

    for fn in (slide_1, slide_2, slide_3, slide_4, slide_5, slide_6):
        fn(prs.slides.add_slide(blank))

    prs.save(OUT)
    return OUT


if __name__ == "__main__":
    path = build()
    print(f"Wrote {path} ({date.today().isoformat()})")
