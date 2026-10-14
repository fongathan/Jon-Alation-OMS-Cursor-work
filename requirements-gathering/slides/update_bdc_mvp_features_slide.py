#!/usr/bin/env python3
"""Replace slide 3 in Jon - DG Initiatives Roadmap with MVP vs Beyond MVP Alation features."""

from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.util import Inches, Pt

DECK = Path(__file__).resolve().parent / "Jon - DG Initiatives Roadmap.pptx"
SLIDE_INDEX = 2  # slide 3 (0-based)

# Deck palette (from existing BDC roadmap slides)
BLUE = RGBColor(0x5F, 0xA5, 0xFF)
GREEN = RGBColor(0x61, 0xEB, 0xA7)
LIGHT = RGBColor(0xE2, 0xE9, 0xEC)
MUTED = RGBColor(0x6A, 0x80, 0x95)
TITLE = RGBColor(0xB8, 0xC7, 0xD0)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
PANEL = RGBColor(0x12, 0x1E, 0x2B)
PANEL_LINE = RGBColor(0x2A, 0x3D, 0x52)
FONT = "Montserrat"
FONT_MED = "Montserrat Medium"

MVP_FEATURES = [
    "Search & discovery — keyword search across tables, glossary, articles",
    "Business glossary — terms, definitions, synonyms, links to assets",
    "Catalog browse & asset detail — datasource → schema → table/column",
    "Tags, domains & custom fields — classification and domain scoping",
    "Trust flags — endorse, warn, deprecate",
    "Ownership & stewardship — business and technical stewards",
    "Data lineage — table-level impact and upstream/downstream context",
    "Documentation / articles & folder navigation",
    "Dataflows, bulk edit & dictionary import/export",
    "SSO, RBAC & connected Snowflake sources",
]

BEYOND_MVP_FEATURES = [
    "Data Product marketplace / product pages — curated product docs, metrics, stewards",
    "AI chat assistant & description suggestions — NL discovery beyond search",
    "Admin analytics dashboard — MAU, top searches, usage telemetry, CSV export",
    "Mass set rules / catalog sets — policy and ownership at scale",
    "Policy center — compliance rules, risk classification, governance workflows",
    "Expanded connectors — non-Snowflake / activation tools (e.g. HighTouch)",
    "Saved searches, favorites & social / Q&A on assets",
    "Persona-based catalog views & in-app help / AI how-to",
    "Automated metadata curation (626) — scan sources, provenance, steward review",
    "Product pages, semantic layer & enterprise governance shell (Navigator FY27+)",
]


def _delete_all_shapes(slide) -> None:
    sp_tree = slide.shapes._spTree
    for shape in list(slide.shapes):
        sp_tree.remove(shape._element)


def _add_panel(slide, left, top, width, height, accent: RGBColor):
    panel = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    panel.fill.solid()
    panel.fill.fore_color.rgb = PANEL
    panel.line.color.rgb = PANEL_LINE
    panel.line.width = Pt(1)
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, Inches(0.12), height)
    bar.fill.solid()
    bar.fill.fore_color.rgb = accent
    bar.line.fill.background()
    return panel


def _text_box(slide, left, top, width, height, paragraphs, anchor=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    first = True
    for spec in paragraphs:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.alignment = spec.get("align", PP_ALIGN.LEFT)
        p.space_after = spec.get("space_after", Pt(4))
        p.line_spacing = spec.get("line_spacing", 1.05)
        parts = spec.get("parts") or [(spec.get("text", ""), spec)]
        for text, style in parts:
            run = p.add_run()
            run.text = text
            f = run.font
            f.name = style.get("font", FONT_MED)
            f.size = style.get("size", Pt(14))
            f.bold = style.get("bold", False)
            f.italic = style.get("italic", False)
            if style.get("color"):
                f.color.rgb = style["color"]
    return box


def _bullet_paragraphs(items, color=LIGHT, size=Pt(13.5), marker="•  "):
    out = []
    for item in items:
        out.append(
            {
                "parts": [(marker, {"size": size, "color": color, "font": FONT_MED}), (item, {"size": size, "color": color, "font": FONT_MED})],
                "space_after": Pt(5),
                "line_spacing": 1.08,
            }
        )
    return out


def rebuild_slide_3(prs: Presentation) -> None:
    slide = prs.slides[SLIDE_INDEX]
    _delete_all_shapes(slide)
    sw, sh = prs.slide_width, prs.slide_height

    _text_box(
        slide,
        Inches(0.55),
        Inches(0.35),
        Inches(12.2),
        Inches(0.95),
        [
            {
                "text": "Alation Feature Landscape — MVP vs. Beyond",
                "size": Pt(28),
                "bold": True,
                "color": WHITE,
                "font": FONT,
                "space_after": Pt(2),
            },
            {
                "text": "We only used a subset of Alation today. Additional vendor capabilities are candidates to map into Navigator / BDC.",
                "size": Pt(14),
                "color": MUTED,
                "font": FONT_MED,
                "italic": True,
            },
        ],
    )

    col_top = Inches(1.45)
    col_h = Inches(5.55)
    col_w = Inches(5.95)
    gap = Inches(0.35)
    lx = Inches(0.55)
    rx = lx + col_w + gap

    _add_panel(slide, lx, col_top, col_w, col_h, BLUE)
    _text_box(
        slide,
        lx + Inches(0.28),
        col_top + Inches(0.18),
        col_w - Inches(0.45),
        Inches(0.9),
        [
            {"text": "MVP — In use today", "size": Pt(22), "bold": True, "color": BLUE, "font": FONT},
            {
                "text": "Must migrate for Alation exit · validated in feature inventory",
                "size": Pt(12.5),
                "color": MUTED,
                "font": FONT_MED,
            },
        ],
    )
    _text_box(
        slide,
        lx + Inches(0.28),
        col_top + Inches(1.05),
        col_w - Inches(0.45),
        col_h - Inches(1.2),
        _bullet_paragraphs(MVP_FEATURES),
    )

    _add_panel(slide, rx, col_top, col_w, col_h, GREEN)
    _text_box(
        slide,
        rx + Inches(0.28),
        col_top + Inches(0.18),
        col_w - Inches(0.45),
        Inches(0.9),
        [
            {"text": "Beyond MVP — Map & evaluate", "size": Pt(22), "bold": True, "color": GREEN, "font": FONT},
            {
                "text": "Licensed in Alation but lightly used or aspirational — roadmap input",
                "size": Pt(12.5),
                "color": MUTED,
                "font": FONT_MED,
            },
        ],
    )
    _text_box(
        slide,
        rx + Inches(0.28),
        col_top + Inches(1.05),
        col_w - Inches(0.45),
        col_h - Inches(1.2),
        _bullet_paragraphs(BEYOND_MVP_FEATURES, color=LIGHT),
    )

    _text_box(
        slide,
        Inches(0.55),
        Inches(7.05),
        Inches(12.2),
        Inches(0.35),
        [
            {
                "text": "Source: Alation feature inventory matrix · capability comparison chart · stakeholder interviews",
                "size": Pt(9.5),
                "color": MUTED,
                "font": FONT_MED,
                "align": PP_ALIGN.CENTER,
            }
        ],
        anchor=MSO_ANCHOR.MIDDLE,
    )


def main() -> None:
    prs = Presentation(DECK)
    rebuild_slide_3(prs)
    prs.save(DECK)
    print(f"Updated slide 3 in {DECK}")


if __name__ == "__main__":
    main()
