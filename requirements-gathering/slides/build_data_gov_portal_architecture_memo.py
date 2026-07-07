#!/usr/bin/env python3
"""Build Data Governance Portal architecture Q&A memo (Word) and deck (PowerPoint)."""
from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from docx.shared import Inches, Pt, RGBColor
from pptx import Presentation
from pptx.dml.color import RGBColor as PptxRGB
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches as PInches, Pt as PPt

OUT_DIR = Path(__file__).resolve().parent
MEMO_PATH = OUT_DIR / "Data-Gov-Portal-Architecture-QA-Memo.docx"
DECK_PATH = OUT_DIR / "Data-Gov-Portal-Architecture-QA-Deck.pptx"
SHOTS_DIR = OUT_DIR / "ui-walkthrough-shots"

NAVY = RGBColor(0x0F, 0x27, 0x44)
BLUE = RGBColor(0x15, 0x65, 0xC0)
MUTED = RGBColor(0x5C, 0x65, 0x70)
SLATE = RGBColor(0x33, 0x41, 0x55)

DEET_NAVY = PptxRGB(0x15, 0x23, 0x42)
DEET_BLUE = PptxRGB(0x1E, 0x40, 0xAF)
DEET_SLATE = PptxRGB(0x33, 0x41, 0x55)
DEET_MUTED = PptxRGB(0x47, 0x55, 0x69)
DEET_FOOT_MUTED = PptxRGB(0x64, 0x74, 0x8B)
DEET_TOP_BG = PptxRGB(0xEE, 0xF2, 0xFF)
DEET_FOOTER_BG = PptxRGB(0xF1, 0xF5, 0xF9)
DEET_WHITE = PptxRGB(0xFF, 0xFF, 0xFF)
DEET_BORDER = PptxRGB(0xE2, 0xE8, 0xF0)
DEET_GREEN = PptxRGB(0x15, 0x6F, 0x4A)
DEET_TEAL = PptxRGB(0x0E, 0x94, 0x9A)
DEET_AMBER = PptxRGB(0xB4, 0x53, 0x09)
DEET_SLIDE_W = PInches(13.333333333333334)
DEET_FOOTER = "DEET Data Governance  |  June 2026  |  Data Gov Portal Architecture Q&A"


def set_run(p, text, size=11, color=SLATE, bold=False, font="Calibri"):
    p.clear()
    run = p.add_run(text)
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.font.bold = bold
    run.font.name = font
    return p


def add_heading(doc, text, level=1):
    h = doc.add_heading(text, level=level)
    for run in h.runs:
        run.font.color.rgb = NAVY
        run.font.name = "Calibri Light" if level == 1 else "Calibri"
    return h


def add_bullets(doc, items, size=11):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        set_run(p, item, size=size)


def add_table(doc, headers, rows):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        set_run(hdr[i].paragraphs[0], h, size=10, color=RGBColor(0xFF, 0xFF, 0xFF), bold=True)
        shading = hdr[i]._tc.get_or_add_tcPr()
        from docx.oxml.ns import nsdecls
        from docx.oxml import parse_xml

        shading.append(parse_xml(f'<w:shd {nsdecls("w")} w:fill="0F2744"/>'))
    for r_idx, row in enumerate(rows):
        cells = table.rows[r_idx + 1].cells
        for c_idx, val in enumerate(row):
            set_run(cells[c_idx].paragraphs[0], val, size=10)
    doc.add_paragraph()
    return table


def build_memo() -> Path:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)

    title = doc.add_paragraph()
    title.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    set_run(title, "Data Governance Portal — Architecture Q&A", size=22, color=NAVY, bold=True, font="Calibri Light")

    sub = doc.add_paragraph()
    sub.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    set_run(
        sub,
        "Personas · Portal separation · Integration · Dev ownership\n"
        "Jonathan Fong · Data Governance · June 29, 2026 · INTERNAL",
        size=11,
        color=MUTED,
    )
    doc.add_paragraph()

    add_heading(doc, "Executive summary", 2)
    p = doc.add_paragraph()
    set_run(
        p,
        "The Governance Platform architecture shows a dedicated Governance Console & Applications layer "
        "above platform metadata stores and compliance services. Our program aligns with that model: "
        "one shared metadata spine (626 → OMS → governance repository), multiple purpose-built experiences. "
        "The Data Governance Portal is the command center for posture, exceptions, and evidence — not a "
        "replacement for specialist privacy, access, or inventory tools. Business Data Navigator (catalog) "
        "serves discovery and thin stewardship now; the full Governance Portal hub is phase 2 alongside "
        "the Alation exit (Oct 1, 2026).",
        size=11,
    )

    add_heading(doc, "1. Which personas does the Data Gov Portal serve?", 1)
    p = doc.add_paragraph()
    set_run(
        p,
        "Primary users operate governance programs — they need triage, attestation, and cross-pillar "
        "visibility. Secondary users consume governed metadata through Navigator. Engineers use the "
        "Foundations OMS UI, not the Portal.",
        size=11,
    )

    add_table(
        doc,
        ["Persona", "Portal role", "Primary surfaces (architecture)"],
        [
            (
                "Privacy & compliance officers",
                "Primary",
                "DSR Mgmt UI, Consent Mgmt UI, tracker remediation queues, compliance reporting",
            ),
            (
                "Governance program leads & executives",
                "Primary",
                "Governance Overview (PCI + section indices), metrics governance, evidence export",
            ),
            (
                "Data stewards / DPMs",
                "Primary (workflows)",
                "Classifications Mgmt UI, tag & policy workflows, classification review",
            ),
            (
                "Access managers & GIS / security",
                "Primary",
                "Access manager, data-use approvals, security dimension queues",
            ),
            (
                "Partner & API managers",
                "Primary",
                "Partner API Mgmt UI, partner data transfer framework handoffs",
            ),
            (
                "Records / inventory owners",
                "Primary",
                "Business Inventory Mgmt UI, DLM workflows, Kronos inventory link (Regulatory)",
            ),
            (
                "Internal audit & legal",
                "Primary (read/export)",
                "Evidence aggregation, auditor view, regulatory incident posture",
            ),
            (
                "Business & analytics consumers",
                "Indirect",
                "Navigator catalog (discovery, glossary, classifications read) — not Portal MVP",
            ),
            (
                "Data engineers / platform operators",
                "Not Portal",
                "Foundations OMS UI, lineage, technical metadata — deep link from Navigator",
            ),
        ],
    )

    p = doc.add_paragraph()
    set_run(
        p,
        "Organizing principle: Portal IA follows four dimensions — Privacy, Platform, Security, Regulatory — "
        "mapped to how DEEPT runs governance, not org-chart names.",
        size=11,
        bold=True,
    )

    add_heading(doc, "2. Does the Data Gov Portal need to be separate from the Data Foundation Platform?", 1)
    p = doc.add_paragraph()
    set_run(p, "Short answer: ", size=11, bold=True)
    set_run(
        p,
        "Logically yes — physically it should feel like one product family in a shared shell, not one "
        "monolithic app owned by Foundations.",
        size=11,
    )

    add_heading(doc, "Why separation matters", 2)
    add_bullets(
        doc,
        [
            "Different jobs to be done — Portal answers “Are we in policy? What needs a human this week?” "
            "DFP/OMS answers “What exists technically? How do pipelines connect?”",
            "Different systems of record — Governance Metadata Repository (Postgres) holds classifications, "
            "policies, and metrics definitions. OMS is the operational metadata store. Foundations should not "
            "own governance SoR or our program roadmap.",
            "Different personas — stewards and privacy leads are not daily OMS operators. Forcing one UI "
            "risks engineer-first density or governance features buried in platform chrome.",
            "Political independence — leadership asked for a governance program that is tool-agnostic and not "
            "captive to the Foundations OMS UI overhaul (same surface as Alation migration).",
            "Federated execution — each workstream (consent, DSR, tracker, access, GIS) keeps its authoritative "
            "tool; the Portal unifies posture and routing without rebuilding backends.",
        ],
    )

    add_heading(doc, "What “separate” does not mean", 2)
    add_bullets(
        doc,
        [
            "Not a second metadata store — one spine: 626/MCI supplies, OMS stores, governance repo defines.",
            "Not five disconnected URLs forever — shared shell, global search, deep links, same asset IDs.",
            "Not duplicating Navigator — Portal = govern & prove; Navigator = discover & steward.",
        ],
    )

    add_heading(doc, "3. How can they be incorporated?", 1)
    add_heading(doc, "Recommended pattern: dual product, single spine", 2)
    add_bullets(
        doc,
        [
            "Shared OMS shell — Governance Portal, Navigator, and specialist tools feel like one navigation "
            "story with role-based landing pages.",
            "Deep links — Portal exception → Navigator asset detail → “Open in OMS” for full technical context.",
            "API boundaries — Navigator and Portal UIs call Gov Application Services and OMS APIs; browsers "
            "do not write directly to Postgres governance repo.",
            "OMS projection — OMS reads/syncs governance definitions onto catalog assets; governance repo "
            "stays tool-agnostic.",
            "Phased delivery — Phase 1 (now): Navigator/BDC MVP on OMS APIs for Alation exit. Phase 2: "
            "Governance Portal hub with live KPI feeds and work queues. Phase 3: metrics governance UI, "
            "evidence automation, resolved OMS↔governance read model.",
        ],
    )

    add_heading(doc, "Architecture layers (from capability diagram)", 2)
    add_table(
        doc,
        ["Layer", "What it is", "How UI connects"],
        [
            (
                "Governance Console & Applications",
                "Portal, workflow UIs, management consoles",
                "Read/update via Gov Application Services",
            ),
            (
                "Gov Metadata Management Services",
                "Gov app backend, evidence capture, enterprise integrations",
                "APIs for Portal and Navigator workflow modules",
            ),
            (
                "Governance Metadata Repository",
                "Postgres — classifications, policies, taxonomies, business inventory defs",
                "SoR for governance definitions (not OMS)",
            ),
            (
                "OMS + platforms (SF, Databricks)",
                "Operational metadata, Horizon catalog, tagging enforcement",
                "Navigator reads OMS; engineers use Foundations OMS UI",
            ),
            (
                "Compliance & monitoring services",
                "Scanning, DSR orchestration, anonymization, conflict detection",
                "Portal aggregates signals; specialists work in downstream tools",
            ),
            (
                "Policy engines (SAME, DSR, RIM)",
                "Automated enforcement",
                "Read classifications, policies, and rules from governance services",
            ),
        ],
    )

    p = doc.add_paragraph()
    set_run(
        p,
        "Design principle: unified experience, federated execution — one front door for the story; "
        "specialist systems stay authoritative for deep work.",
        size=11,
        bold=True,
        color=BLUE,
    )

    add_heading(doc, "4. Which dev team builds which parts?", 1)
    add_table(
        doc,
        ["Component", "Owning team / partner", "Notes"],
        [
            (
                "Governance Portal & Navigator/BDC UX",
                "Data Governance (DEET)",
                "Product, UX, program delivery; not the metadata engine",
            ),
            (
                "OMS store, connectors, catalog APIs",
                "OMS Foundations (Data Foundation Platform)",
                "Operational metadata hub; Neptune graph; SF/DBX relationships",
            ),
            (
                "Foundations OMS UI (engineer path)",
                "OMS Foundations",
                "Technical browse, lineage, SQL — not a substitute for Navigator",
            ),
            (
                "Governance Metadata Repository (Postgres)",
                "SISU partnership + Data Governance program",
                "Classifications, policies, metrics definitions, taxonomies",
            ),
            (
                "Gov Application Services & workflow backend",
                "SISU partnership + Data Governance",
                "APIs for Portal UIs; audit logging; enterprise integrations",
            ),
            (
                "Access Manager",
                "Platform / access program (with governance alignment)",
                "Bidirectional read/update with classifications",
            ),
            (
                "Business inventory (Kronos UI)",
                "Inventory program on SISU Postgres",
                "Linked from Portal Regulatory dimension",
            ),
            (
                "Metadata supply (scanning, reconciliation, AI enrichment)",
                "Data 626 / MCI + metadata working group",
                "Feeds OMS; Navigator surfaces outcomes — does not own scrapers",
            ),
            (
                "Privacy pillar tools (consent, DSR, tracker)",
                "Privacy engineering & operations",
                "Portal deep-links and aggregates KPIs — does not replace",
            ),
            (
                "Compliance & monitoring services",
                "Compliance / security engineering",
                "Scanning, DSR orchestration, anonymization, evidence aggregation",
            ),
            (
                "Policy engines (SAME, DSR, RIM)",
                "Compliance / records programs",
                "Consume governance definitions; automated enforcement",
            ),
        ],
    )

    add_heading(doc, "RACI guardrails", 2)
    add_bullets(
        doc,
        [
            "When Foundations and Navigator scope collide → metadata working group + escalation (Suman + Charles).",
            "OMS UI unification ≠ Navigator — no duplicate Alation-migration UI in Foundations’ OMS overhaul.",
            "Open design question: confirm OMS reads governance definitions directly from Postgres or via Gov App API — affects classification display and schema ownership.",
        ],
    )

    add_heading(doc, "Decisions to confirm with leadership", 2)
    add_bullets(
        doc,
        [
            "Persona-first MVP scope locked for Navigator (business/analyst primary; Portal phase 2).",
            "Foundations boundary: platform store + engineer UI; Data Governance owns catalog program UX.",
            "Single architecture story across Foundations, OMS, metadata WG, and Data Governance.",
            "Navigation IA: prefer v2 governance-first shell with BDC nested under Platform (see section 5).",
        ],
    )

    add_heading(doc, "5. UI perspective — how it looks", 1)
    p = doc.add_paragraph()
    set_run(
        p,
        "Users should see one Navigator-branded product family with shared chrome (top bar, search, "
        "breadcrumbs, collapsible left rail). Three front doors share the same asset IDs and deep links — "
        "they do not share one mega-screen.",
        size=11,
    )

    add_heading(doc, "Shared shell (Data Governance product)", 2)
    add_table(
        doc,
        ["UI element", "What users see"],
        [
            ("Top bar", "Product breadcrumb, global search, role context"),
            ("Left rail", "Mode switch or governance dimensions (depends on IA version)"),
            ("Main panel", "Catalog browse OR governance overview OR specialist handoff"),
        ],
    )

    add_heading(doc, "Three front doors", 2)
    add_table(
        doc,
        ["Door", "Default user", "Main content", "Prototype"],
        [
            (
                "Business Data Catalog",
                "Analysts, business consumers, stewards",
                "Search, catalog, glossary, business-first asset detail, Open in OMS",
                "data-navigator-governance · oms-ui-on-top.html",
            ),
            (
                "Data Governance Portal",
                "Privacy, compliance, execs, audit",
                "Governance Overview, four dimensions, work queues, evidence export",
                "data-navigator-governance · navigator-version-2",
            ),
            (
                "Foundations OMS UI",
                "Data engineers, platform operators",
                "Technical tabs: SQL, lineage graph, ops metadata",
                "oms-ui-integrated.html",
            ),
        ],
    )

    add_heading(doc, "Navigation IA — v1 vs v2", 2)
    add_table(
        doc,
        ["", "v1 — Two top-level tabs", "v2 — Governance-first (recommended)"],
        [
            ("Left nav", "Business Data Catalog | Data Governance Portal", "Governance Overview + Privacy / Platform / Security / Regulatory sections"),
            ("Catalog entry", "Top-level tab", "Platform → Business Data Catalog (nested shell)"),
            ("Best for", "Clear mode switch; Alation-exit MVP", "Privacy/compliance default landing; exec posture first"),
            ("Risk", "Users may live only in catalog tab", "Catalog one click deeper — mitigated by search + Platform CTA"),
        ],
    )
    p = doc.add_paragraph()
    set_run(
        p,
        "Recommendation: ship v2 governance-first IA for the Portal phase-2 hub; keep v1 two-tab pattern "
        "as the simpler MVP path for catalog-only users until Portal launches. Both use the same shell and APIs.",
        size=11,
        bold=True,
        color=BLUE,
    )

    add_heading(doc, "Example user journeys (UI)", 2)
    add_bullets(
        doc,
        [
            "Privacy officer: Governance Overview → Privacy queue row → open DSR/MAUI with context preserved.",
            "Analyst: Business Data Catalog → search → asset detail (no governance dashboard unless entitled).",
            "Steward: Catalog for descriptions; Platform → Classifications for review workflows.",
            "Engineer: OMS UI directly, or Open in OMS from a catalog asset.",
        ],
    )

    p = doc.add_paragraph()
    set_run(
        p,
        "Screenshots for stakeholder walkthrough: requirements-gathering/slides/ui-walkthrough-shots/",
        size=10,
        color=MUTED,
    )

    foot = doc.add_paragraph()
    foot.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    set_run(foot, "INTERNAL — Data Governance Program · References: UDGE, Navigator intake, governance capability architecture", size=9, color=MUTED)

    doc.save(MEMO_PATH)
    return MEMO_PATH


# --- PowerPoint helpers (DEET style) -----------------------------------------

def deet_para(p, text, size=11, color=DEET_SLATE, bold=False, font="Calibri", align=None):
    p.text = text
    p.font.size = PPt(size)
    p.font.color.rgb = color
    p.font.bold = bold
    p.font.name = font
    if align is not None:
        p.alignment = align


def deet_send_back(slide, shape):
    sp_tree = slide.shapes._spTree  # noqa: SLF001
    el = shape._element  # noqa: SLF001
    sp_tree.remove(el)
    sp_tree.insert(2, el)


def deet_rect(slide, left, top, width, height, fill, line=None, radius=False):
    kind = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(kind, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = line if line else fill
    if not line:
        shape.line.fill.background()
    return shape


def deet_footer(slide):
    deet_rect(slide, PInches(0), PInches(7.0), DEET_SLIDE_W, PInches(0.5), DEET_FOOTER_BG)
    tb = slide.shapes.add_textbox(PInches(0.5), PInches(7.08), PInches(12.33), PInches(0.34))
    deet_para(tb.text_frame.paragraphs[0], DEET_FOOTER, size=9, color=DEET_FOOT_MUTED)


def deet_title_slide(prs, title, subtitle):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    top = deet_rect(slide, PInches(0), PInches(0), DEET_SLIDE_W, PInches(2.4), DEET_TOP_BG)
    deet_send_back(slide, top)
    bar = deet_rect(slide, PInches(0), PInches(2.4), DEET_SLIDE_W, PInches(0.06), DEET_NAVY)
    deet_send_back(slide, bar)
    tb = slide.shapes.add_textbox(PInches(0.7), PInches(2.9), PInches(11.9), PInches(1.4))
    deet_para(tb.text_frame.paragraphs[0], title, size=36, color=DEET_NAVY, font="Calibri Light")
    st = slide.shapes.add_textbox(PInches(0.7), PInches(4.05), PInches(11.9), PInches(0.9))
    deet_para(st.text_frame.paragraphs[0], subtitle, size=16, color=DEET_MUTED)
    deet_footer(slide)
    return slide


def deet_content_slide(prs, title, subtitle=None):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    header = deet_rect(slide, PInches(0), PInches(0), DEET_SLIDE_W, PInches(0.55), DEET_TOP_BG)
    deet_send_back(slide, header)
    tb = slide.shapes.add_textbox(PInches(0.5), PInches(0.88), PInches(12.5), PInches(0.44))
    deet_para(tb.text_frame.paragraphs[0], title, size=24, color=DEET_NAVY, font="Calibri Light")
    top = PInches(1.36)
    if subtitle:
        st = slide.shapes.add_textbox(PInches(0.5), top, PInches(12.33), PInches(0.3))
        deet_para(st.text_frame.paragraphs[0], subtitle, size=12, color=DEET_MUTED)
        top = PInches(1.72)
    deet_footer(slide)
    return slide, top


def deet_bullets(slide, left, top, width, height, items, size=11):
    tb = slide.shapes.add_textbox(left, top, width, height)
    tf = tb.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = f"• {item}"
        p.font.size = PPt(size)
        p.font.color.rgb = DEET_SLATE
        p.font.name = "Calibri"
        p.space_before = PPt(4)


def deet_callout(slide, top, text, height=PInches(0.58)):
    deet_rect(slide, PInches(0.5), top, PInches(12.33), height, PptxRGB(0xE0, 0xE7, 0xFF), line=DEET_BLUE, radius=True)
    tb = slide.shapes.add_textbox(PInches(0.65), top + PInches(0.08), PInches(12.0), height - PInches(0.12))
    tf = tb.text_frame
    tf.word_wrap = True
    deet_para(tf.paragraphs[0], text, size=12, color=DEET_NAVY, bold=True, font="Calibri Light")


def deet_image_card(slide, left, top, width, height, image_path: Path, caption: str, accent=DEET_BLUE):
    deet_rect(slide, left, top, width, height, DEET_WHITE, line=accent, radius=True)
    cap_h = PInches(0.42)
    img_h = height - cap_h - PInches(0.12)
    if image_path.exists():
        slide.shapes.add_picture(
            str(image_path),
            left + PInches(0.08),
            top + PInches(0.08),
            width=width - PInches(0.16),
            height=img_h,
        )
    else:
        ph = deet_rect(slide, left + PInches(0.08), top + PInches(0.08), width - PInches(0.16), img_h, DEET_FOOTER_BG, line=DEET_BORDER)
        deet_para(ph.text_frame.paragraphs[0], "Screenshot pending", size=9, color=DEET_MUTED, align=PP_ALIGN.CENTER)
    tb = slide.shapes.add_textbox(left + PInches(0.1), top + height - cap_h, width - PInches(0.2), cap_h)
    tf = tb.text_frame
    tf.word_wrap = True
    deet_para(tf.paragraphs[0], caption, size=9, color=DEET_NAVY, bold=True, font="Calibri Light")


def deet_mini_wireframe(slide, left, top, width, height):
    """Simple three-door wireframe for UI walkthrough slide."""
    deet_rect(slide, left, top, width, height, DEET_WHITE, line=DEET_BORDER, radius=True)
    bar_h = PInches(0.28)
    deet_rect(slide, left + PInches(0.1), top + PInches(0.1), width - PInches(0.2), bar_h, DEET_TOP_BG, line=DEET_BORDER)
    tb = slide.shapes.add_textbox(left + PInches(0.15), top + PInches(0.14), width - PInches(0.3), PInches(0.2))
    deet_para(tb.text_frame.paragraphs[0], "Navigator  /  Search  /  User", size=7, color=DEET_MUTED)
    rail_w = PInches(0.72)
    main_left = left + PInches(0.1) + rail_w + PInches(0.08)
    main_w = width - rail_w - PInches(0.28)
    body_top = top + PInches(0.48)
    body_h = height - PInches(0.58)
    deet_rect(slide, left + PInches(0.1), body_top, rail_w, body_h, PptxRGB(0xF8, 0xFA, 0xFC), line=DEET_BORDER)
    doors = [
        ("BDC", PptxRGB(0xDB, 0xEA, 0xFE), DEET_BLUE),
        ("Portal", PptxRGB(0xCC, 0xFB, 0xF1), DEET_TEAL),
        ("OMS →", PptxRGB(0xFE, 0xF3, 0xC7), DEET_AMBER),
    ]
    for i, (label, fill, accent) in enumerate(doors):
        y = body_top + PInches(0.08) + i * PInches(0.55)
        box = deet_rect(slide, left + PInches(0.16), y, rail_w - PInches(0.12), PInches(0.46), fill, line=accent, radius=True)
        deet_para(box.text_frame.paragraphs[0], label, size=7, color=DEET_NAVY, bold=True, align=PP_ALIGN.CENTER)
    deet_rect(slide, main_left, body_top, main_w, body_h, DEET_WHITE, line=DEET_BORDER)
    tb2 = slide.shapes.add_textbox(main_left + PInches(0.1), body_top + PInches(0.12), main_w - PInches(0.2), PInches(0.5))
    deet_para(tb2.text_frame.paragraphs[0], "Main content swaps by mode", size=8, color=DEET_SLATE, bold=True)
    tb3 = slide.shapes.add_textbox(main_left + PInches(0.1), body_top + PInches(0.55), main_w - PInches(0.2), body_h - PInches(0.65))
    tf = tb3.text_frame
    tf.word_wrap = True
    for line in ["Catalog grid / asset detail", "— or —", "Governance Overview + queues", "— or deep link —", "OMS technical tabs"]:
        p = tf.paragraphs[0] if line == "Catalog grid / asset detail" else tf.add_paragraph()
        p.text = line
        p.font.size = PPt(7)
        p.font.color.rgb = DEET_MUTED
        p.font.name = "Calibri"


def build_ui_walkthrough_slides(prs: Presentation) -> None:
    slide, top = deet_content_slide(
        prs,
        "UI walkthrough — three doors, one shell",
        "Same chrome and asset IDs · different default landing by persona · specialist tools open on demand",
    )
    deet_mini_wireframe(slide, PInches(0.5), top, PInches(2.55), PInches(4.35))
    cards = [
        (SHOTS_DIR / "bdc-catalog-mode.png", "Business Data Catalog\nDiscover & steward · Catalog · Glossary", DEET_BLUE),
        (SHOTS_DIR / "v2-governance-overview.png", "Data Governance Portal\nGovern & prove · Overview · Queues", DEET_TEAL),
        (SHOTS_DIR / "oms-engineer-ui.png", "Foundations OMS UI\nEngineer path · SQL · Lineage", DEET_AMBER),
    ]
    cw, ch = PInches(3.15), PInches(4.35)
    for i, (img, caption, accent) in enumerate(cards):
        left = PInches(3.25) + i * PInches(3.28)
        deet_image_card(slide, left, top, cw, ch, img, caption, accent)
    deet_callout(
        slide,
        PInches(6.05),
        "Deep link chain: Portal exception → Catalog asset → Open in OMS (one asset ID end to end)",
        height=PInches(0.48),
    )

    slide, top = deet_content_slide(
        prs,
        "Navigation IA — v1 vs v2 (recommendation)",
        "Both use the same shell and APIs · choice is default landing and left-rail structure",
    )
    col_w = PInches(5.95)
    variants = [
        (
            "v1 — Two top-level tabs",
            "Simpler MVP · clear mode switch",
            DEET_BLUE,
            [
                "Left rail: Business Data Catalog | Data Governance Portal",
                "Click tab → entire main panel swaps",
                "Best for: catalog-first users during Alation exit",
                "Prototype: data-navigator-governance/index.html",
            ],
            SHOTS_DIR / "v1-governance-portal.png",
        ),
        (
            "v2 — Governance-first (recommended for Portal)",
            "Privacy/compliance default · exec posture first",
            DEET_TEAL,
            [
                "Left rail: Overview + Privacy / Platform / Security / Regulatory",
                "BDC under Platform → opens nested catalog shell",
                "Best for: governance program operators & leadership",
                "Prototype: navigator-version-2/index.html",
            ],
            SHOTS_DIR / "v2-governance-overview.png",
        ),
    ]
    for i, (title, subtitle, accent, bullets, img) in enumerate(variants):
        left = PInches(0.5) + i * PInches(6.35)
        deet_rect(slide, left, top, col_w, PInches(4.55), DEET_WHITE, line=accent, radius=True)
        tb = slide.shapes.add_textbox(left + PInches(0.18), top + PInches(0.12), col_w - PInches(0.36), PInches(0.55))
        tf = tb.text_frame
        deet_para(tf.paragraphs[0], title, size=14, color=DEET_NAVY, bold=True, font="Calibri Light")
        deet_para(tf.add_paragraph(), subtitle, size=10, color=DEET_MUTED)
        deet_bullets(slide, left + PInches(0.18), top + PInches(0.72), col_w - PInches(0.36), PInches(1.55), bullets, size=9)
        if img.exists():
            slide.shapes.add_picture(
                str(img),
                left + PInches(0.18),
                top + PInches(2.35),
                width=col_w - PInches(0.36),
                height=PInches(2.0),
            )
    deet_callout(
        slide,
        PInches(6.05),
        "Recommendation: v1 for Navigator MVP (catalog) · v2 for Governance Portal phase 2 · shared search, breadcrumbs, and deep links in both",
        height=PInches(0.55),
    )


def build_deck() -> Path:
    prs = Presentation()
    prs.slide_width = DEET_SLIDE_W
    prs.slide_height = PInches(7.5)

    deet_title_slide(
        prs,
        "Data Governance Portal — Architecture Q&A",
        "Personas · Separation from DFP · Integration · Dev ownership\nJune 29, 2026 · INTERNAL",
    )

    slide, top = deet_content_slide(
        prs,
        "1 — Which personas does the Portal serve?",
        "Primary = govern, triage, attest. Secondary = discover via Navigator. Engineers → OMS UI.",
    )
    personas = [
        ("Privacy & compliance", "DSR, consent, tracker queues, compliance reporting", DEET_TEAL),
        ("Governance leads & execs", "Overview indices, metrics gov, evidence export", DEET_BLUE),
        ("Data stewards / DPMs", "Classifications, tag & policy workflows, review", DEET_GREEN),
        ("Access & GIS / security", "Access manager, data-use approvals", PptxRGB(0xF7, 0x96, 0x46)),
        ("Partner & inventory owners", "Partner API mgmt, business inventory, DLM", PptxRGB(0x80, 0x64, 0xA2)),
        ("Audit & legal", "Evidence packages, auditor view", DEET_AMBER),
    ]
    cw, ch = PInches(3.95), PInches(1.15)
    for i, (name, detail, accent) in enumerate(personas):
        col, row = i % 3, i // 3
        left = PInches(0.5) + col * PInches(4.15)
        y = top + row * PInches(1.35)
        deet_rect(slide, left, y, cw, ch, DEET_WHITE, line=accent, radius=True)
        tb = slide.shapes.add_textbox(left + PInches(0.15), y + PInches(0.1), cw - PInches(0.25), ch - PInches(0.15))
        tf = tb.text_frame
        tf.word_wrap = True
        deet_para(tf.paragraphs[0], name, size=13, color=DEET_NAVY, bold=True, font="Calibri Light")
        deet_para(tf.add_paragraph(), detail, size=9, color=DEET_MUTED)
    deet_callout(
        slide,
        top + PInches(2.85),
        "Not primary Portal users: business/analyst consumers (Navigator MVP) · data engineers (Foundations OMS UI)",
        height=PInches(0.5),
    )
    deet_callout(
        slide,
        PInches(5.55),
        "Portal IA = Privacy · Platform · Security · Regulatory — how DEEPT runs governance",
        height=PInches(0.5),
    )

    slide, top = deet_content_slide(
        prs,
        "2 — Separate from Data Foundation Platform?",
        "Logically yes. One product family in a shared shell — not one monolith owned by Foundations.",
    )
    deet_bullets(
        slide,
        PInches(0.5),
        top,
        PInches(5.9),
        PInches(4.5),
        [
            "Different jobs — posture & triage vs technical metadata & pipelines",
            "Different SoR — Governance Postgres vs OMS operational store",
            "Different personas — stewards ≠ daily OMS operators",
            "Program independence — not captive to Foundations OMS UI roadmap",
            "Federated execution — unify story; keep specialist tools authoritative",
        ],
        size=11,
    )
    deet_rect(slide, PInches(6.65), top, PInches(6.18), PInches(4.5), DEET_WHITE, line=DEET_BORDER, radius=True)
    tb = slide.shapes.add_textbox(PInches(6.85), top + PInches(0.15), PInches(5.8), PInches(4.2))
    tf = tb.text_frame
    tf.word_wrap = True
    deet_para(tf.paragraphs[0], "“Separate” does NOT mean", size=14, color=DEET_NAVY, bold=True, font="Calibri Light")
    for item in [
        "Second metadata store — one spine: 626 → OMS → governance repo",
        "Five URLs forever — shared shell, deep links, same asset IDs",
        "Duplicating Navigator — Portal = govern & prove; Navigator = discover & steward",
    ]:
        p = tf.add_paragraph()
        p.text = f"• {item}"
        p.font.size = PPt(10)
        p.font.color.rgb = DEET_SLATE
        p.space_before = PPt(6)
    deet_callout(slide, PInches(6.15), "Risk if merged: Navigator becomes “the OMS app” — Foundations owns our Alation exit scope")

    slide, top = deet_content_slide(
        prs,
        "3 — How to incorporate: dual product, single spine",
        "Unified experience · federated execution",
    )
    layers = [
        ("Governance Console", "Portal + workflow UIs", "Gov App Services"),
        ("Gov services", "Backend, evidence, integrations", "APIs for Portal & Navigator"),
        ("Governance repo", "Postgres — policies, tags, taxonomies", "SoR (not OMS)"),
        ("OMS + platforms", "Catalog APIs, SF/DBX, Horizon", "Navigator reads; engineers use OMS UI"),
        ("Compliance services", "Scanning, DSR, anonymization", "Portal aggregates; tools stay deep"),
    ]
    for i, (layer, desc, conn) in enumerate(layers):
        y = top + i * PInches(0.95)
        deet_rect(slide, PInches(0.5), y, PInches(2.2), PInches(0.78), DEET_NAVY, radius=True)
        tb = slide.shapes.add_textbox(PInches(0.6), y + PInches(0.12), PInches(2.0), PInches(0.55))
        deet_para(tb.text_frame.paragraphs[0], layer, size=10, color=DEET_WHITE, bold=True, align=PP_ALIGN.CENTER)
        deet_rect(slide, PInches(2.85), y, PInches(5.5), PInches(0.78), DEET_WHITE, line=DEET_BORDER, radius=True)
        tb2 = slide.shapes.add_textbox(PInches(3.0), y + PInches(0.1), PInches(5.2), PInches(0.6))
        deet_para(tb2.text_frame.paragraphs[0], desc, size=10, color=DEET_SLATE)
        deet_rect(slide, PInches(8.5), y, PInches(4.33), PInches(0.78), PptxRGB(0xE8, 0xF5, 0xE9), line=DEET_GREEN, radius=True)
        tb3 = slide.shapes.add_textbox(PInches(8.65), y + PInches(0.1), PInches(4.0), PInches(0.6))
        deet_para(tb3.text_frame.paragraphs[0], conn, size=9, color=DEET_GREEN, bold=True)
    deet_callout(
        slide,
        PInches(6.0),
        "Phasing: Now = Navigator/BDC MVP · Phase 2 = Portal hub · Phase 3 = metrics gov + evidence automation",
        height=PInches(0.5),
    )

    slide, top = deet_content_slide(
        prs,
        "4 — Which dev team builds which parts?",
        "Clear RACI protects Alation exit scope and prevents scope collision with Foundations",
    )
    teams = [
        ("Data Governance (DEET)", "Governance Portal, Navigator/BDC UX, program delivery"),
        ("OMS Foundations (DFP)", "OMS store, APIs, connectors, engineer OMS UI"),
        ("SISU + DG program", "Governance Postgres repo, Gov App backend"),
        ("Data 626 / MCI", "Metadata supply — feeds OMS; Navigator surfaces"),
        ("Privacy / compliance eng", "Consent, DSR, tracker tools — Portal integrates"),
        ("Inventory program", "Kronos UI on SISU Postgres"),
        ("Compliance / security eng", "Scanning, policy engines, monitoring services"),
    ]
    col_w = PInches(6.04)
    for i, (team, scope) in enumerate(teams):
        col, row = i % 2, i // 2
        left = PInches(0.5) + col * PInches(6.3)
        y = top + row * PInches(1.35)
        deet_rect(slide, left, y, col_w, PInches(1.15), DEET_WHITE, line=DEET_BORDER, radius=True)
        tb = slide.shapes.add_textbox(left + PInches(0.18), y + PInches(0.1), col_w - PInches(0.3), PInches(0.95))
        tf = tb.text_frame
        tf.word_wrap = True
        deet_para(tf.paragraphs[0], team, size=12, color=DEET_NAVY, bold=True, font="Calibri Light")
        deet_para(tf.add_paragraph(), scope, size=10, color=DEET_SLATE)
    deet_callout(
        slide,
        PInches(6.05),
        "Escalation: metadata WG when Foundations & Navigator scope collide · Open: OMS reads governance repo directly or via Gov API?",
        height=PInches(0.62),
    )

    slide, top = deet_content_slide(prs, "Summary — four answers in one view")
    answers = [
        ("1 Personas", "Privacy, stewards, access, partners, inventory, audit — plus exec overview. Consumers & engineers use Navigator / OMS.", DEET_TEAL),
        ("2 Separate?", "Yes logically — different job, SoR, persona, and program ownership. No second metadata store.", DEET_BLUE),
        ("3 How?", "Shared shell + deep links + API boundaries. 626 → OMS → governance repo. Portal phase 2.", DEET_GREEN),
        ("4 Dev teams", "DG = Portal & Navigator UX · Foundations = OMS · SISU = gov repo/backend · pillar teams own specialist tools.", DEET_AMBER),
    ]
    for i, (q, a, accent) in enumerate(answers):
        left = PInches(0.5) + (i % 2) * PInches(6.3)
        y = top + (i // 2) * PInches(2.35)
        deet_rect(slide, left, y, PInches(6.04), PInches(2.1), DEET_WHITE, line=accent, radius=True)
        tb = slide.shapes.add_textbox(left + PInches(0.2), y + PInches(0.15), PInches(5.65), PInches(1.8))
        tf = tb.text_frame
        tf.word_wrap = True
        deet_para(tf.paragraphs[0], q, size=15, color=DEET_NAVY, bold=True, font="Calibri Light")
        deet_para(tf.add_paragraph(), a, size=10, color=DEET_SLATE)

    build_ui_walkthrough_slides(prs)

    prs.save(DECK_PATH)
    return DECK_PATH


def main():
    memo = build_memo()
    deck = build_deck()
    print(f"Wrote {memo}")
    print(f"Wrote {deck} ({len(Presentation(str(deck)).slides)} slides)")


if __name__ == "__main__":
    main()
