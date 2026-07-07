"""Generate Kronos User Stories as a Word document."""

from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

doc = Document()
doc.add_heading("Kronos User Stories", 0)
doc.add_paragraph(
    "Reverse-engineered from the Sustainable Inventory Management UI. "
    "Organized by persona and feature area."
)

# 1. Discovery & Search
doc.add_heading("1. Discovery & Search", level=1)
table = doc.add_table(rows=1, cols=4)
table.style = "Table Grid"
hdr = table.rows[0].cells
hdr[0].text = "ID"
hdr[1].text = "As a…"
hdr[2].text = "I want to…"
hdr[3].text = "So that…"
for row in [
    ("D1", "Governance Steward", "Search across all vendors, trackers, systems, and use cases in one place", "I can quickly find assets for audit requests without checking multiple spreadsheets"),
    ("D2", "Privacy Analyst", "Filter search results by asset type (vendor, tracker, system, use case)", "I only see the asset types relevant to my investigation"),
    ("D3", "DPM / Engineer", "Use example or suggested queries (e.g., “Which platforms use Adobe’s systems & trackers?”)", "I can discover relationships without knowing exact asset names"),
    ("D4", "Any stakeholder", "See search results with asset name, status, type, owner, and relationship count", "I can decide which result to open without extra clicks"),
]:
    r = table.add_row().cells
    r[0].text, r[1].text, r[2].text, r[3].text = row

# 2. Inventory Views
doc.add_heading("2. Inventory Views", level=1)
table = doc.add_table(rows=1, cols=4)
table.style = "Table Grid"
hdr = table.rows[0].cells
hdr[0].text, hdr[1].text, hdr[2].text, hdr[3].text = "ID", "As a…", "I want to…", "So that…"
for row in [
    ("I1", "Governance Steward", "View a list of all vendors with ID, name, type, category, risk rating, DPA status, owner, and status", "I can review the vendor inventory for audits"),
    ("I2", "Governance Steward", "View a list of all digital trackers with type, platform, purpose, vendor, consent requirement, owner, and status", "I can review the tracker inventory for compliance"),
    ("I3", "Governance Steward", "View a list of all systems with type, business unit, data classification, PII flag, owner, and status", "I can review the system inventory for governance"),
    ("I4", "Governance Steward", "View a list of all business use cases with category, subcategory, legal basis, cross-border flag, owner, and status", "I can review use cases for governance and compliance"),
    ("I5", "Any stakeholder", "Filter inventory lists by status (active, inactive, deprecated, pending, under_review)", "I can focus on active or deprecated assets"),
    ("I6", "Any stakeholder", "Filter vendors by risk rating (low, medium, high, critical)", "I can prioritize high-risk vendors"),
    ("I7", "Any stakeholder", "Filter inventory by owner", "I can see assets for a specific team"),
    ("I8", "Any stakeholder", "Sort inventory tables by ID or name", "I can scan or browse the list more easily"),
    ("I9", "Any stakeholder", "Search/filter within a table using a text field", "I can find assets without leaving the inventory view"),
    ("I10", "Any stakeholder", "Export inventory to CSV (per asset type or full audit)", "I can use the data in audits, reports, or other tools"),
]:
    r = table.add_row().cells
    r[0].text, r[1].text, r[2].text, r[3].text = row

# 3. Asset Detail & Relationships
doc.add_heading("3. Asset Detail & Relationships", level=1)
table = doc.add_table(rows=1, cols=4)
table.style = "Table Grid"
hdr = table.rows[0].cells
hdr[0].text, hdr[1].text, hdr[2].text, hdr[3].text = "ID", "As a…", "I want to…", "So that…"
for row in [
    ("A1", "Privacy Analyst", "Open a vendor, tracker, system, or use case and see its properties", "I can understand what each asset represents"),
    ("A2", "Privacy Analyst", "See related assets (vendors, trackers, systems, use cases) for any asset", "I can trace data flows and dependencies"),
    ("A3", "DPM", "Navigate from a use case to its related vendors, trackers, and systems", "I can assess impact and compliance"),
    ("A4", "DPM", "Navigate from a vendor to its related use cases, trackers, and systems", "I can see where a vendor is used"),
    ("A5", "Legal Counsel", "View metadata (region, status, notes) for any asset", "I can assess regulatory and compliance context"),
    ("A6", "Any stakeholder", "See change history for an asset (created, updated, deboarded)", "I can audit changes over time"),
]:
    r = table.add_row().cells
    r[0].text, r[1].text, r[2].text, r[3].text = row

# 4. Relationship Explorer
doc.add_heading("4. Relationship Explorer", level=1)
table = doc.add_table(rows=1, cols=4)
table.style = "Table Grid"
hdr = table.rows[0].cells
hdr[0].text, hdr[1].text, hdr[2].text, hdr[3].text = "ID", "As a…", "I want to…", "So that…"
for row in [
    ("R1", "Governance Steward", "Select an asset type and then a specific asset to see its relationships", "I can explore dependency graphs"),
    ("R2", "Privacy Analyst", "See relationships grouped by type (vendors, trackers, systems, use cases)", "I can quickly find related assets"),
    ("R3", "DPM", "View a summary of relationship counts (e.g., use cases ↔ vendors, use cases ↔ systems)", "I can understand overall connectivity"),
    ("R4", "Any stakeholder", "Click a related asset to navigate to its detail page", "I can traverse the graph without extra steps"),
]:
    r = table.add_row().cells
    r[0].text, r[1].text, r[2].text, r[3].text = row

# 5. Asset Registration & Onboarding
doc.add_heading("5. Asset Registration & Onboarding", level=1)
table = doc.add_table(rows=1, cols=4)
table.style = "Table Grid"
hdr = table.rows[0].cells
hdr[0].text, hdr[1].text, hdr[2].text, hdr[3].text = "ID", "As a…", "I want to…", "So that…"
for row in [
    ("O1", "Asset Owner (Eng, Legal, etc.)", "Register a new vendor with name, type, category, risk rating, DPA status, contract expiry, owner, and region", "I can add vendors to the inventory"),
    ("O2", "Asset Owner", "Register a new digital tracker with vendor, type, platform, purpose, placement, consent requirement, and owner", "I can add trackers to the inventory"),
    ("O3", "Asset Owner", "Register a new system with vendor, type, business unit, data classification, hosting, PII flag, and owner", "I can add systems to the inventory"),
    ("O4", "Asset Owner", "Register a new business use case with category, subcategory, legal basis, cross-border flag, and owner", "I can add use cases to the inventory"),
    ("O5", "Asset Owner", "Link a new use case to related vendors, systems, and trackers during registration", "I can model relationships in one step"),
    ("O6", "Data Governance", "Have asset registration requests come through JIRA or another request portal", "I can review and approve before adding to inventory"),
    ("O7", "Data Governance", "See the Register Asset flow in the UI even when using a read-only demo", "I can understand how registration works"),
]:
    r = table.add_row().cells
    r[0].text, r[1].text, r[2].text, r[3].text = row

# 6. Asset Deboarding & Lifecycle
doc.add_heading("6. Asset Deboarding & Lifecycle", level=1)
table = doc.add_table(rows=1, cols=4)
table.style = "Table Grid"
hdr = table.rows[0].cells
hdr[0].text, hdr[1].text, hdr[2].text, hdr[3].text = "ID", "As a…", "I want to…", "So that…"
for row in [
    ("L1", "Data Governance", "Deboard (deprecate) an asset when it is no longer in use", "I can keep the inventory accurate"),
    ("L2", "Data Governance", "Confirm before deboarding with a confirmation dialog", "I avoid accidental deboarding"),
    ("L3", "Any stakeholder", "See deboarding recorded in the asset change history", "I can audit lifecycle changes"),
    ("L4", "Data Governance", "Edit an asset's name, description, owner, status, region, and notes", "I can correct or update metadata"),
]:
    r = table.add_row().cells
    r[0].text, r[1].text, r[2].text, r[3].text = row

# 7. Dashboard & Overview
doc.add_heading("7. Dashboard & Overview", level=1)
table = doc.add_table(rows=1, cols=4)
table.style = "Table Grid"
hdr = table.rows[0].cells
hdr[0].text, hdr[1].text, hdr[2].text, hdr[3].text = "ID", "As a…", "I want to…", "So that…"
for row in [
    ("DB1", "Governance Steward", "See counts of vendors, trackers, systems, and use cases (total and active)", "I can monitor inventory size"),
    ("DB2", "Governance Steward", "See counts of legal exceptions and prior audits", "I can track governance workload"),
    ("DB3", "Governance Steward", "See total relationship count (M:M connections)", "I can understand overall connectivity"),
    ("DB4", "Governance Steward", "See recent activity (creates, updates, deboards)", "I can stay current on changes"),
    ("DB5", "Governance Steward", "See a status breakdown (active, inactive, deprecated, pending, under_review)", "I can focus on active or problematic assets"),
    ("DB6", "Governance Steward", "See high-risk vendors (high/critical) in a dedicated section", "I can prioritize review and remediation"),
    ("DB7", "Any stakeholder", "Click a stat card to navigate to the related inventory view", "I can quickly drill down"),
]:
    r = table.add_row().cells
    r[0].text, r[1].text, r[2].text, r[3].text = row

# 8. Reports & Audits
doc.add_heading("8. Reports & Audits", level=1)
table = doc.add_table(rows=1, cols=4)
table.style = "Table Grid"
hdr = table.rows[0].cells
hdr[0].text, hdr[1].text, hdr[2].text, hdr[3].text = "ID", "As a…", "I want to…", "So that…"
for row in [
    ("RP1", "Governance Steward", "View prior audits with name, type, date, auditor, findings, critical findings, status, and remediation due", "I can track audit progress"),
    ("RP2", "Governance Steward", "View a full activity log", "I can audit all changes"),
    ("RP3", "Governance Steward", "Export vendor, tracker, or full audit reports as CSV", "I can share with auditors or use in reports"),
    ("RP4", "Legal Counsel", "See audit status (open, closed, in_progress) for each audit", "I can track compliance status"),
]:
    r = table.add_row().cells
    r[0].text, r[1].text, r[2].text, r[3].text = row

# 9. Legal Exceptions
doc.add_heading("9. Legal Exceptions", level=1)
table = doc.add_table(rows=1, cols=4)
table.style = "Table Grid"
hdr = table.rows[0].cells
hdr[0].text, hdr[1].text, hdr[2].text, hdr[3].text = "ID", "As a…", "I want to…", "So that…"
for row in [
    ("LE1", "Privacy Legal", "View legal exceptions with name, type, related assets, policy, approver, expiry, and justification", "I can manage compliance exceptions"),
    ("LE2", "Privacy Legal", "See exceptions that are past expiry highlighted in red", "I can prioritize remediation"),
    ("LE3", "Governance Steward", "See how many legal exceptions are active", "I can track exception volume"),
]:
    r = table.add_row().cells
    r[0].text, r[1].text, r[2].text, r[3].text = row

# 10. Data & Technical
doc.add_heading("10. Data & Technical", level=1)
table = doc.add_table(rows=1, cols=4)
table.style = "Table Grid"
hdr = table.rows[0].cells
hdr[0].text, hdr[1].text, hdr[2].text, hdr[3].text = "ID", "As a…", "I want to…", "So that…"
for row in [
    ("T1", "Data / Platform Engineer", "Store inventory in a relational database (PostgreSQL, Snowflake, or SQLite)", "I can support multiple backends and environments"),
    ("T2", "Data / Platform Engineer", "Use M:M relationships between use cases and vendors, trackers, and systems", "I can model dependencies correctly"),
    ("T3", "Data / Platform Engineer", "Have a change log for every create, update, and deboard", "I can audit all changes"),
    ("T4", "Data / Platform Engineer", "Generate unique asset IDs (e.g., VND-XXXXXX, TRK-XXXXXX)", "I can avoid duplicates"),
    ("T5", "Data / Platform Engineer", "Seed the database with sample data for demos", "I can demo without manual setup"),
    ("T6", "Data / Platform Engineer", "Export seed data as JSON for embedding in a static HTML file", "I can share a read-only demo without a server"),
]:
    r = table.add_row().cells
    r[0].text, r[1].text, r[2].text, r[3].text = row

# 11. Sharing & Distribution
doc.add_heading("11. Sharing & Distribution", level=1)
table = doc.add_table(rows=1, cols=4)
table.style = "Table Grid"
hdr = table.rows[0].cells
hdr[0].text, hdr[1].text, hdr[2].text, hdr[3].text = "ID", "As a…", "I want to…", "So that…"
for row in [
    ("S1", "Data Governance", "Share a single HTML file that works without a server", "Colleagues can view the demo without setup"),
    ("S2", "Data Governance", "Use the same UI in the shared file as in the full app", "I can demo without explaining differences"),
    ("S3", "Data Governance", "See a clear message when the shared file is read-only (e.g., for Register Asset)", "I know how to register assets in the real system"),
]:
    r = table.add_row().cells
    r[0].text, r[1].text, r[2].text, r[3].text = row

# Summary Epic table
doc.add_page_break()
doc.add_heading("Summary by Epic", level=1)
table = doc.add_table(rows=1, cols=2)
table.style = "Table Grid"
hdr = table.rows[0].cells
hdr[0].text, hdr[1].text = "Epic", "User Stories"
for row in [
    ("Discovery & Search", "D1–D4"),
    ("Inventory Views", "I1–I10"),
    ("Asset Detail & Relationships", "A1–A6"),
    ("Relationship Explorer", "R1–R4"),
    ("Asset Registration & Onboarding", "O1–O7"),
    ("Asset Deboarding & Lifecycle", "L1–L4"),
    ("Dashboard & Overview", "DB1–DB7"),
    ("Reports & Audits", "RP1–RP4"),
    ("Legal Exceptions", "LE1–LE3"),
    ("Data & Technical", "T1–T6"),
    ("Sharing & Distribution", "S1–S3"),
]:
    r = table.add_row().cells
    r[0].text, r[1].text = row

out_path = "/Users/jonathan.fong/Documents/Jon's Team Docs/kronos-ui/Kronos_User_Stories.docx"
doc.save(out_path)
print(f"Saved: {out_path}")
