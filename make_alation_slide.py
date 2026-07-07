from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

# Brand-ish palette
NAVY = RGBColor(0x0B, 0x1F, 0x3A)
INDIGO = RGBColor(0x1B, 0x3A, 0x6B)
BLUE = RGBColor(0x2E, 0x6F, 0xD6)
LIGHT = RGBColor(0xF2, 0xF5, 0xFA)
GREY = RGBColor(0x55, 0x5E, 0x6B)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
CARD_BORDER = RGBColor(0xD8, 0xDF, 0xEA)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
SW, SH = prs.slide_width, prs.slide_height

slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank


def add_box(left, top, width, height, fill=None, line=None, line_w=None):
    shp = slide.shapes.add_shape(1, left, top, width, height)  # rectangle
    shp.shadow.inherit = False
    if fill is None:
        shp.fill.background()
    else:
        shp.fill.solid()
        shp.fill.fore_color.rgb = fill
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line
        shp.line.width = line_w or Pt(1)
    return shp


def set_text(tf, runs, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
             space_after=Pt(6), line_spacing=1.0):
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    first = True
    for item in runs:
        text = item.get("text", "")
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.alignment = item.get("align", align)
        p.space_after = item.get("space_after", space_after)
        p.space_before = item.get("space_before", Pt(0))
        p.line_spacing = item.get("line_spacing", line_spacing)
        r = p.add_run()
        r.text = text
        f = r.font
        f.size = item.get("size", Pt(14))
        f.bold = item.get("bold", False)
        f.italic = item.get("italic", False)
        f.color.rgb = item.get("color", NAVY)
        f.name = "Calibri"


# ---- Header band ----
add_box(0, 0, SW, Inches(1.30), fill=NAVY)
# accent stripe
add_box(0, Inches(1.30), SW, Inches(0.07), fill=BLUE)

title = add_box(Inches(0.55), Inches(0.16), Inches(12.2), Inches(0.66))
set_text(title.text_frame, [
    {"text": "Migrating from Alation \u2192 Business Data Catalog", "size": Pt(29), "bold": True, "color": WHITE},
], anchor=MSO_ANCHOR.MIDDLE, space_after=Pt(0))

sub = add_box(Inches(0.57), Inches(0.80), Inches(12.2), Inches(0.45))
set_text(sub.text_frame, [
    {"text": "Business Data Catalog  \u2022  Consolidating onto an in-house, DEEPT-owned platform \u2014 currently in discovery & requirements",
     "size": Pt(13), "color": RGBColor(0xC9, 0xD6, 0xEC)},
], anchor=MSO_ANCHOR.MIDDLE, space_after=Pt(0))

# ---- Core principle banner ----
cp_top = Inches(1.47)
cp_h = Inches(0.82)
add_box(Inches(0.55), cp_top, Inches(12.23), cp_h, fill=RGBColor(0xEA, 0xF1, 0xFB), line=BLUE, line_w=Pt(1.25))
add_box(Inches(0.55), cp_top, Inches(0.16), cp_h, fill=BLUE)  # left accent
cptf = add_box(Inches(0.95), cp_top, Inches(11.7), cp_h)
set_text(cptf.text_frame, [
    {"text": "CORE PRINCIPLE", "size": Pt(11), "bold": True, "color": BLUE, "space_after": Pt(2)},
    {"text": "Alation can\u2019t automate our catalog at the scale we need. We\u2019re moving to one enterprise-owned, "
             "extensible catalog with automation and governance built in.",
     "size": Pt(15), "bold": True, "color": NAVY, "line_spacing": 1.05, "space_after": Pt(0)},
], anchor=MSO_ANCHOR.MIDDLE)

# ---- Two columns ----
col_top = Inches(2.45)
col_h = Inches(3.75)
col_w = Inches(5.95)
gap = Inches(0.5)
left1 = Inches(0.55)
left2 = left1 + col_w + gap

cards = [
    (left1, "WHY WE'RE DOING THIS", [
        "Legacy Alation leans on heavy manual upkeep \u2014 stale descriptions, limited lineage, an aging catalog",
        "Owning the platform unlocks recent AI advances + direct access to our data sources no vendor can match",
        "One single, extensible, enterprise-owned catalog \u2014 business + technical metadata together on OMS",
        "Deeper customization & automation, aligned with internal metadata & governance standards",
    ]),
    (left2, "HIGH-LEVEL PLAN  (WHERE WE ARE)", [
        "In discovery & requirements now \u2014 documenting the most important Alation features in use so nothing is lost",
        "Working with all stakeholders: stewards, analysts, engineers, governance, leadership",
        "Phased delivery: MVP parity \u2192 desired enhancements",
        "Metadata-first migration with a parallel run before cutover; Alation retired after validation",
    ]),
]

for left, header, bullets in cards:
    # card
    add_box(left, col_top, col_w, col_h, fill=WHITE, line=CARD_BORDER, line_w=Pt(1))
    # header strip
    add_box(left, col_top, col_w, Inches(0.62), fill=INDIGO)
    htf = add_box(left + Inches(0.05), col_top, col_w - Inches(0.1), Inches(0.62))
    set_text(htf.text_frame, [
        {"text": header, "size": Pt(15), "bold": True, "color": WHITE, "align": PP_ALIGN.CENTER},
    ], anchor=MSO_ANCHOR.MIDDLE, space_after=Pt(0))
    # body bullets
    body = add_box(left + Inches(0.35), col_top + Inches(0.80), col_w - Inches(0.7), col_h - Inches(0.95))
    runs = []
    for b in bullets:
        runs.append({"text": "\u25B8  " + b, "size": Pt(14), "color": NAVY,
                     "space_after": Pt(11), "line_spacing": 1.04})
    set_text(body.text_frame, runs, anchor=MSO_ANCHOR.TOP)

# ---- Footer goal bar ----
foot_top = col_top + col_h + Inches(0.18)
add_box(Inches(0.55), foot_top, Inches(12.23), Inches(0.52), fill=LIGHT, line=CARD_BORDER, line_w=Pt(1))
ftf = add_box(Inches(0.75), foot_top, Inches(11.83), Inches(0.52))
set_text(ftf.text_frame, [
    {"text": "Goal:  retain most valuable services  \u2022  evolve past traditional manual catalog  \u2022  migrate to provide minimal disruption",
     "size": Pt(14), "bold": True, "color": INDIGO, "align": PP_ALIGN.CENTER},
], anchor=MSO_ANCHOR.MIDDLE, space_after=Pt(0))

# tag line bottom-right
tag = add_box(Inches(8.0), Inches(6.98), Inches(4.8), Inches(0.32))
set_text(tag.text_frame, [
    {"text": "DnA <> DP Monthly Sync \u2014 June 22, 2026  |  4 min  |  Jon Fong, Data Governance",
     "size": Pt(9.5), "italic": True, "color": GREY, "align": PP_ALIGN.RIGHT},
], anchor=MSO_ANCHOR.MIDDLE, space_after=Pt(0))

# ---- Speaker notes ----
notes = slide.notes_slide.notes_text_frame
notes.text = (
    "CORE PRINCIPLE: Alation doesn't automate our catalog processes to the level we need across all data platforms. "
    "We're migrating off to establish a single, extensible, enterprise-owned catalog with deeper customization, "
    "automation, and alignment with internal metadata and governance standards.\n\n"
    "WHY (~45s): We're moving our business data catalog off Alation onto an in-house Business Data Catalog built on OMS. "
    "Alation is a legacy tool that leans on heavy manual upkeep, so descriptions go stale and lineage is limited. "
    "Owning the platform lets us tap recent AI advances and direct access to our own data that no vendor can match \u2014 "
    "and bring business metadata (glossary, ownership, tags, descriptions) together with the technical metadata already in OMS "
    "to create one single source of truth.\n\n"
    "OPPORTUNITY (~45s): Not a like-for-like rebuild. This is our chance to establish a single, extensible, enterprise-owned "
    "catalog \u2014 using AI to keep it current and make search dramatically faster, with deeper customization aligned to our "
    "governance standards.\n\n"
    "WHERE WE ARE / PLAN (~1.5m): We're in discovery and requirements now, documenting every Alation feature in use so we "
    "don't lose functionality. We're interviewing all stakeholder groups \u2014 stewards, analysts, engineers, governance, leadership. "
    "Then a phased rollout: MVP parity on critical features (search, glossary, lineage, ownership), then full parity, then enhancements. "
    "Migration is metadata-first and low-risk: extract from Alation, validate in the new Business Data Catalog, run both in parallel, "
    "then cut over and retire Alation after validation.\n\n"
    "CLOSE / ASK (~30s): Headline \u2014 no functionality lost, minimal disruption, a modern catalog DEEPT fully owns. "
    "Over the next few weeks we'll reach out for stakeholder time; engaging now is the best way to get your team's needs captured."
)

out = "Alation-to-Data-Navigator-Slide.pptx"
prs.save(out)
print("saved", out)
