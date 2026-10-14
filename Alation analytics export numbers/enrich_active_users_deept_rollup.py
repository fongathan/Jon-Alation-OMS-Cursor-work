#!/usr/bin/env python3
"""Add column N (DEEPT EVP/SVP rollup) to the active-users workbook."""

from __future__ import annotations

import csv
import json
import re
import subprocess
import sys
import tempfile
import threading
from concurrent.futures import ThreadPoolExecutor, as_completed
from pathlib import Path

import openpyxl

BASE = Path(__file__).resolve().parent
XLSX_PATH = BASE / "Active users with catalog activity (visits) in the past 6 months.xlsx"
CSV_PATH = BASE / "Active users with catalog activity (visits) in the past 6 months.csv"
CACHE_PATH = BASE / "active_users.workforce_cache.json"
CORRUPT_CACHE = BASE / "Users with catalog activity (visits) in the past 6 months.workforce_cache.json"
ROLLUP_COLUMN = "deept_evp_svp_rollup"
WORKERS = 1

TONY_ID = "dbc99bfe-b28e-490f-9106-d740b58fcfdd"
CACHE_LOCK = threading.Lock()


def wf(cmd: list[str], timeout: int = 45) -> dict:
    result = subprocess.run(
        ["workforce", "--agent", *cmd],
        capture_output=True,
        text=True,
        check=False,
        timeout=timeout,
    )
    if result.returncode != 0:
        raise RuntimeError(result.stderr.strip() or result.stdout.strip() or "workforce failed")
    payload = json.loads(result.stdout)
    if payload.get("status") != "success":
        raise RuntimeError(json.dumps(payload)[:500])
    return payload["data"]


def atomic_save_cache(cache: dict) -> None:
    with CACHE_LOCK:
        fd, tmp = tempfile.mkstemp(suffix=".json", dir=CACHE_PATH.parent)
        try:
            with open(fd, "w", encoding="utf-8") as fh:
                json.dump(cache, fh, indent=2)
            Path(tmp).replace(CACHE_PATH)
        finally:
            if Path(tmp).exists():
                Path(tmp).unlink(missing_ok=True)


def load_cache() -> dict:
    if CACHE_PATH.exists():
        with CACHE_PATH.open(encoding="utf-8") as fh:
            cache = json.load(fh)
    else:
        cache = {"email_to_uuid": {}, "person": {}, "rollup": {}, "search": {}}

    if not cache.get("email_to_uuid") and CORRUPT_CACHE.exists():
        text = CORRUPT_CACHE.read_text(encoding="utf-8", errors="ignore")
        match = re.search(r'"email_to_uuid"\s*:\s*\{(.*?)\}\s*,\s*"person"', text, re.S)
        if match:
            for email, uuid in re.findall(
                r'"([^"]+)"\s*:\s*(?:"([0-9a-f-]{36})"|null)', match.group(1)
            ):
                cache["email_to_uuid"][email] = uuid or None

    for key in ("email_to_uuid", "person", "rollup", "search"):
        cache.setdefault(key, {})
    return cache


def title_of(person: dict) -> str:
    return (person.get("positionTitle") or person.get("title") or "").strip()


def person_name(person: dict) -> str:
    return (person.get("fullName") or person.get("displayName") or "").strip()


def is_evpsvp(title: str) -> bool:
    upper = title.upper()
    return "EVP" in upper or upper.startswith("SVP") or " SVP" in upper


def format_leader_label(person: dict) -> str:
    title = title_of(person)
    name = person_name(person).split("(")[0].strip()
    level = "EVP" if "EVP" in title.upper() else "SVP"
    return f"{level} {name}"


def get_person(uuid: str, cache: dict) -> dict:
    with CACHE_LOCK:
        cached = cache["person"].get(uuid)
        if cached and cached.get("leaderId") is not None:
            return cached
    person = wf(["person", "get", "--uuid", uuid])
    with CACHE_LOCK:
        cache["person"][uuid] = person
    return person


def resolve_uuid(email: str, display_name: str, cache: dict) -> str | None:
    key = email.strip().lower()
    search_key = f"{email}|{display_name}"
    with CACHE_LOCK:
        if key in cache["email_to_uuid"]:
            return cache["email_to_uuid"][key]
        if search_key in cache["search"]:
            uuid = cache["search"][search_key]
            cache["email_to_uuid"][key] = uuid
            return uuid

    candidates: list[dict] = []
    for query in (email, display_name):
        if not query:
            continue
        data = wf(["search", "query", "--query", query])
        persons = data.get("persons") or []
        if persons:
            candidates = persons
            break

    uuid = None
    if candidates:
        email_lower = key
        exact = [p for p in candidates if (p.get("email") or "").lower() == email_lower]
        if len(exact) == 1:
            uuid = exact[0]["id"]
        elif len(candidates) == 1:
            uuid = candidates[0]["id"]
        else:
            rostr = email_lower.upper()
            rostr_matches = [
                p for p in candidates if (p.get("rostrEmail") or "").upper() == rostr
            ]
            if len(rostr_matches) == 1:
                uuid = rostr_matches[0]["id"]

    with CACHE_LOCK:
        cache["search"][search_key] = uuid
        cache["email_to_uuid"][key] = uuid
    return uuid


def rollup_leader(uuid: str | None, cache: dict) -> str | None:
    """Return nearest EVP/SVP leader; None when Workforce has no match."""
    if not uuid:
        return None
    with CACHE_LOCK:
        if uuid in cache["rollup"]:
            return cache["rollup"][uuid]

    seen: set[str] = set()
    current = uuid
    first_evpsvp: str | None = None

    while current and current not in seen:
        seen.add(current)
        person = get_person(current, cache)
        pid = person.get("id") or current
        ptitle = title_of(person)

        if first_evpsvp is None and is_evpsvp(ptitle):
            first_evpsvp = format_leader_label(person)

        leader_id = person.get("leaderId")
        if not leader_id:
            break
        current = leader_id

    with CACHE_LOCK:
        cache["rollup"][uuid] = first_evpsvp
    return first_evpsvp


def column_value(leader: str | None) -> str:
    if leader:
        return leader
    return "Non-DEEPT"


def enrich(email: str, display_name: str, cache: dict) -> str:
    uuid = resolve_uuid(email, display_name, cache)
    leader = rollup_leader(uuid, cache)
    return column_value(leader)


def write_xlsx(rows: list[tuple[str, str, str]]) -> None:
    wb = openpyxl.load_workbook(XLSX_PATH)
    ws = wb.active
    ws.cell(row=1, column=14, value=ROLLUP_COLUMN)
    email_to_value = {email.lower(): value for email, _, value in rows}
    for row_idx in range(2, ws.max_row + 1):
        email = (ws.cell(row=row_idx, column=3).value or "").strip().lower()
        ws.cell(row=row_idx, column=14, value=email_to_value.get(email, "Non-DEEPT"))
    wb.save(XLSX_PATH)


def write_csv(rows: list[tuple[str, str, str]]) -> None:
    with CSV_PATH.open(newline="", encoding="utf-8") as fh:
        reader = csv.DictReader(fh)
        fieldnames = list(reader.fieldnames or [])
        data = list(reader)
    if ROLLUP_COLUMN not in fieldnames:
        fieldnames.append(ROLLUP_COLUMN)
    lookup = {email.lower(): value for email, _, value in rows}
    for row in data:
        row[ROLLUP_COLUMN] = lookup.get((row.get("user_email") or "").lower(), "Non-DEEPT")
    with CSV_PATH.open("w", newline="", encoding="utf-8") as fh:
        writer = csv.DictWriter(fh, fieldnames=fieldnames, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(data)


def resolve_only(email: str, display_name: str, cache: dict) -> str | None:
    return resolve_uuid(email, display_name, cache)


def main() -> int:
    cache = load_cache()
    with CSV_PATH.open(newline="", encoding="utf-8") as fh:
        source_rows = list(csv.DictReader(fh))

    total = len(source_rows)
    results: list[tuple[str, str, str]] = []

    # Phase 1: resolve any missing Workforce UUIDs sequentially.
    pending_search = [
        row for row in source_rows
        if (row.get("user_email") or "").strip().lower() not in cache["email_to_uuid"]
    ]

    print(
        f"UUID cache hits: {total - len(pending_search)}/{total}; "
        f"searches remaining: {len(pending_search)}",
        flush=True,
    )

    for idx, row in enumerate(pending_search, start=1):
        resolve_only(
            (row.get("user_email") or "").strip(),
            (row.get("display_name") or "").strip(),
            cache,
        )
        if idx % 25 == 0 or idx == len(pending_search):
            atomic_save_cache(cache)
            print(f"Resolved UUIDs {total - len(pending_search) + idx}/{total}", flush=True)

    # Phase 2: sequential org rollups (Workforce person/get is more reliable serially).
    for idx, row in enumerate(source_rows, start=1):
        email = (row.get("user_email") or "").strip()
        display_name = (row.get("display_name") or "").strip()
        uuid = cache["email_to_uuid"].get(email.lower())
        try:
            if uuid and uuid in cache["rollup"]:
                value = column_value(cache["rollup"][uuid])
            else:
                value = column_value(rollup_leader(uuid, cache))
        except Exception as exc:  # noqa: BLE001
            print(f"Rollup failed for {email}: {exc}", flush=True)
            value = "Non-DEEPT"
        results.append((email, display_name, value))
        if idx % 50 == 0 or idx == total:
            atomic_save_cache(cache)
            print(f"Rolled up {idx}/{total}", flush=True)

    write_xlsx(results)
    write_csv(results)
    atomic_save_cache(cache)

    counts: dict[str, int] = {}
    for _, _, value in results:
        counts[value] = counts.get(value, 0) + 1

    print(f"Updated {XLSX_PATH.name} column N ({ROLLUP_COLUMN})")
    for label, count in sorted(counts.items(), key=lambda item: (-item[1], item[0])):
        print(f"  {count:4d}  {label}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
